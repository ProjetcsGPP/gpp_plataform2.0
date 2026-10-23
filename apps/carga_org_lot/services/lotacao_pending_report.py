"""Relatório objetivo das pendências de uma importação de lotação."""

from __future__ import annotations

from collections import Counter
from io import BytesIO

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from apps.carga_org_lot.models import StatusCasoValidacao


ISSUE_LABELS = {
    "cpf_digitos_verificadores_invalidos": "CPF inválido",
    "cpf_invalido": "CPF inválido",
    "cpf_ausente": "CPF não informado",
    "cpf_duplicado_planilha": "CPF repetido no mesmo órgão e unidade",
    "nome_servidor_ausente": "Nome do servidor não informado",
    "cargo_ausente": "Cargo não informado",
    "cargo_correspondencia_sugerida": "Cargo precisa de confirmação",
    "orgao_ausente": "Órgão não informado",
    "orgao_invalido": "Órgão não encontrado no organograma",
    "unidade_ausente": "Unidade não informada",
    "unidade_invalida": "Unidade não encontrada no organograma",
    "unidade_ambigua": "Unidade encontrada mais de uma vez",
    "nome_orgao_divergente": "Órgão divergente do organograma",
    "nome_unidade_divergente": "Unidade divergente do organograma",
    "sigla_orgao_muito_longa": "Sigla do órgão acima do tamanho permitido",
    "sigla_unidade_muito_longa": "Sigla da unidade acima do tamanho permitido",
    "cargo_muito_longo": "Cargo acima do tamanho permitido",
    "nome_orgao_muito_longo": "Nome do órgão acima do tamanho permitido",
    "nome_servidor_muito_longo": "Nome do servidor acima do tamanho permitido",
    "nome_unidade_muito_longo": "Nome da unidade acima do tamanho permitido",
}

CORRECTION_GUIDANCE = {
    "cpf_digitos_verificadores_invalidos": "Corrija o CPF informado.",
    "cpf_invalido": "Corrija o CPF informado.",
    "cpf_ausente": "Informe o CPF do servidor.",
    "cpf_duplicado_planilha": (
        "Confira os vínculos repetidos e corrija ou remova o que estiver incorreto."
    ),
    "nome_servidor_ausente": "Informe o nome do servidor.",
    "cargo_ausente": "Informe o cargo do servidor.",
    "cargo_correspondencia_sugerida": "Confirme e padronize a descrição do cargo.",
    "orgao_ausente": "Informe a sigla do órgão existente no organograma.",
    "orgao_invalido": "Use a sigla de um órgão existente no organograma.",
    "unidade_ausente": "Informe a sigla da unidade existente no organograma.",
    "unidade_invalida": "Use a sigla de uma unidade existente no organograma.",
    "unidade_ambigua": "Identifique sem ambiguidade a unidade correta.",
    "nome_orgao_divergente": "Use a sigla do órgão cadastrada no organograma.",
    "nome_unidade_divergente": "Use a sigla da unidade cadastrada no organograma.",
    "sigla_orgao_muito_longa": "Reduza a sigla do órgão para até 20 caracteres.",
    "sigla_unidade_muito_longa": (
        "Reduza a sigla da unidade para até 25 caracteres."
    ),
    "cargo_muito_longo": "Reduza a descrição do cargo para até 100 caracteres.",
    "nome_orgao_muito_longo": (
        "Reduza o nome do órgão para até 500 caracteres."
    ),
    "nome_servidor_muito_longo": (
        "Reduza o nome do servidor para até 255 caracteres."
    ),
    "nome_unidade_muito_longo": (
        "Reduza o nome da unidade para até 500 caracteres."
    ),
}


def _safe_cell(value):
    """Evita que conteúdo vindo da planilha seja interpretado como fórmula."""
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return f"'{value}"
    return value


def _friendly_label(code: str) -> str:
    return ISSUE_LABELS.get(code, code.replace("_", " ").capitalize())


def _guidance(code: str, fallback: str) -> str:
    return CORRECTION_GUIDANCE.get(code, fallback)


def _pending_rows(importation) -> list[dict]:
    cases = (
        importation.relatorio_validacao.casos.filter(
            strstatus=StatusCasoValidacao.PENDENTE
        )
        .prefetch_related("ocorrencias")
        .order_by("strseveridade", "idcasovalidacao")
    )
    rows = []
    for case in cases:
        for occurrence in case.ocorrencias.all():
            rows.append(
                {
                    "severidade": occurrence.get_strseveridade_display(),
                    "severidade_codigo": occurrence.strseveridade,
                    "aba": occurrence.straba,
                    "linha": occurrence.intlinha,
                    "campo": occurrence.strcampo,
                    "valor": occurrence.strvalororiginal,
                    "problema": _friendly_label(case.strcodigo),
                    "orientacao": _guidance(
                        case.strcodigo,
                        occurrence.strmensagem,
                    ),
                }
            )
    severity_order = {"ERRO": 0, "AVISO": 1}
    return sorted(
        rows,
        key=lambda row: (
            severity_order.get(row["severidade_codigo"], 2),
            row["aba"],
            row["linha"] or 0,
            row["problema"],
        ),
    )


def build_pending_lotation_report_xlsx(importation) -> tuple[bytes, dict]:
    """Gera XLSX apenas com casos ainda pendentes e suas ocorrências."""
    rows = _pending_rows(importation)
    counts = Counter(row["severidade_codigo"] for row in rows)

    workbook = Workbook()
    summary = workbook.active
    summary.title = "RESUMO"
    summary_rows = [
        ("Prefeitura/autarquia", importation.idpatriarca.strnome),
        ("Arquivo analisado", importation.strnomearquivo),
        ("Referência", importation.strreferencia),
        ("Aba importada", importation.strabaimportada or "Todas as abas"),
        (
            "Data da importação",
            timezone.localtime(importation.datimportacao).strftime(
                "%d/%m/%Y %H:%M"
            ),
        ),
        ("Total de pendências", len(rows)),
        ("Erros", counts.get("ERRO", 0)),
        ("Avisos", counts.get("AVISO", 0)),
        (
            "Orientação",
            "Corrija as linhas indicadas e envie uma nova versão da planilha.",
        ),
    ]
    for label, value in summary_rows:
        summary.append([label, _safe_cell(value)])
    summary.column_dimensions["A"].width = 26
    summary.column_dimensions["B"].width = 80
    for cell in summary["A"]:
        cell.font = Font(bold=True)
    for row in summary.iter_rows():
        row[1].alignment = Alignment(wrap_text=True, vertical="top")

    pending = workbook.create_sheet("PENDENCIAS")
    headers = [
        "Severidade",
        "Aba",
        "Linha",
        "Campo",
        "Valor informado",
        "Problema",
        "Ação solicitada",
    ]
    pending.append(headers)
    for cell in pending[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in rows:
        pending.append(
            [
                row["severidade"],
                _safe_cell(row["aba"]),
                row["linha"],
                _safe_cell(row["campo"]),
                _safe_cell(row["valor"]),
                _safe_cell(row["problema"]),
                _safe_cell(row["orientacao"]),
            ]
        )
    pending.freeze_panes = "A2"
    pending.auto_filter.ref = pending.dimensions
    widths = [14, 25, 10, 24, 35, 42, 65]
    for index, width in enumerate(widths, start=1):
        pending.column_dimensions[pending.cell(1, index).column_letter].width = width
    for row in pending.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    output = BytesIO()
    workbook.save(output)
    return output.getvalue(), {
        "total": len(rows),
        "erros": counts.get("ERRO", 0),
        "avisos": counts.get("AVISO", 0),
    }
