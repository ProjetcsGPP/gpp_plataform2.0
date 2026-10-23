"""Ciclo de vida, auditoria e exportação dos conjuntos de lotação."""

from __future__ import annotations

from io import BytesIO

from django.db import transaction
from django.utils import timezone
from openpyxl import Workbook
from rest_framework import serializers

from apps.carga_org_lot.models import (
    ConjuntoLotacao,
    HistoricoConjuntoLotacao,
    ImportacaoPlanilha,
    ItemLotacao,
    ItemLotacaoImportado,
    SituacaoConjuntoLotacao,
    StatusCasoValidacao,
    StatusImportacaoLocal,
    StatusItemLotacaoImportado,
)
from apps.carga_org_lot.services.workflow import AuditActor


TERMINAL_SET_STATUSES = {
    SituacaoConjuntoLotacao.CANCELADO,
    SituacaoConjuntoLotacao.SUBSTITUIDO,
}
IMPORT_CLOSED_SET_STATUSES = TERMINAL_SET_STATUSES | {
    SituacaoConjuntoLotacao.HOMOLOGADO,
    SituacaoConjuntoLotacao.ATIVO,
}
PENDING_REVIEW_IMPORT_STATUSES = {
    StatusImportacaoLocal.AGUARDANDO_REVISAO,
    StatusImportacaoLocal.AGUARDANDO_REVISAO_CARGOS,
}
IMPORT_IN_PROGRESS_STATUSES = {
    StatusImportacaoLocal.RECEBIDA,
    StatusImportacaoLocal.PROCESSANDO,
    StatusImportacaoLocal.ANALISANDO_CARGOS,
    StatusImportacaoLocal.PRONTA_PARA_INCORPORAR,
    StatusImportacaoLocal.INCORPORANDO,
}


def _audit_fields(actor: AuditActor) -> dict:
    return {
        "created_by_id": actor.user_id,
        "created_by_name": actor.user_name,
        "updated_by_id": actor.user_id,
        "updated_by_name": actor.user_name,
    }


def _record_history(
    conjunto: ConjuntoLotacao,
    *,
    action: str,
    previous_status: str,
    actor: AuditActor,
    message: str = "",
    details: dict | None = None,
) -> HistoricoConjuntoLotacao:
    return HistoricoConjuntoLotacao.objects.create(
        idconjuntolotacao=conjunto,
        stracao=action,
        strsituacaoanterior=previous_status,
        strsituacaonova=conjunto.strsituacao,
        strmensagem=message,
        jsondetalhes=details or {},
        **_audit_fields(actor),
    )


def validate_set_organogram(patriarca, organograma: ImportacaoPlanilha) -> None:
    if organograma.idtipocarga_id != 3:
        raise serializers.ValidationError(
            {
                "idimportacaoorganograma_id": (
                    "A referência informada não é um organograma."
                )
            }
        )
    if organograma.idpatriarca_id != patriarca.pk:
        raise serializers.ValidationError(
            {
                "idimportacaoorganograma_id": (
                    "O organograma deve pertencer ao mesmo patriarca do conjunto."
                )
            }
        )
    if organograma.strstatusimportacao == StatusImportacaoLocal.REJEITADA:
        raise serializers.ValidationError(
            {
                "idimportacaoorganograma_id": (
                    "Um organograma rejeitado não pode ser usado."
                )
            }
        )


@transaction.atomic
def create_lotation_set(
    *,
    patriarca,
    organograma: ImportacaoPlanilha,
    reference: str,
    actor: AuditActor,
) -> ConjuntoLotacao:
    validate_set_organogram(patriarca, organograma)
    conjunto = ConjuntoLotacao.objects.create(
        idpatriarca=patriarca,
        idimportacaoorganograma=organograma,
        strreferencia=reference,
        **_audit_fields(actor),
    )
    _record_history(
        conjunto,
        action="CONJUNTO_CRIADO",
        previous_status="",
        actor=actor,
        message="Conjunto de lotação criado.",
    )
    return conjunto


def ensure_set_accepts_import(conjunto: ConjuntoLotacao) -> None:
    if conjunto.strsituacao in IMPORT_CLOSED_SET_STATUSES:
        raise serializers.ValidationError(
            {
                "idconjuntolotacao_id": (
                    "O conjunto homologado, ativo, cancelado ou substituído está "
                    "fechado para novos arquivos. Crie uma nova alternativa."
                )
            }
        )


def available_set_actions(conjunto: ConjuntoLotacao) -> dict:
    """Expõe a elegibilidade de negócio sem substituir a autorização."""
    closed_for_imports = conjunto.strsituacao in IMPORT_CLOSED_SET_STATUSES
    terminal = conjunto.strsituacao in TERMINAL_SET_STATUSES
    return {
        "adicionar_planilha": {
            "permitida": not closed_for_imports,
            "motivo": (
                ""
                if not closed_for_imports
                else "O conjunto está fechado para novas planilhas."
            ),
        },
        "homologar": {
            "permitida": (
                conjunto.strsituacao
                == SituacaoConjuntoLotacao.PRONTO_PARA_HOMOLOGACAO
            ),
            "motivo": (
                ""
                if conjunto.strsituacao
                == SituacaoConjuntoLotacao.PRONTO_PARA_HOMOLOGACAO
                else "O conjunto ainda não está pronto para homologação."
            ),
        },
        "ativar": {
            "permitida": conjunto.strsituacao == SituacaoConjuntoLotacao.HOMOLOGADO,
            "motivo": (
                ""
                if conjunto.strsituacao == SituacaoConjuntoLotacao.HOMOLOGADO
                else "Somente um conjunto homologado pode ser ativado."
            ),
        },
        "cancelar": {
            "permitida": not terminal,
            "motivo": "" if not terminal else "O conjunto já está encerrado.",
        },
    }


@transaction.atomic
def register_set_import(
    conjunto: ConjuntoLotacao,
    importacao: ImportacaoPlanilha,
    *,
    actor: AuditActor,
) -> ConjuntoLotacao:
    _record_history(
        conjunto,
        action="PLANILHA_IMPORTADA",
        previous_status=conjunto.strsituacao,
        actor=actor,
        message=f'Planilha "{importacao.strnomearquivo}" adicionada ao conjunto.',
        details={
            "idimportacaoplanilha": importacao.pk,
            "tipo_operacao": importacao.strtipooperacaolotacao,
            "status_importacao": importacao.strstatusimportacao,
        },
    )
    return sync_lotation_set(
        conjunto,
        actor=actor,
        action="IMPORTACAO_PROCESSADA",
    )


@transaction.atomic
def sync_lotation_set(
    conjunto: ConjuntoLotacao,
    *,
    actor: AuditActor,
    action: str = "SITUACAO_RECALCULADA",
) -> ConjuntoLotacao:
    conjunto = ConjuntoLotacao.objects.select_for_update().get(pk=conjunto.pk)
    if conjunto.strsituacao in TERMINAL_SET_STATUSES:
        return conjunto
    if conjunto.strsituacao == SituacaoConjuntoLotacao.ATIVO:
        return conjunto

    imports = conjunto.importacoes.all()
    previous = conjunto.strsituacao
    if not imports.exists():
        new_status = SituacaoConjuntoLotacao.RASCUNHO
    elif imports.filter(
        strstatusimportacao__in=PENDING_REVIEW_IMPORT_STATUSES
    ).exists():
        new_status = SituacaoConjuntoLotacao.COM_PENDENCIAS
    elif imports.filter(
        strstatusimportacao__in=IMPORT_IN_PROGRESS_STATUSES
    ).exists():
        new_status = SituacaoConjuntoLotacao.EM_VALIDACAO
    elif ItemLotacao.objects.filter(
        idimportacaoplanilha__idconjuntolotacao=conjunto
    ).exists():
        new_status = SituacaoConjuntoLotacao.PRONTO_PARA_HOMOLOGACAO
    else:
        new_status = SituacaoConjuntoLotacao.EM_VALIDACAO

    if previous != new_status:
        conjunto.strsituacao = new_status
        conjunto.updated_by_id = actor.user_id
        conjunto.updated_by_name = actor.user_name
        conjunto.save()
        _record_history(
            conjunto,
            action=action,
            previous_status=previous,
            actor=actor,
            message="Situação atualizada conforme as importações do conjunto.",
        )
    return conjunto


@transaction.atomic
def homologate_lotation_set(
    conjunto: ConjuntoLotacao,
    *,
    justification: str,
    actor: AuditActor,
) -> ConjuntoLotacao:
    conjunto = ConjuntoLotacao.objects.select_for_update().get(pk=conjunto.pk)
    sync_lotation_set(conjunto, actor=actor)
    conjunto.refresh_from_db()
    if conjunto.strsituacao != SituacaoConjuntoLotacao.PRONTO_PARA_HOMOLOGACAO:
        raise serializers.ValidationError(
            "O conjunto ainda não está pronto para homologação."
        )
    if not justification.strip():
        raise serializers.ValidationError(
            {"justificativa": "Informe a justificativa da homologação."}
        )
    previous = conjunto.strsituacao
    conjunto.strsituacao = SituacaoConjuntoLotacao.HOMOLOGADO
    conjunto.strjustificativaultimaacao = justification.strip()
    conjunto.dathomologacao = timezone.now()
    conjunto.updated_by_id = actor.user_id
    conjunto.updated_by_name = actor.user_name
    conjunto.save()
    _record_history(
        conjunto,
        action="CONJUNTO_HOMOLOGADO",
        previous_status=previous,
        actor=actor,
        message=justification.strip(),
    )
    return conjunto


@transaction.atomic
def activate_lotation_set(
    conjunto: ConjuntoLotacao,
    *,
    justification: str,
    actor: AuditActor,
) -> ConjuntoLotacao:
    conjunto = ConjuntoLotacao.objects.select_for_update().get(pk=conjunto.pk)
    if conjunto.strsituacao != SituacaoConjuntoLotacao.HOMOLOGADO:
        raise serializers.ValidationError(
            "Somente um conjunto homologado pode ser ativado."
        )
    if not justification.strip():
        raise serializers.ValidationError(
            {"justificativa": "Informe a justificativa da ativação."}
        )
    ImportacaoPlanilha.objects.select_for_update().get(
        pk=conjunto.idimportacaoorganograma_id
    )
    current = (
        ConjuntoLotacao.objects.select_for_update()
        .filter(
            idimportacaoorganograma=conjunto.idimportacaoorganograma,
            strsituacao=SituacaoConjuntoLotacao.ATIVO,
        )
        .exclude(pk=conjunto.pk)
        .first()
    )
    if current is not None:
        old_status = current.strsituacao
        current.strsituacao = SituacaoConjuntoLotacao.SUBSTITUIDO
        current.strjustificativaultimaacao = justification.strip()
        current.updated_by_id = actor.user_id
        current.updated_by_name = actor.user_name
        current.save()
        _record_history(
            current,
            action="CONJUNTO_SUBSTITUIDO",
            previous_status=old_status,
            actor=actor,
            message=justification.strip(),
            details={"idconjuntosubstituto": conjunto.pk},
        )
        conjunto.idconjuntoanterior = current

    previous = conjunto.strsituacao
    conjunto.strsituacao = SituacaoConjuntoLotacao.ATIVO
    conjunto.strjustificativaultimaacao = justification.strip()
    conjunto.datativacao = timezone.now()
    conjunto.updated_by_id = actor.user_id
    conjunto.updated_by_name = actor.user_name
    conjunto.save()
    _record_history(
        conjunto,
        action="CONJUNTO_ATIVADO",
        previous_status=previous,
        actor=actor,
        message=justification.strip(),
        details={"idconjuntoanterior": conjunto.idconjuntoanterior_id},
    )
    return conjunto


@transaction.atomic
def cancel_lotation_set(
    conjunto: ConjuntoLotacao,
    *,
    justification: str,
    actor: AuditActor,
) -> ConjuntoLotacao:
    conjunto = ConjuntoLotacao.objects.select_for_update().get(pk=conjunto.pk)
    if conjunto.strsituacao in TERMINAL_SET_STATUSES:
        raise serializers.ValidationError("O conjunto já está encerrado.")
    if not justification.strip():
        raise serializers.ValidationError(
            {"justificativa": "Informe a justificativa do cancelamento."}
        )
    previous = conjunto.strsituacao
    conjunto.strsituacao = SituacaoConjuntoLotacao.CANCELADO
    conjunto.strjustificativaultimaacao = justification.strip()
    conjunto.datcancelamento = timezone.now()
    conjunto.updated_by_id = actor.user_id
    conjunto.updated_by_name = actor.user_name
    conjunto.save()
    _record_history(
        conjunto,
        action="CONJUNTO_CANCELADO",
        previous_status=previous,
        actor=actor,
        message=justification.strip(),
    )
    return conjunto


def build_consolidated_lotation_xlsx(conjunto: ConjuntoLotacao) -> bytes:
    workbook = Workbook()
    active_sheet = workbook.active
    active_sheet.title = "LOTACOES_CONSOLIDADAS"
    active_sheet.append(
        [
            "Importação",
            "Aba",
            "Linha",
            "Sigla do órgão",
            "Nome do órgão",
            "Nome do servidor",
            "CPF",
            "Cargo original",
            "Cargo normalizado",
            "Sigla da unidade",
            "Nome da unidade",
        ]
    )
    items = ItemLotacao.objects.filter(
        idimportacaoplanilha__idconjuntolotacao=conjunto
    ).order_by(
        "idimportacaoplanilha_id",
        "strabaorigem",
        "intlinhaorigem",
        "iditemlotacao",
    )
    for item in items:
        active_sheet.append(
            [
                item.idimportacaoplanilha_id,
                item.strabaorigem,
                item.intlinhaorigem,
                item.strsiglaorgao,
                item.strnomeorgao,
                item.strnomeservidor,
                item.strcpf,
                item.strcargooriginal,
                item.strcargonormalizado,
                item.strsiglaunidade,
                item.strnomeunidade,
            ]
        )

    discarded_sheet = workbook.create_sheet("REGISTROS_DESCARTADOS")
    discarded_sheet.append(
        ["Importação", "Aba", "Linha", "Servidor", "CPF", "Cargo", "Situação"]
    )
    discarded = ItemLotacaoImportado.objects.filter(
        idimportacaoplanilha__idconjuntolotacao=conjunto,
        strsituacao=StatusItemLotacaoImportado.DESCARTADO,
    ).order_by("idimportacaoplanilha_id", "strabaorigem", "intlinhaorigem")
    for item in discarded:
        discarded_sheet.append(
            [
                item.idimportacaoplanilha_id,
                item.strabaorigem,
                item.intlinhaorigem,
                item.strnomeservidor,
                item.strcpf,
                item.strcargooriginal,
                item.strsituacao,
            ]
        )

    summary = workbook.create_sheet("RESUMO")
    pending_cases = sum(
        importacao.relatorio_validacao.casos.filter(
            strstatus=StatusCasoValidacao.PENDENTE
        ).count()
        for importacao in conjunto.importacoes.select_related(
            "relatorio_validacao"
        )
        if hasattr(importacao, "relatorio_validacao")
    )
    summary.append(["Conjunto", conjunto.pk])
    summary.append(["Referência", conjunto.strreferencia])
    summary.append(["Situação", conjunto.strsituacao])
    summary.append(["Patriarca", conjunto.idpatriarca.strnome])
    summary.append(["Organograma", conjunto.idimportacaoorganograma_id])
    summary.append(["Importações", conjunto.importacoes.count()])
    summary.append(["Registros consolidados", items.count()])
    summary.append(["Registros descartados", discarded.count()])
    summary.append(["Casos pendentes", pending_cases])

    history_sheet = workbook.create_sheet("HISTORICO")
    history_sheet.append(
        ["Data", "Ação", "Situação anterior", "Situação nova", "Usuário", "Mensagem"]
    )
    for history in conjunto.historico.all():
        history_sheet.append(
            [
                history.created_at.isoformat(),
                history.stracao,
                history.strsituacaoanterior,
                history.strsituacaonova,
                history.created_by_name,
                history.strmensagem,
            ]
        )

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
