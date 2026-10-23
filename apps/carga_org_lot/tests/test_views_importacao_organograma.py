"""Testes de views: importacao organograma."""

import hashlib
from io import BytesIO

import pytest
from openpyxl import load_workbook

from apps.carga_org_lot.models import (CasoValidacaoImportacao,
                                       FaseProcessoCargaInicial,
                                       ImportacaoPlanilha, ItemOrganograma,
                                       RelatorioValidacaoImportacao,
                                       ResultadoRelatorioValidacao,
                                       StatusImportacaoLocal)
from apps.carga_org_lot.tests.view_helpers import (IMPORTACOES_URL,
                                                   ITENS_ORGANOGRAMA_URL,
                                                   arquivo_planilha,
                                                   conteudo_organograma)


@pytest.mark.django_db(transaction=True)
class TestImportacaoOrganograma:
    def test_cria_organograma_inicial_e_importa_o_xlsx_gerado(
        self,
        client_gestor_carga_lot,
        gestor_carga_lot,
        patriarca,
        tipo_organograma,
    ):
        response = client_gestor_carga_lot.post(
            f"{IMPORTACOES_URL}criar-organograma-inicial/",
            {
                "idpatriarca_id": patriarca.pk,
                "strreferencia": "Estrutura inicial padrão",
            },
            format="json",
        )

        assert response.status_code == 201
        importacao_id = response.data["idimportacaoplanilha"]
        assert response.data["strnomearquivo"] == ("organograma_inicial_PMVV.xlsx")
        assert response.data["strreferencia"] == "Estrutura inicial padrão"
        assert response.data["strstatusimportacao"] == (
            StatusImportacaoLocal.INCORPORADA
        )
        assert response.data["flgpossuialteracoes"] is False
        assert response.data["flgcriarunidadesespelho"] is False
        assert response.data["total_itens_organograma"] == 9
        assert response.data["created_by_id"] == gestor_carga_lot.pk
        assert response.data["relatorio_validacao"]["inttotallinhas"] == 9
        assert response.data["relatorio_validacao"]["inttotalerros"] == 0

        importacao = ImportacaoPlanilha.objects.get(pk=importacao_id)
        assert importacao.idtipocarga_id == tipo_organograma.pk
        assert hashlib.sha256(bytes(importacao.binconteudo)).hexdigest() == (
            importacao.strhasharquivo
        )
        items = list(
            importacao.itens_organograma.select_related("idpai").order_by("strnumero")
        )
        assert len(items) == 9
        assert not any(item.flggeradoautomaticamente for item in items)
        root = next(item for item in items if item.idpai_id is None)
        assert (root.strnome, root.strsigla) == (
            "PREFEITURA MUNICIPAL DE VILA VELHA",
            "PMVV",
        )
        for acronym in ("GAB", "SEFAZ", "SESA", "SEMED"):
            organ = next(
                item
                for item in items
                if item.strsigla == acronym and item.idpai_id == root.pk
            )
            mirror = next(
                item
                for item in items
                if item.strsigla == acronym and item.idpai_id == organ.pk
            )
            assert mirror.strnome == organ.strnome

        download = client_gestor_carga_lot.get(
            f"{IMPORTACOES_URL}{importacao_id}/download/"
        )
        assert download.status_code == 200
        assert download.content == bytes(importacao.binconteudo)
        workbook = load_workbook(BytesIO(download.content), data_only=True)
        assert workbook.sheetnames == ["Organograma"]
        assert list(workbook["Organograma"].values) == [
            ("Numero", "Nome", "Sigla", "Pai_Sigla", "Responsavel", "Cargo"),
            (
                "1",
                "Prefeitura Municipal de Vila Velha",
                "PMVV",
                None,
                None,
                None,
            ),
            ("1.1", "Gabinete do Prefeito", "GAB", "PMVV", None, None),
            ("1.1.1", "Gabinete do Prefeito", "GAB", "GAB", None, None),
            (
                "1.2",
                "Secretaria Municipal de Fazenda",
                "SEFAZ",
                "PMVV",
                None,
                None,
            ),
            (
                "1.2.1",
                "Secretaria Municipal de Fazenda",
                "SEFAZ",
                "SEFAZ",
                None,
                None,
            ),
            (
                "1.3",
                "Secretaria Municipal de Saúde",
                "SESA",
                "PMVV",
                None,
                None,
            ),
            (
                "1.3.1",
                "Secretaria Municipal de Saúde",
                "SESA",
                "SESA",
                None,
                None,
            ),
            (
                "1.4",
                "Secretaria Municipal de Educação",
                "SEMED",
                "PMVV",
                None,
                None,
            ),
            (
                "1.4.1",
                "Secretaria Municipal de Educação",
                "SEMED",
                "SEMED",
                None,
                None,
            ),
        ]

    def test_criar_organograma_inicial_exige_patriarca_valido(
        self,
        client_gestor_carga_lot,
    ):
        response = client_gestor_carga_lot.post(
            f"{IMPORTACOES_URL}criar-organograma-inicial/",
            {"idpatriarca_id": 999999},
            format="json",
        )

        assert response.status_code == 400
        assert "idpatriarca_id" in response.data["errors"]
        assert not ImportacaoPlanilha.objects.exists()

    def test_upload_lista_detalha_e_download(
        self,
        client_gestor_carga_lot,
        gestor_carga_lot,
        patriarca,
        tipo_organograma,
    ):
        conteudo = conteudo_organograma()
        upload_response = client_gestor_carga_lot.post(
            f"{IMPORTACOES_URL}upload/",
            {
                "idpatriarca_id": patriarca.pk,
                "idtipocarga_id": tipo_organograma.pk,
                "strreferencia": "Processo E-Docs 2026-ABC123",
                "arquivo": arquivo_planilha(conteudo=conteudo),
            },
            format="multipart",
        )

        assert upload_response.status_code == 201
        pk = upload_response.data["idimportacaoplanilha"]
        assert upload_response.data["strnomearquivo"] == "organograma.xlsx"
        assert upload_response.data["inttamanhoarquivo"] == len(conteudo)
        assert (
            upload_response.data["strhasharquivo"]
            == hashlib.sha256(conteudo).hexdigest()
        )
        assert upload_response.data["status"]["strdescricao"] == "Novo"
        assert upload_response.data["strreferencia"] == ("Processo E-Docs 2026-ABC123")
        assert upload_response.data["created_by_id"] == gestor_carga_lot.pk
        assert "binconteudo" not in upload_response.data
        assert (
            upload_response.data["strstatusimportacao"]
            == StatusImportacaoLocal.INCORPORADA
        )
        assert (
            upload_response.data["processo"]["strfaseatual"]
            == FaseProcessoCargaInicial.PREPARANDO_ORGANOGRAMA
        )
        assert (
            upload_response.data["relatorio_validacao"]["strresultado"]
            == ResultadoRelatorioValidacao.VALIDO
        )

        importacao = ImportacaoPlanilha.objects.get(pk=pk)
        assert bytes(importacao.binconteudo) == conteudo
        assert importacao.relatorio_validacao.inttotalerros == 0

        list_response = client_gestor_carga_lot.get(IMPORTACOES_URL)
        assert list_response.status_code == 200
        assert list_response.data["count"] == 1
        listed_import = list_response.data["results"][0]
        assert "binconteudo" not in listed_import
        report_summary = listed_import["relatorio_validacao"]
        assert report_summary["strresultado"] == ResultadoRelatorioValidacao.VALIDO
        assert "jsonresumo" not in report_summary
        assert "ocorrencias" not in report_summary
        assert "casos" not in report_summary
        assert "historico" not in listed_import["processo"]
        assert "processo_ativo" not in listed_import["patriarca"]

        filtered_response = client_gestor_carga_lot.get(
            IMPORTACOES_URL,
            {"patriarca": patriarca.pk, "tipo_carga": tipo_organograma.pk},
        )
        assert filtered_response.status_code == 200
        assert filtered_response.data["count"] == 1

        empty_response = client_gestor_carga_lot.get(
            IMPORTACOES_URL,
            {"patriarca": patriarca.pk + 1},
        )
        assert empty_response.status_code == 200
        assert empty_response.data["count"] == 0

        detail_response = client_gestor_carga_lot.get(f"{IMPORTACOES_URL}{pk}/")
        assert detail_response.status_code == 200
        assert detail_response.data["patriarca"]["idpatriarca"] == patriarca.pk

        download_response = client_gestor_carga_lot.get(
            f"{IMPORTACOES_URL}{pk}/download/"
        )
        assert download_response.status_code == 200
        assert download_response.content == conteudo
        assert download_response["Content-Length"] == str(len(conteudo))
        assert "organograma.xlsx" in download_response["Content-Disposition"]
        assert download_response["X-Content-Type-Options"] == "nosniff"

        report_response = client_gestor_carga_lot.get(
            f"{IMPORTACOES_URL}{pk}/relatorio/"
        )
        assert report_response.status_code == 200
        assert report_response.data["strresultado"] == "VALIDO"
        assert "jsonresumo" in report_response.data
        assert report_response.data["ocorrencias"] == []

    def test_listagem_nao_expande_casos_de_validacao(
        self,
        client_gestor_carga_lot,
        patriarca,
        tipo_organograma,
        django_assert_max_num_queries,
    ):
        uploaded = client_gestor_carga_lot.post(
            f"{IMPORTACOES_URL}upload/",
            {
                "idpatriarca_id": patriarca.pk,
                "idtipocarga_id": tipo_organograma.pk,
                "arquivo": arquivo_planilha(),
            },
            format="multipart",
        )
        assert uploaded.status_code == 201
        report = RelatorioValidacaoImportacao.objects.get(
            idimportacaoplanilha_id=uploaded.data["idimportacaoplanilha"]
        )
        CasoValidacaoImportacao.objects.bulk_create(
            [
                CasoValidacaoImportacao(
                    idrelatoriovalidacao=report,
                    strchaveagrupamento=f"teste:{index}",
                    strcodigo="caso_teste",
                    strseveridade="AVISO",
                    strtitulo=f"Caso {index}",
                    strmensagem="Caso criado para testar a listagem resumida.",
                )
                for index in range(50)
            ]
        )

        with django_assert_max_num_queries(20):
            response = client_gestor_carga_lot.get(
                IMPORTACOES_URL,
                {
                    "patriarca": patriarca.pk,
                    "tipo_carga": tipo_organograma.pk,
                },
            )

        assert response.status_code == 200
        item = response.data["results"][0]
        assert item["total_casos_pendentes"] == 50
        assert "casos" not in item["relatorio_validacao"]

    def test_rejeita_filtro_nao_numerico(self, client_gestor_carga_lot):
        response = client_gestor_carga_lot.get(
            IMPORTACOES_URL,
            {"patriarca": "qualquer"},
        )
        assert response.status_code == 400
        assert "patriarca" in response.data["errors"]

    def test_post_padrao_tambem_realiza_upload(
        self,
        client_gestor_carga_lot,
        patriarca,
        tipo_organograma,
    ):
        response = client_gestor_carga_lot.post(
            IMPORTACOES_URL,
            {
                "idpatriarca_id": patriarca.pk,
                "idtipocarga_id": tipo_organograma.pk,
                "arquivo": arquivo_planilha(),
            },
            format="multipart",
        )
        assert response.status_code == 201

    def test_rejeita_extensao_invalida(
        self,
        client_gestor_carga_lot,
        patriarca,
        tipo_organograma,
    ):
        response = client_gestor_carga_lot.post(
            f"{IMPORTACOES_URL}upload/",
            {
                "idpatriarca_id": patriarca.pk,
                "idtipocarga_id": tipo_organograma.pk,
                "arquivo": arquivo_planilha(nome="dados.exe"),
            },
            format="multipart",
        )
        assert response.status_code == 400
        assert "arquivo" in response.data["errors"]

    def test_rejeita_arquivo_acima_do_limite(
        self,
        client_gestor_carga_lot,
        patriarca,
        tipo_organograma,
        settings,
    ):
        settings.CARGA_ORG_LOT_MAX_UPLOAD_BYTES = 4
        response = client_gestor_carga_lot.post(
            f"{IMPORTACOES_URL}upload/",
            {
                "idpatriarca_id": patriarca.pk,
                "idtipocarga_id": tipo_organograma.pk,
                "arquivo": arquivo_planilha(conteudo=b"12345"),
            },
            format="multipart",
        )
        assert response.status_code == 400
        assert "arquivo" in response.data["errors"]

    @pytest.mark.parametrize(
        ("cabecalho", "linhas"),
        [
            (
                ("Numero", "Nome", "Sigla", "Pai_Sigla"),
                [
                    (1, "Prefeitura", "PMVV", ""),
                    (2, "Secretaria", "SEC", "PMVV"),
                    (3, "Gerência", "GER", "SEC"),
                ],
            ),
            (
                ("Nome", "Sigla", "Pai_Sigla"),
                [
                    ("Prefeitura", "PMVV", ""),
                    ("Secretaria", "SEC", "PMVV"),
                    ("Gerência", "GER", "SEC"),
                ],
            ),
            (
                (
                    "pode ser isso",
                    "ESTRUTURA ORGANIZACIONAL DO MUNICÍPIO",
                    "SIGLAS ATUAIS",
                ),
                [
                    ("", "Prefeitura", "PMVV"),
                    (1, "Secretaria", "SEC"),
                    ("1.1", "Gerência", "GER"),
                ],
            ),
        ],
    )
    def test_aceita_formatos_flexiveis_de_hierarquia(
        self,
        client_gestor_carga_lot,
        patriarca,
        tipo_organograma,
        cabecalho,
        linhas,
    ):
        response = client_gestor_carga_lot.post(
            f"{IMPORTACOES_URL}upload/",
            {
                "idpatriarca_id": patriarca.pk,
                "idtipocarga_id": tipo_organograma.pk,
                "arquivo": arquivo_planilha(
                    conteudo=conteudo_organograma(linhas, cabecalho)
                ),
            },
            format="multipart",
        )

        assert response.status_code == 201
        assert ImportacaoPlanilha.objects.count() == 1

    @pytest.mark.parametrize(
        "linhas",
        [
            [
                (1, "Prefeitura", "PMVV"),
                ("1.1", "Secretaria de Administração", "SEMAD"),
                ("1.1.1", "Secretaria de Administração", "SEMAD"),
            ],
            [
                (1, "Prefeitura", "PMVV"),
                ("1.1", "Órgão A", "ORGA"),
                ("1.1.1", "Unidade Comum A", "COMUM"),
                ("1.2", "Órgão B", "ORGB"),
                ("1.2.1", "Unidade Comum B", "COMUM"),
            ],
        ],
    )
    def test_aceita_sigla_repetida_com_pai_ou_em_outro_ramo(
        self,
        client_gestor_carga_lot,
        patriarca,
        tipo_organograma,
        linhas,
    ):
        response = client_gestor_carga_lot.post(
            f"{IMPORTACOES_URL}upload/",
            {
                "idpatriarca_id": patriarca.pk,
                "idtipocarga_id": tipo_organograma.pk,
                "arquivo": arquivo_planilha(
                    conteudo=conteudo_organograma(linhas=linhas)
                ),
            },
            format="multipart",
        )

        assert response.status_code == 201

    def test_cria_unidades_espelho_quando_opcao_esta_marcada(
        self,
        client_gestor_carga_lot,
        patriarca,
        tipo_organograma,
    ):
        content = conteudo_organograma(
            linhas=[
                (1, "Prefeitura", "PMVV"),
                ("1.1", "Secretaria de Administração", "SEMAD"),
                ("1.2", "Secretaria de Saúde", "SESA"),
            ]
        )

        response = client_gestor_carga_lot.post(
            f"{IMPORTACOES_URL}upload/",
            {
                "idpatriarca_id": patriarca.pk,
                "idtipocarga_id": tipo_organograma.pk,
                "criar_unidades_espelho": "true",
                "arquivo": arquivo_planilha(conteudo=content),
            },
            format="multipart",
        )

        assert response.status_code == 201
        assert response.data["flgcriarunidadesespelho"] is True
        assert response.data["flgpossuialteracoes"] is True
        assert response.data["strstatusimportacao"] == StatusImportacaoLocal.ALTERADA
        assert response.data["organograma_atualizado_download_url"] is not None
        assert response.data["total_itens_organograma"] == 5
        summary = response.data["relatorio_validacao"]["jsonresumo"]
        assert summary["criar_unidades_espelho"] is True
        assert summary["total_unidades_espelho_criadas"] == 2

        generated = ItemOrganograma.objects.filter(
            idimportacaoplanilha_id=response.data["idimportacaoplanilha"],
            flggeradoautomaticamente=True,
        ).order_by("strnumero")
        assert list(
            generated.values_list(
                "strnome",
                "strsigla",
                "idpai__strsigla",
                "strnumero",
                "strabaorigem",
                "intlinhaorigem",
            )
        ) == [
            (
                "SECRETARIA DE ADMINISTRAÇÃO",
                "SEMAD",
                "SEMAD",
                "1.1.1",
                "Gerada automaticamente",
                0,
            ),
            (
                "SECRETARIA DE SAÚDE",
                "SESA",
                "SESA",
                "1.2.1",
                "Gerada automaticamente",
                0,
            ),
        ]
        items_response = client_gestor_carga_lot.get(
            ITENS_ORGANOGRAMA_URL,
            {"importacao": response.data["idimportacaoplanilha"]},
        )
        assert items_response.status_code == 200
        assert (
            sum(
                item["flggeradoautomaticamente"]
                for item in items_response.data["results"]
            )
            == 2
        )

    def test_nao_duplica_unidade_espelho_que_ja_existe(
        self,
        client_gestor_carga_lot,
        patriarca,
        tipo_organograma,
    ):
        content = conteudo_organograma(
            linhas=[
                (1, "Prefeitura", "PMVV"),
                ("1.1", "Secretaria de Administração", "SEMAD"),
                ("1.1.1", "Secretaria de Administração", "SEMAD"),
            ]
        )

        response = client_gestor_carga_lot.post(
            f"{IMPORTACOES_URL}upload/",
            {
                "idpatriarca_id": patriarca.pk,
                "idtipocarga_id": tipo_organograma.pk,
                "criar_unidades_espelho": True,
                "arquivo": arquivo_planilha(conteudo=content),
            },
            format="multipart",
        )

        assert response.status_code == 201
        assert response.data["total_itens_organograma"] == 3
        assert response.data["flgcriarunidadesespelho"] is True
        assert response.data["flgpossuialteracoes"] is False
        assert response.data["strstatusimportacao"] == (
            StatusImportacaoLocal.INCORPORADA
        )
        assert not ItemOrganograma.objects.filter(
            flggeradoautomaticamente=True
        ).exists()
        summary = response.data["relatorio_validacao"]["jsonresumo"]
        assert summary["total_unidades_espelho_criadas"] == 0

    def test_nao_cria_unidade_espelho_sem_a_opcao(
        self,
        client_gestor_carga_lot,
        patriarca,
        tipo_organograma,
    ):
        content = conteudo_organograma(
            linhas=[
                (1, "Prefeitura", "PMVV"),
                ("1.1", "Secretaria de Administração", "SEMAD"),
            ]
        )

        response = client_gestor_carga_lot.post(
            f"{IMPORTACOES_URL}upload/",
            {
                "idpatriarca_id": patriarca.pk,
                "idtipocarga_id": tipo_organograma.pk,
                "arquivo": arquivo_planilha(conteudo=content),
            },
            format="multipart",
        )

        assert response.status_code == 201
        assert response.data["flgcriarunidadesespelho"] is False
        assert response.data["total_itens_organograma"] == 2
        assert not ItemOrganograma.objects.filter(
            flggeradoautomaticamente=True
        ).exists()

    def test_rejeita_opcao_de_unidade_espelho_em_importacao_de_lotacao(
        self,
        client_gestor_carga_lot,
        patriarca,
        tipo_organograma,
        tipo_lotacao,
    ):
        organogram = client_gestor_carga_lot.post(
            f"{IMPORTACOES_URL}upload/",
            {
                "idpatriarca_id": patriarca.pk,
                "idtipocarga_id": tipo_organograma.pk,
                "arquivo": arquivo_planilha(),
            },
            format="multipart",
        )
        assert organogram.status_code == 201

        response = client_gestor_carga_lot.post(
            f"{IMPORTACOES_URL}upload/",
            {
                "idpatriarca_id": patriarca.pk,
                "idtipocarga_id": tipo_lotacao.pk,
                "idimportacaoorganograma_id": organogram.data["idimportacaoplanilha"],
                "criar_unidades_espelho": True,
                "arquivo": arquivo_planilha(nome="lotacao.xlsx"),
            },
            format="multipart",
        )

        assert response.status_code == 400
        assert "criar_unidades_espelho" in response.data["errors"]

    def test_rejeita_conflito_parcial_ao_criar_unidade_espelho(
        self,
        client_gestor_carga_lot,
        patriarca,
        tipo_organograma,
    ):
        content = conteudo_organograma(
            linhas=[
                (1, "Prefeitura", "PMVV"),
                ("1.1", "Secretaria de Administração", "SEMAD"),
                ("1.1.1", "Secretaria de Administração", "OUTRA"),
            ]
        )

        response = client_gestor_carga_lot.post(
            f"{IMPORTACOES_URL}upload/",
            {
                "idpatriarca_id": patriarca.pk,
                "idtipocarga_id": tipo_organograma.pk,
                "criar_unidades_espelho": True,
                "arquivo": arquivo_planilha(conteudo=content),
            },
            format="multipart",
        )

        assert response.status_code == 400
        report = response.data["errors"]["arquivo"]
        assert "conflito_unidade_espelho" in {
            error["codigo"] for error in report["erros"]
        }
        importacao = ImportacaoPlanilha.objects.get(pk=report["idimportacaoplanilha"])
        assert importacao.flgcriarunidadesespelho is True
        assert not importacao.itens_organograma.exists()
