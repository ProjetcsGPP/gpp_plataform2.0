"""Testes de views: importacao lotacao."""

# O cenário de ciclo de vida completo mantém todas as asserções em um único teste.
# pylint: disable=too-many-locals,too-many-statements

from io import BytesIO

import pytest
from django.core.exceptions import ValidationError as DjangoValidationError
from openpyxl import load_workbook

from apps.carga_org_lot.models import (ConjuntoLotacao,
                                       FaseProcessoCargaInicial,
                                       ImportacaoPlanilha, ItemLotacao,
                                       ItemLotacaoImportado,
                                       ProcessoCargaInicial,
                                       ResolucaoValidacao,
                                       ResultadoRelatorioValidacao,
                                       SituacaoConjuntoLotacao,
                                       StatusImportacaoLocal,
                                       StatusItemLotacaoImportado,
                                       StatusLotacaoProcesso,
                                       StatusOrganogramaProcesso)
from apps.carga_org_lot.tests.view_helpers import (CASOS_VALIDACAO_URL,
                                                   CONJUNTOS_LOTACAO_URL,
                                                   IMPORTACOES_URL,
                                                   ITENS_LOTACAO_URL,
                                                   arquivo_planilha,
                                                   conteudo_lotacao)


@pytest.mark.django_db(transaction=True)
class TestImportacaoLotacao:
    def test_lotacao_exige_organograma_do_mesmo_patriarca(
        self,
        client_gestor_carga_lot,
        patriarca,
        tipo_organograma,
        tipo_lotacao,
    ):
        organograma = client_gestor_carga_lot.post(
            f"{IMPORTACOES_URL}upload/",
            {
                "idpatriarca_id": patriarca.pk,
                "idtipocarga_id": tipo_organograma.pk,
                "arquivo": arquivo_planilha(),
            },
            format="multipart",
        )
        assert organograma.status_code == 201
        sem_referencia = client_gestor_carga_lot.post(
            f"{IMPORTACOES_URL}upload/",
            {
                "idpatriarca_id": patriarca.pk,
                "idtipocarga_id": tipo_lotacao.pk,
                "arquivo": arquivo_planilha(nome="lotacao.xlsx"),
            },
            format="multipart",
        )
        assert sem_referencia.status_code == 400

        response = client_gestor_carga_lot.post(
            f"{IMPORTACOES_URL}upload/",
            {
                "idpatriarca_id": patriarca.pk,
                "idtipocarga_id": tipo_lotacao.pk,
                "idimportacaoorganograma_id": organograma.data["idimportacaoplanilha"],
                "arquivo": arquivo_planilha(
                    nome="lotacao.xlsx",
                    conteudo=conteudo_lotacao(),
                ),
            },
            format="multipart",
        )
        assert response.status_code == 201
        assert (
            response.data["idimportacaoorganograma_id"]
            == organograma.data["idimportacaoplanilha"]
        )
        assert response.data["total_itens_lotacao"] == 1
        process = ProcessoCargaInicial.objects.get(idpatriarca=patriarca)
        assert process.strfaseatual == FaseProcessoCargaInicial.PREPARANDO_LOTACOES
        assert process.strstatusorganograma == StatusOrganogramaProcesso.PRONTO
        assert process.strstatuslotacao == StatusLotacaoProcesso.EM_PREPARACAO

    def test_lotacao_processa_todas_as_abas_e_persiste_itens_normalizados(
        self,
        client_gestor_carga_lot,
        patriarca,
        tipo_organograma,
        tipo_lotacao,
    ):
        organograma = client_gestor_carga_lot.post(
            f"{IMPORTACOES_URL}upload/",
            {
                "idpatriarca_id": patriarca.pk,
                "idtipocarga_id": tipo_organograma.pk,
                "arquivo": arquivo_planilha(),
            },
            format="multipart",
        )
        segunda_linha = (
            "SEMAD",
            "  secretaria de administracao  ",
            "João Souza",
            1234567890,
            "Analista Administrativo",
            "SEMAD",
            "secretaria de administracao",
        )

        response = client_gestor_carga_lot.post(
            f"{IMPORTACOES_URL}upload/",
            {
                "idpatriarca_id": patriarca.pk,
                "idtipocarga_id": tipo_lotacao.pk,
                "idimportacaoorganograma_id": organograma.data["idimportacaoplanilha"],
                "arquivo": arquivo_planilha(
                    nome="lotacao.xlsx",
                    conteudo=conteudo_lotacao(segunda_aba=[segunda_linha]),
                ),
            },
            format="multipart",
        )

        assert response.status_code == 201
        assert response.data["total_itens_lotacao"] == 2
        report = response.data["relatorio_validacao"]
        assert report["strresultado"] == ResultadoRelatorioValidacao.VALIDO
        assert report["strversaoregras"] == "lotacao-v4"
        assert report["inttotalabas"] == 2
        assert report["inttotallinhas"] == 2
        assert report["inttotalvalidas"] == 2
        assert report["inttotalnormalizacoes"] >= 2

        items = ItemLotacao.objects.order_by("strabaorigem", "intlinhaorigem")
        assert list(items.values_list("strcpf", flat=True)) == [
            "90396100457",
            "01234567890",
        ]
        assert items[0].strcargooriginal == "auxiliar administrativo"
        assert items[0].strcargonormalizado == "AUXILIAR ADMINISTRATIVO"
        assert items[0].idorgaoorganograma_id == items[0].idunidadeorganograma_id

        listed = client_gestor_carga_lot.get(
            ITENS_LOTACAO_URL,
            {"importacao": response.data["idimportacaoplanilha"]},
        )
        assert listed.status_code == 200
        assert listed.data["count"] == 2

    def test_conjunto_lotacao_acrescenta_homologa_ativa_e_exporta(
        self,
        client_gestor_carga_lot,
        patriarca,
        tipo_organograma,
        tipo_lotacao,
    ):
        organograma = client_gestor_carga_lot.post(
            f"{IMPORTACOES_URL}upload/",
            {
                "idpatriarca_id": patriarca.pk,
                "idtipocarga_id": tipo_organograma.pk,
                "arquivo": arquivo_planilha(),
            },
            format="multipart",
        )
        first_upload = client_gestor_carga_lot.post(
            f"{IMPORTACOES_URL}upload/",
            {
                "idpatriarca_id": patriarca.pk,
                "idtipocarga_id": tipo_lotacao.pk,
                "idimportacaoorganograma_id": organograma.data["idimportacaoplanilha"],
                "strreferencia": "Versão discutida com o RH",
                "arquivo": arquivo_planilha(
                    nome="lotacao-inicial.xlsx",
                    conteudo=conteudo_lotacao(),
                ),
            },
            format="multipart",
        )
        assert first_upload.status_code == 201
        set_id = first_upload.data["idconjuntolotacao_id"]
        conjunto = ConjuntoLotacao.objects.get(pk=set_id)
        assert conjunto.strsituacao == SituacaoConjuntoLotacao.PRONTO_PARA_HOMOLOGACAO
        assert first_upload.data["strtipooperacaolotacao"] == "INICIAL"

        second_row = (
            "SEMAD",
            "Secretaria de Administração",
            "João Souza",
            "01234567890",
            "Analista Administrativo",
            "SEMAD",
            "Secretaria de Administração",
        )
        addition = client_gestor_carga_lot.post(
            f"{CONJUNTOS_LOTACAO_URL}{set_id}/importacoes/",
            {
                "idtipocarga_id": tipo_lotacao.pk,
                "strreferencia": "Complemento recebido da secretaria",
                "arquivo": arquivo_planilha(
                    nome="lotacao-complementar.xlsx",
                    conteudo=conteudo_lotacao(linhas=[second_row]),
                ),
            },
            format="multipart",
        )
        assert addition.status_code == 201
        assert addition.data["idconjuntolotacao_id"] == set_id
        assert addition.data["strtipooperacaolotacao"] == "ACRESCIMO"
        assert conjunto.importacoes.count() == 2
        assert (
            ItemLotacao.objects.filter(
                idimportacaoplanilha__idconjuntolotacao=conjunto
            ).count()
            == 2
        )

        listed = client_gestor_carga_lot.get(
            CONJUNTOS_LOTACAO_URL,
            {
                "organograma": organograma.data["idimportacaoplanilha"],
            },
        )
        assert listed.status_code == 200
        assert listed.data["count"] == 1
        assert listed.data["results"][0]["total_importacoes"] == 2
        assert listed.data["results"][0]["total_registros"] == 2

        summary = client_gestor_carga_lot.get(
            f"{CONJUNTOS_LOTACAO_URL}resumo-organograma/",
            {"organograma": organograma.data["idimportacaoplanilha"]},
        )
        assert summary.status_code == 200
        assert summary.data["total_conjuntos"] == 1
        assert summary.data["total_importacoes"] == 2
        assert summary.data["conjunto_ativo"] is None

        homologated = client_gestor_carga_lot.post(
            f"{CONJUNTOS_LOTACAO_URL}{set_id}/homologar/",
            {"justificativa": "Conteúdo confirmado com o RH municipal."},
            format="json",
        )
        assert homologated.status_code == 200
        assert homologated.data["strsituacao"] == "HOMOLOGADO"

        activated = client_gestor_carga_lot.post(
            f"{CONJUNTOS_LOTACAO_URL}{set_id}/ativar/",
            {"justificativa": "Primeira versão homologada."},
            format="json",
        )
        assert activated.status_code == 200
        assert activated.data["strsituacao"] == "ATIVO"

        closed_addition = client_gestor_carga_lot.post(
            f"{CONJUNTOS_LOTACAO_URL}{set_id}/importacoes/",
            {
                "idtipocarga_id": tipo_lotacao.pk,
                "arquivo": arquivo_planilha(
                    nome="nao-pode-alterar-ativo.xlsx",
                    conteudo=conteudo_lotacao(linhas=[second_row]),
                ),
            },
            format="multipart",
        )
        assert closed_addition.status_code == 400

        protected_original = client_gestor_carga_lot.delete(
            (f"{IMPORTACOES_URL}" f"{first_upload.data['idimportacaoplanilha']}/")
        )
        assert protected_original.status_code == 400
        assert ImportacaoPlanilha.objects.filter(
            pk=first_upload.data["idimportacaoplanilha"]
        ).exists()

        current_items = client_gestor_carga_lot.get(
            ITENS_LOTACAO_URL,
            {
                "conjunto_lotacao": set_id,
                "vigente": "true",
            },
        )
        assert current_items.status_code == 200
        assert current_items.data["count"] == 2

        history = client_gestor_carga_lot.get(
            f"{CONJUNTOS_LOTACAO_URL}{set_id}/historico/"
        )
        assert history.status_code == 200
        assert history.data["count"] >= 4
        assert {event["stracao"] for event in history.data["results"]} >= {
            "CONJUNTO_CRIADO",
            "PLANILHA_IMPORTADA",
            "CONJUNTO_HOMOLOGADO",
            "CONJUNTO_ATIVADO",
        }

        exported = client_gestor_carga_lot.get(
            f"{CONJUNTOS_LOTACAO_URL}{set_id}/download-consolidado/"
        )
        assert exported.status_code == 200
        workbook = load_workbook(BytesIO(exported.content))
        assert workbook.sheetnames == [
            "LOTACOES_CONSOLIDADAS",
            "REGISTROS_DESCARTADOS",
            "RESUMO",
            "HISTORICO",
        ]
        assert workbook["LOTACOES_CONSOLIDADAS"].max_row == 3

        original = client_gestor_carga_lot.get(first_upload.data["download_url"])
        assert original.status_code == 200
        assert (
            original.content
            == ImportacaoPlanilha.objects.get(
                pk=first_upload.data["idimportacaoplanilha"]
            ).binconteudo
        )

    def test_ativar_novo_conjunto_substitui_anterior_e_preserva_historico(
        self,
        client_gestor_carga_lot,
        patriarca,
        tipo_organograma,
        tipo_lotacao,
    ):
        organograma = client_gestor_carga_lot.post(
            f"{IMPORTACOES_URL}upload/",
            {
                "idpatriarca_id": patriarca.pk,
                "idtipocarga_id": tipo_organograma.pk,
                "arquivo": arquivo_planilha(),
            },
            format="multipart",
        )

        def create_ready_set(reference, filename, person_name):
            created = client_gestor_carga_lot.post(
                CONJUNTOS_LOTACAO_URL,
                {
                    "idpatriarca_id": patriarca.pk,
                    "idimportacaoorganograma_id": organograma.data[
                        "idimportacaoplanilha"
                    ],
                    "strreferencia": reference,
                },
                format="json",
            )
            assert created.status_code == 201
            row = (
                "SEMAD",
                "Secretaria de Administração",
                person_name,
                "90396100457",
                "Auxiliar Administrativo",
                "SEMAD",
                "Secretaria de Administração",
            )
            uploaded = client_gestor_carga_lot.post(
                (
                    f"{CONJUNTOS_LOTACAO_URL}"
                    f"{created.data['idconjuntolotacao']}/importacoes/"
                ),
                {
                    "idtipocarga_id": tipo_lotacao.pk,
                    "arquivo": arquivo_planilha(
                        nome=filename,
                        conteudo=conteudo_lotacao(linhas=[row]),
                    ),
                },
                format="multipart",
            )
            assert uploaded.status_code == 201
            expected_operation = (
                "SUBSTITUICAO"
                if ConjuntoLotacao.objects.filter(
                    idimportacaoorganograma_id=organograma.data["idimportacaoplanilha"],
                    strsituacao=SituacaoConjuntoLotacao.ATIVO,
                )
                .exclude(pk=created.data["idconjuntolotacao"])
                .exists()
                else "INICIAL"
            )
            assert uploaded.data["strtipooperacaolotacao"] == expected_operation
            homologated = client_gestor_carga_lot.post(
                (
                    f"{CONJUNTOS_LOTACAO_URL}"
                    f"{created.data['idconjuntolotacao']}/homologar/"
                ),
                {"justificativa": "Conferido com a prefeitura."},
                format="json",
            )
            assert homologated.status_code == 200
            return created.data["idconjuntolotacao"]

        first_id = create_ready_set("Alternativa inicial", "primeira.xlsx", "Maria")
        first_active = client_gestor_carga_lot.post(
            f"{CONJUNTOS_LOTACAO_URL}{first_id}/ativar/",
            {"justificativa": "Primeira versão vigente."},
            format="json",
        )
        assert first_active.status_code == 200

        replacement_id = create_ready_set(
            "Nova consolidação",
            "substituta.xlsx",
            "Joana",
        )
        replacement = client_gestor_carga_lot.post(
            f"{CONJUNTOS_LOTACAO_URL}{replacement_id}/ativar/",
            {"justificativa": "Nova consolidação aprovada pela prefeitura."},
            format="json",
        )
        assert replacement.status_code == 200
        assert replacement.data["strsituacao"] == "ATIVO"
        assert replacement.data["idconjuntoanterior_id"] == first_id

        first = ConjuntoLotacao.objects.get(pk=first_id)
        assert first.strsituacao == SituacaoConjuntoLotacao.SUBSTITUIDO
        assert first.importacoes.count() == 1
        assert first.historico.filter(stracao="CONJUNTO_SUBSTITUIDO").exists()
        assert (
            ConjuntoLotacao.objects.filter(
                idimportacaoorganograma_id=organograma.data["idimportacaoplanilha"],
                strsituacao=SituacaoConjuntoLotacao.ATIVO,
            ).count()
            == 1
        )

        first_detail = client_gestor_carga_lot.get(
            f"{CONJUNTOS_LOTACAO_URL}{first_id}/"
        )
        assert first_detail.status_code == 200
        assert first_detail.data["idconjuntosubstituto"] == replacement_id

        old_items = client_gestor_carga_lot.get(
            ITENS_LOTACAO_URL,
            {"conjunto_lotacao": first_id, "vigente": "false"},
        )
        assert old_items.status_code == 200
        assert old_items.data["count"] == 1

        cancelled = client_gestor_carga_lot.post(
            f"{CONJUNTOS_LOTACAO_URL}{replacement_id}/cancelar/",
            {"justificativa": "Prefeitura solicitou nova revisão."},
            format="json",
        )
        assert cancelled.status_code == 200
        assert cancelled.data["strsituacao"] == "CANCELADO"
        assert ConjuntoLotacao.objects.get(pk=replacement_id).importacoes.count() == 1

    def test_lotacao_rejeita_cpf_com_digitos_verificadores_invalidos(
        self,
        client_gestor_carga_lot,
        patriarca,
        tipo_organograma,
        tipo_lotacao,
    ):
        organograma = client_gestor_carga_lot.post(
            f"{IMPORTACOES_URL}upload/",
            {
                "idpatriarca_id": patriarca.pk,
                "idtipocarga_id": tipo_organograma.pk,
                "arquivo": arquivo_planilha(),
            },
            format="multipart",
        )
        linha = (
            "SEMAD",
            "Secretaria de Administração",
            "Maria da Silva",
            "123.456.789-01",
            "Auxiliar Administrativo",
            "SEMAD",
            "Secretaria de Administração",
        )

        response = client_gestor_carga_lot.post(
            f"{IMPORTACOES_URL}upload/",
            {
                "idpatriarca_id": patriarca.pk,
                "idtipocarga_id": tipo_lotacao.pk,
                "idimportacaoorganograma_id": organograma.data["idimportacaoplanilha"],
                "arquivo": arquivo_planilha(
                    nome="lotacao-cpf-invalido.xlsx",
                    conteudo=conteudo_lotacao(linhas=[linha]),
                ),
            },
            format="multipart",
        )

        assert response.status_code == 201
        assert response.data["strstatusimportacao"] == "AGUARDANDO_REVISAO"
        assert {
            error["codigo"]
            for error in response.data["relatorio_validacao"]["jsonresumo"]["erros"]
        } == {"cpf_digitos_verificadores_invalidos"}
        staged = ItemLotacaoImportado.objects.get(
            idimportacaoplanilha_id=response.data["idimportacaoplanilha"]
        )
        assert staged.strsituacao == StatusItemLotacaoImportado.COM_ERROS
        assert not ItemLotacao.objects.filter(
            idimportacaoplanilha_id=response.data["idimportacaoplanilha"]
        ).exists()

    def test_lotacao_avisa_cpf_repetido_inclusive_entre_abas(
        self,
        client_gestor_carga_lot,
        patriarca,
        tipo_organograma,
        tipo_lotacao,
    ):
        organograma = client_gestor_carga_lot.post(
            f"{IMPORTACOES_URL}upload/",
            {
                "idpatriarca_id": patriarca.pk,
                "idtipocarga_id": tipo_organograma.pk,
                "arquivo": arquivo_planilha(),
            },
            format="multipart",
        )
        repetida = (
            "SEMAD",
            "Secretaria de Administração",
            "Maria da Silva",
            "90396100457",
            "Outro cargo",
            "SEMAD",
            "Secretaria de Administração",
        )

        response = client_gestor_carga_lot.post(
            f"{IMPORTACOES_URL}upload/",
            {
                "idpatriarca_id": patriarca.pk,
                "idtipocarga_id": tipo_lotacao.pk,
                "idimportacaoorganograma_id": organograma.data["idimportacaoplanilha"],
                "arquivo": arquivo_planilha(
                    nome="lotacao-cpf-repetido.xlsx",
                    conteudo=conteudo_lotacao(segunda_aba=[repetida]),
                ),
            },
            format="multipart",
        )

        assert response.status_code == 201
        report = response.data["relatorio_validacao"]
        assert report["inttotalavisos"] == 2
        assert {
            warning["strcodigo"]
            for warning in report["ocorrencias"]
            if warning["strseveridade"] == "AVISO"
        } == {"cpf_duplicado_planilha"}
        assert response.data["strstatusimportacao"] == "AGUARDANDO_REVISAO"
        assert response.data["total_itens_lotacao"] == 0
        assert response.data["total_itens_lotacao_importados"] == 2
        conjunto = ConjuntoLotacao.objects.get(pk=response.data["idconjuntolotacao_id"])
        assert conjunto.strsituacao == SituacaoConjuntoLotacao.COM_PENDENCIAS
        cases = report["casos"]
        assert len(cases) == 1
        assert len(cases[0]["item_ids"]) == 2

        resolution = client_gestor_carga_lot.post(
            f"{CASOS_VALIDACAO_URL}{cases[0]['idcasovalidacao']}/resolver/",
            {
                "acao": "ACEITAR_COMO_INFORMADO",
                "justificativa": "O servidor possui duas lotações válidas.",
            },
            format="json",
        )
        assert resolution.status_code == 200
        assert resolution.data["strstatus"] == "ACEITO_COMO_INFORMADO"
        importacao = ImportacaoPlanilha.objects.get(
            pk=response.data["idimportacaoplanilha"]
        )
        assert (
            importacao.strstatusimportacao
            == StatusImportacaoLocal.PRONTA_PARA_INCORPORAR
        )
        conjunto.refresh_from_db()
        assert conjunto.strsituacao == SituacaoConjuntoLotacao.EM_VALIDACAO

        inherited_item = ItemLotacaoImportado.objects.filter(
            idimportacaoplanilha=importacao
        ).first()
        inherited_item.strcpf = "903.961.004-57"
        inherited_item.save(update_fields=["strcpf", "updated_at"])

        incorporated = client_gestor_carga_lot.post(
            f"{IMPORTACOES_URL}{importacao.pk}/incorporar-lotacao/",
            {},
            format="json",
        )
        assert incorporated.status_code == 200
        assert incorporated.data["strstatusimportacao"] == "INCORPORADA"
        assert incorporated.data["total_itens_lotacao"] == 2
        inherited_item.refresh_from_db()
        assert inherited_item.strcpf == "90396100457"
        conjunto.refresh_from_db()
        assert conjunto.strsituacao == SituacaoConjuntoLotacao.PRONTO_PARA_HOMOLOGACAO
        assert set(
            ItemLotacaoImportado.objects.filter(
                idimportacaoplanilha=importacao
            ).values_list("strsituacao", flat=True)
        ) == {StatusItemLotacaoImportado.INCORPORADO}

    def test_correcao_manual_normaliza_cpf_formatado_antes_de_incorporar(
        self,
        client_gestor_carga_lot,
        patriarca,
        tipo_organograma,
        tipo_lotacao,
    ):
        organograma = client_gestor_carga_lot.post(
            f"{IMPORTACOES_URL}upload/",
            {
                "idpatriarca_id": patriarca.pk,
                "idtipocarga_id": tipo_organograma.pk,
                "arquivo": arquivo_planilha(),
            },
            format="multipart",
        )
        repetida = (
            "SEMAD",
            "Secretaria de Administração",
            "Maria da Silva",
            "90396100457",
            "Outro cargo",
            "SEMAD",
            "Secretaria de Administração",
        )
        uploaded = client_gestor_carga_lot.post(
            f"{IMPORTACOES_URL}upload/",
            {
                "idpatriarca_id": patriarca.pk,
                "idtipocarga_id": tipo_lotacao.pk,
                "idimportacaoorganograma_id": organograma.data["idimportacaoplanilha"],
                "arquivo": arquivo_planilha(
                    nome="lotacao-correcao-cpf.xlsx",
                    conteudo=conteudo_lotacao(segunda_aba=[repetida]),
                ),
            },
            format="multipart",
        )
        case = uploaded.data["relatorio_validacao"]["casos"][0]
        item_ids = case["item_ids"]

        corrected = client_gestor_carga_lot.post(
            f"{CASOS_VALIDACAO_URL}{case['idcasovalidacao']}/resolver/",
            {
                "acao": "CORRIGIR_MANUALMENTE",
                "alteracoes": [
                    {
                        "iditemlotacaoimportado": item_ids[0],
                        "campos": {"strcpf": "980.545.250-60"},
                    },
                    {
                        "iditemlotacaoimportado": item_ids[1],
                        "campos": {"strcpf": "359.644.720-87"},
                    },
                ],
            },
            format="json",
        )

        assert corrected.status_code == 200
        importacao_id = uploaded.data["idimportacaoplanilha"]
        assert set(
            ItemLotacaoImportado.objects.filter(
                idimportacaoplanilha_id=importacao_id
            ).values_list("strcpf", flat=True)
        ) == {"98054525060", "35964472087"}
        resolution = ResolucaoValidacao.objects.get(
            idcasovalidacao_id=case["idcasovalidacao"]
        )
        assert {item["strcpf"] for item in resolution.jsondepois} == {
            "98054525060",
            "35964472087",
        }

        incorporated = client_gestor_carga_lot.post(
            f"{IMPORTACOES_URL}{importacao_id}/incorporar-lotacao/",
            {},
            format="json",
        )

        assert incorporated.status_code == 200
        assert set(
            ItemLotacao.objects.filter(
                idimportacaoplanilha_id=importacao_id
            ).values_list("strcpf", flat=True)
        ) == {"98054525060", "35964472087"}

    def test_incorporacao_converte_erro_de_modelo_em_resposta_400(
        self,
        client_gestor_carga_lot,
        patriarca,
        tipo_organograma,
        tipo_lotacao,
        monkeypatch,
    ):
        organograma = client_gestor_carga_lot.post(
            f"{IMPORTACOES_URL}upload/",
            {
                "idpatriarca_id": patriarca.pk,
                "idtipocarga_id": tipo_organograma.pk,
                "arquivo": arquivo_planilha(),
            },
            format="multipart",
        )
        repetida = (
            "SEMAD",
            "Secretaria de Administração",
            "Maria da Silva",
            "90396100457",
            "Outro cargo",
            "SEMAD",
            "Secretaria de Administração",
        )
        uploaded = client_gestor_carga_lot.post(
            f"{IMPORTACOES_URL}upload/",
            {
                "idpatriarca_id": patriarca.pk,
                "idtipocarga_id": tipo_lotacao.pk,
                "idimportacaoorganograma_id": organograma.data["idimportacaoplanilha"],
                "arquivo": arquivo_planilha(
                    nome="lotacao-erro-modelo.xlsx",
                    conteudo=conteudo_lotacao(segunda_aba=[repetida]),
                ),
            },
            format="multipart",
        )
        case = uploaded.data["relatorio_validacao"]["casos"][0]
        accepted = client_gestor_carga_lot.post(
            f"{CASOS_VALIDACAO_URL}{case['idcasovalidacao']}/resolver/",
            {
                "acao": "ACEITAR_COMO_INFORMADO",
                "justificativa": "As duas lotações são intencionais.",
            },
            format="json",
        )
        assert accepted.status_code == 200
        importacao_id = uploaded.data["idimportacaoplanilha"]

        def invalid_model(_instance):
            raise DjangoValidationError(
                {"strnomeservidor": ["Valor rejeitado pelo modelo."]}
            )

        monkeypatch.setattr(ItemLotacao, "full_clean", invalid_model)

        response = client_gestor_carga_lot.post(
            f"{IMPORTACOES_URL}{importacao_id}/incorporar-lotacao/",
            {},
            format="json",
        )

        assert response.status_code == 400
        assert not ItemLotacao.objects.filter(
            idimportacaoplanilha_id=importacao_id
        ).exists()
        importacao = ImportacaoPlanilha.objects.get(pk=importacao_id)
        assert (
            importacao.strstatusimportacao
            == StatusImportacaoLocal.PRONTA_PARA_INCORPORAR
        )
