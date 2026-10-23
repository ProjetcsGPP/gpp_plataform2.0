"""Testes de views: consultas lotacao."""

from io import BytesIO

import pytest
from django.utils import timezone
from openpyxl import load_workbook

from apps.carga_org_lot.tests.view_helpers import (
    CASOS_VALIDACAO_URL, CONJUNTOS_LOTACAO_URL, IMPORTACOES_URL,
    ITENS_LOTACAO_IMPORTADOS_URL, ITENS_LOTACAO_URL, arquivo_planilha,
    conteudo_lotacao)


@pytest.mark.django_db(transaction=True)
class TestConsultasOperacionaisLotacao:
    @staticmethod
    def _importar_organograma(client, patriarca, tipo_organograma):
        response = client.post(
            f"{IMPORTACOES_URL}upload/",
            {
                "idpatriarca_id": patriarca.pk,
                "idtipocarga_id": tipo_organograma.pk,
                "strreferencia": "Estrutura vigente",
                "arquivo": arquivo_planilha(),
            },
            format="multipart",
        )
        assert response.status_code == 201
        return response

    def test_conjunto_expoe_contexto_acoes_filtros_e_resumo(
        self,
        client_gestor_carga_lot,
        patriarca,
        tipo_organograma,
        tipo_lotacao,
    ):
        organogram = self._importar_organograma(
            client_gestor_carga_lot,
            patriarca,
            tipo_organograma,
        )
        imported = client_gestor_carga_lot.post(
            f"{IMPORTACOES_URL}upload/",
            {
                "idpatriarca_id": patriarca.pk,
                "idtipocarga_id": tipo_lotacao.pk,
                "idimportacaoorganograma_id": organogram.data["idimportacaoplanilha"],
                "strreferencia": "Lotacao Outubro",
                "arquivo": arquivo_planilha(
                    nome="lotacao-outubro.xlsx",
                    conteudo=conteudo_lotacao(),
                ),
            },
            format="multipart",
        )
        assert imported.status_code == 201
        set_id = imported.data["idconjuntolotacao_id"]

        detail = client_gestor_carga_lot.get(f"{CONJUNTOS_LOTACAO_URL}{set_id}/")
        assert detail.status_code == 200
        assert detail.data["patriarca_resumo"] == {
            "idpatriarca": patriarca.pk,
            "sigla": patriarca.strsiglapatriarca,
            "nome": patriarca.strnome,
        }
        assert detail.data["organograma_resumo"]["idimportacaoplanilha"] == (
            organogram.data["idimportacaoplanilha"]
        )
        actions = detail.data["acoes_disponiveis"]
        assert actions["adicionar_planilha"]["permitida"] is True
        assert actions["homologar"]["permitida"] is True
        assert actions["ativar"]["permitida"] is False

        listed = client_gestor_carga_lot.get(
            CONJUNTOS_LOTACAO_URL,
            {
                "patriarca": patriarca.pk,
                "referencia": "outubro",
                "ordering": "strreferencia",
                "data_inicio": timezone.localdate().isoformat(),
            },
        )
        assert listed.status_code == 200
        assert listed.data["count"] == 1

        summary = client_gestor_carga_lot.get(
            f"{CONJUNTOS_LOTACAO_URL}resumo/",
            {"patriarca": patriarca.pk},
        )
        assert summary.status_code == 200
        assert summary.data["total_conjuntos"] == 1
        ready = next(
            entry
            for entry in summary.data["por_situacao"]
            if entry["situacao"] == "PRONTO_PARA_HOMOLOGACAO"
        )
        assert ready["total"] == 1

    def test_central_pendencias_expoe_contexto_filtros_e_resumo(
        self,
        client_gestor_carga_lot,
        patriarca,
        tipo_organograma,
        tipo_lotacao,
    ):
        organogram = self._importar_organograma(
            client_gestor_carga_lot,
            patriarca,
            tipo_organograma,
        )
        repeated = (
            "SEMAD",
            "Secretaria de Administração",
            "Maria da Silva",
            "90396100457",
            "Auxiliar Administrativo",
            "SEMAD",
            "Secretaria de Administração",
        )
        imported = client_gestor_carga_lot.post(
            f"{IMPORTACOES_URL}upload/",
            {
                "idpatriarca_id": patriarca.pk,
                "idtipocarga_id": tipo_lotacao.pk,
                "idimportacaoorganograma_id": organogram.data["idimportacaoplanilha"],
                "strreferencia": "Conferencia RH",
                "arquivo": arquivo_planilha(
                    nome="lotacao-pendente.xlsx",
                    conteudo=conteudo_lotacao(segunda_aba=[repeated]),
                ),
            },
            format="multipart",
        )
        assert imported.status_code == 201
        case = imported.data["relatorio_validacao"]["casos"][0]

        response = client_gestor_carga_lot.get(
            CASOS_VALIDACAO_URL,
            {
                "patriarca": patriarca.pk,
                "organograma": organogram.data["idimportacaoplanilha"],
                "conjunto": imported.data["idconjuntolotacao_id"],
                "codigo": "cpf_duplicado_planilha",
                "assunto": "CPF",
                "status": "PENDENTE",
                "q": "pendente",
            },
        )
        assert response.status_code == 200
        assert response.data["count"] == 1
        result = response.data["results"][0]
        assert result["idcasovalidacao"] == case["idcasovalidacao"]
        assert result["assunto"] == "CPF"
        assert result["assunto_nome"] == "CPF"
        assert result["contexto"]["patriarca"]["id"] == patriarca.pk
        assert result["contexto"]["importacao"]["nome_arquivo"] == (
            "lotacao-pendente.xlsx"
        )

        summary = client_gestor_carga_lot.get(
            f"{CASOS_VALIDACAO_URL}resumo/",
            {"patriarca": patriarca.pk, "status": "PENDENTE"},
        )
        assert summary.status_code == 200
        assert summary.data["total"] == 1
        cpf = next(
            entry for entry in summary.data["por_assunto"] if entry["assunto"] == "CPF"
        )
        assert cpf["total"] == 1
        assert summary.data["por_severidade"] == [{"severidade": "AVISO", "total": 1}]
        assert summary.data["por_status"] == [{"status": "PENDENTE", "total": 1}]
        assert summary.data["por_patriarca"] == [
            {
                "idpatriarca": patriarca.pk,
                "nome": patriarca.strnome,
                "total": 1,
            }
        ]
        assert summary.data["por_organograma"] == [
            {
                "idimportacaoorganograma": organogram.data["idimportacaoplanilha"],
                "referencia": "Estrutura vigente",
                "total": 1,
            }
        ]
        assert summary.data["por_conjunto"] == [
            {
                "idconjuntolotacao": imported.data["idconjuntolotacao_id"],
                "referencia": "Conferencia RH",
                "total": 1,
            }
        ]
        assert (
            client_gestor_carga_lot.get(
                CASOS_VALIDACAO_URL,
                {"assunto": "INVALIDO"},
            ).status_code
            == 400
        )

    def test_lotacoes_vigentes_suportam_contexto_e_filtros(
        self,
        client_gestor_carga_lot,
        patriarca,
        tipo_organograma,
        tipo_lotacao,
    ):
        organogram = self._importar_organograma(
            client_gestor_carga_lot,
            patriarca,
            tipo_organograma,
        )
        imported = client_gestor_carga_lot.post(
            f"{IMPORTACOES_URL}upload/",
            {
                "idpatriarca_id": patriarca.pk,
                "idtipocarga_id": tipo_lotacao.pk,
                "idimportacaoorganograma_id": organogram.data["idimportacaoplanilha"],
                "strreferencia": "Versao vigente",
                "arquivo": arquivo_planilha(
                    nome="lotacao-vigente.xlsx",
                    conteudo=conteudo_lotacao(),
                ),
            },
            format="multipart",
        )
        set_id = imported.data["idconjuntolotacao_id"]
        assert (
            client_gestor_carga_lot.post(
                f"{CONJUNTOS_LOTACAO_URL}{set_id}/homologar/",
                {"justificativa": "Dados conferidos."},
                format="json",
            ).status_code
            == 200
        )
        assert (
            client_gestor_carga_lot.post(
                f"{CONJUNTOS_LOTACAO_URL}{set_id}/ativar/",
                {"justificativa": "Versão vigente."},
                format="json",
            ).status_code
            == 200
        )

        response = client_gestor_carga_lot.get(
            ITENS_LOTACAO_URL,
            {
                "vigente": "true",
                "patriarca": patriarca.pk,
                "organograma": organogram.data["idimportacaoplanilha"],
                "conjunto_lotacao": set_id,
                "nome": "maria",
                "cargo": "auxiliar",
                "q": "903961",
            },
        )
        assert response.status_code == 200
        assert response.data["count"] == 1
        item = response.data["results"][0]
        assert item["contexto"]["conjunto"]["id"] == set_id
        assert (
            item["contexto"]["organograma"]["id"]
            == organogram.data["idimportacaoplanilha"]
        )
        assert item["contexto"]["patriarca"]["nome"] == patriarca.strnome

    def test_descartados_expoem_motivo_autoria_contexto_e_filtros(
        self,
        client_gestor_carga_lot,
        gestor_carga_lot,
        patriarca,
        tipo_organograma,
        tipo_lotacao,
    ):
        organogram = self._importar_organograma(
            client_gestor_carga_lot,
            patriarca,
            tipo_organograma,
        )
        repeated = (
            "SEMAD",
            "Secretaria de Administração",
            "Maria da Silva",
            "90396100457",
            "Auxiliar Administrativo",
            "SEMAD",
            "Secretaria de Administração",
        )
        imported = client_gestor_carga_lot.post(
            f"{IMPORTACOES_URL}upload/",
            {
                "idpatriarca_id": patriarca.pk,
                "idtipocarga_id": tipo_lotacao.pk,
                "idimportacaoorganograma_id": organogram.data["idimportacaoplanilha"],
                "arquivo": arquivo_planilha(
                    nome="lotacao-descartes.xlsx",
                    conteudo=conteudo_lotacao(segunda_aba=[repeated]),
                ),
            },
            format="multipart",
        )
        case = imported.data["relatorio_validacao"]["casos"][0]
        discarded = client_gestor_carga_lot.post(
            f"{CASOS_VALIDACAO_URL}{case['idcasovalidacao']}/resolver/",
            {
                "acao": "DESCARTAR_LINHAS",
                "justificativa": "Registros removidos após conferência.",
            },
            format="json",
        )
        assert discarded.status_code == 200

        response = client_gestor_carga_lot.get(
            f"{ITENS_LOTACAO_IMPORTADOS_URL}descartados/",
            {
                "patriarca": patriarca.pk,
                "organograma": organogram.data["idimportacaoplanilha"],
                "conjunto": imported.data["idconjuntolotacao_id"],
                "codigo": "cpf_duplicado_planilha",
                "descartado_por": gestor_carga_lot.pk,
                "data_inicio": timezone.localdate().isoformat(),
            },
        )
        assert response.status_code == 200
        assert response.data["count"] == 2
        item = response.data["results"][0]
        assert item["contexto"]["patriarca"]["id"] == patriarca.pk
        assert item["descartes"][0]["motivo"] == (
            "Registros removidos após conferência."
        )
        assert item["descartes"][0]["descartado_por_id"] == gestor_carga_lot.pk

    def test_download_pendencias_gera_relatorio_objetivo_da_prefeitura(
        self,
        client_gestor_carga_lot,
        patriarca,
        tipo_organograma,
        tipo_lotacao,
    ):
        organogram = self._importar_organograma(
            client_gestor_carga_lot,
            patriarca,
            tipo_organograma,
        )
        repeated = (
            "SEMAD",
            "Secretaria de Administração",
            "Maria da Silva",
            "90396100457",
            "Auxiliar Administrativo",
            "SEMAD",
            "Secretaria de Administração",
        )
        imported = client_gestor_carga_lot.post(
            f"{IMPORTACOES_URL}upload/",
            {
                "idpatriarca_id": patriarca.pk,
                "idtipocarga_id": tipo_lotacao.pk,
                "idimportacaoorganograma_id": organogram.data["idimportacaoplanilha"],
                "strreferencia": "Planilha devolvida para correção",
                "arquivo": arquivo_planilha(
                    nome="lotacao-com-pendencias.xlsx",
                    conteudo=conteudo_lotacao(segunda_aba=[repeated]),
                ),
            },
            format="multipart",
        )
        assert imported.status_code == 201

        response = client_gestor_carga_lot.get(
            f"{IMPORTACOES_URL}{imported.data['idimportacaoplanilha']}/"
            "pendencias/download/"
        )

        assert response.status_code == 200
        assert response["Content-Type"] == (
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        workbook = load_workbook(BytesIO(response.content))
        assert workbook.sheetnames == ["RESUMO", "PENDENCIAS"]
        summary = {row[0].value: row[1].value for row in workbook["RESUMO"].iter_rows()}
        assert summary["Prefeitura/autarquia"] == patriarca.strnome
        assert summary["Total de pendências"] == 2
        rows = list(workbook["PENDENCIAS"].iter_rows(values_only=True))
        assert rows[0] == (
            "Severidade",
            "Aba",
            "Linha",
            "Campo",
            "Valor informado",
            "Problema",
            "Ação solicitada",
        )
        assert len(rows) == 3
        assert {row[5] for row in rows[1:]} == {"CPF repetido no mesmo órgão e unidade"}
