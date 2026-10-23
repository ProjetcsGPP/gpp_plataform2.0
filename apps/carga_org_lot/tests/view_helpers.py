"""Constantes e fábricas de arquivos compartilhadas pelos testes de views."""

from io import BytesIO

from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl import Workbook

BASE_URL = "/api/carga-org-lot"

PATRIARCAS_URL = f"{BASE_URL}/patriarcas/"

IMPORTACOES_URL = f"{BASE_URL}/importacoes/"

PROCESSOS_URL = f"{BASE_URL}/processos-carga-inicial/"

ITENS_ORGANOGRAMA_URL = f"{BASE_URL}/itens-organograma/"

ITENS_LOTACAO_URL = f"{BASE_URL}/itens-lotacao/"

ITENS_LOTACAO_IMPORTADOS_URL = f"{BASE_URL}/itens-lotacao-importados/"

CASOS_VALIDACAO_URL = f"{BASE_URL}/casos-validacao-importacao/"

CONJUNTOS_LOTACAO_URL = f"{BASE_URL}/conjuntos-lotacao/"

CARGAS_URL = f"{BASE_URL}/cargas/"

HISTORICO_URL = f"{BASE_URL}/historico-cargas/"

ENVIOS_ORGANOGRAMA_URL = f"{BASE_URL}/envios-organograma-prodest/"

EVENTOS_INTEGRACAO_URL = f"{BASE_URL}/eventos-integracao-prodest/"


def conteudo_organograma(
    linhas=None,
    cabecalho=("Numero", "Nome", "Sigla"),
    segunda_aba=None,
):
    linhas = linhas or [
        (1, "Prefeitura Municipal de Vila Velha", "PMVV"),
        ("1.1", "Secretaria de Administração", "SEMAD"),
    ]
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Organograma"
    worksheet.append(cabecalho)
    for linha in linhas:
        worksheet.append(linha)
    if segunda_aba is not None:
        worksheet = workbook.create_sheet("Segunda aba")
        worksheet.append(segunda_aba[0])
        for linha in segunda_aba[1]:
            worksheet.append(linha)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def arquivo_planilha(nome="organograma.xlsx", conteudo=None):
    conteudo = conteudo if conteudo is not None else conteudo_organograma()
    return SimpleUploadedFile(
        nome,
        conteudo,
        content_type=(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        ),
    )


def conteudo_lotacao(linhas=None, segunda_aba=None):
    linhas = linhas or [
        (
            "SEMAD",
            "Secretaria de Administração",
            "Maria da Silva",
            "903.961.004-57",
            "  auxiliar administrativo  ",
            "SEMAD",
            "Secretaria de Administração",
        )
    ]
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "ADMINISTRACAO"
    header = (
        "Sigla",
        "Secretaria/Órgão",
        "Nome do Servidor",
        "CPF",
        "Ocupação/Cargo",
        "Sigla Setor",
        "Setor de Lotação",
    )
    worksheet.append(("", "Patriarca:", "[Município Lotação]"))
    worksheet.append(header)
    for linha in linhas:
        worksheet.append(linha)
    if segunda_aba is not None:
        worksheet = workbook.create_sheet("SEGUNDA SECRETARIA")
        worksheet.append(header)
        for linha in segunda_aba:
            worksheet.append(linha)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
