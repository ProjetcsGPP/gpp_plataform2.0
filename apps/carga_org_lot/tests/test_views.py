"""Testes das APIs de Carga de Organograma e Lotação."""

import hashlib
from io import BytesIO
from uuid import uuid4

import pytest
from django.conf import settings
from django.core.exceptions import ValidationError as DjangoValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.utils import timezone
from openpyxl import Workbook, load_workbook

from apps.carga_org_lot.models import (
    CasoValidacaoImportacao,
    ConjuntoLotacao,
    DetalheStatusCarga,
    EnvioOrganogramaProdest,
    EventoIntegracaoProdest,
    FaseProcessoCargaInicial,
    ImportacaoPlanilha,
    ItemLotacao,
    ItemLotacaoImportado,
    ItemOrganograma,
    OcorrenciaValidacaoImportacao,
    Patriarca,
    ProcessoCargaInicial,
    RelatorioValidacaoImportacao,
    ResolucaoValidacao,
    ResultadoRelatorioValidacao,
    StatusCarga,
    StatusEnvioOrganograma,
    StatusImportacaoLocal,
    StatusItemLotacaoImportado,
    StatusLotacaoProcesso,
    SituacaoConjuntoLotacao,
    StatusOrganogramaProcesso,
    StatusProgresso,
    TipoCarga,
    TokenEnvioCarga,
)
from apps.carga_org_lot.views import (
    CargaOrgLotViewSet,
    ConjuntoLotacaoViewSet,
    EnvioOrganogramaProdestViewSet,
    EventoIntegracaoProdestViewSet,
    HistoricoCargaViewSet,
    ImportacaoPlanilhaViewSet,
    ItemLotacaoViewSet,
    ItemOrganogramaViewSet,
    PatriarcaViewSet,
    ProcessoCargaInicialViewSet,
    StatusCargaViewSet,
    StatusProgressoViewSet,
    TipoCargaViewSet,
)
from apps.carga_org_lot.services.cargarh import CargaRHRequestError

BASE_URL = "/api/carga-org-lot"
PATRIARCAS_URL = f"{BASE_URL}/patriarcas/"
IMPORTACOES_URL = f"{BASE_URL}/importacoes/"
PROCESSOS_URL = f"{BASE_URL}/processos-carga-inicial/"
ITENS_ORGANOGRAMA_URL = f"{BASE_URL}/itens-organograma/"
ITENS_LOTACAO_URL = f"{BASE_URL}/itens-lotacao/"
ITENS_LOTACAO_IMPORTADOS_URL = f"{BASE_URL}/itens-lotacao-importados/"
CASOS_VALIDACAO_URL = f"{BASE_URL}/casos-validacao-importacao/"
CONJUNTOS_LOTACAO_URL = f"{BASE_URL}/conjuntos-lotacao/"
CARGAS_URL = f"{BASE_URL}/cargas/"
HISTORICO_URL = f"{BASE_URL}/historico-cargas/"
ENVIOS_ORGANOGRAMA_URL = f"{BASE_URL}/envios-organograma-prodest/"
EVENTOS_INTEGRACAO_URL = f"{BASE_URL}/eventos-integracao-prodest/"


@pytest.fixture
def status_novo(db):
    return StatusProgresso.objects.create(
        idstatusprogresso=1,
        strdescricao="Novo",
    )


@pytest.fixture
def status_processando(db):
    return StatusProgresso.objects.create(
        idstatusprogresso=2,
        strdescricao="Em progresso",
    )


@pytest.fixture
def tipo_organograma(db):
    return TipoCarga.objects.create(
        idtipocarga=3,
        strdescricao="Organograma",
    )


@pytest.fixture
def tipo_lotacao(db):
    return TipoCarga.objects.create(
        idtipocarga=4,
        strdescricao="Lotações",
    )


@pytest.fixture
def status_carga_recebida(db):
    return StatusCarga.objects.create(
        idstatuscarga=1,
        strdescricao="Recebida",
        flgsucesso=0,
    )


@pytest.fixture
def patriarca(status_novo):
    return Patriarca.objects.create(
        idexternopatriarca=uuid4(),
        strsiglapatriarca="PMVV",
        strnome="Prefeitura Municipal de Vila Velha",
        datcriacao=timezone.now(),
        idstatusprogresso=status_novo,
    )


@pytest.fixture
def carga(status_processando, tipo_organograma):
    return TokenEnvioCarga.objects.create(
        strtoken="credencial-que-nao-pode-ser-exposta",
        idtipocarga=tipo_organograma,
        idstatusprogresso=status_processando,
    )


def conteudo_organograma(
    linhas=None,
    cabecalho=("Numero", "Nome", "Sigla"),
    segunda_aba=None,
):
    linhas = linhas or [
        (1, "Prefeitura Municipal de Vila Velha", "PMVV"),
        ("1.1", "Secretaria de Administração", "SEMAD"),
    ]
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Organograma"
    worksheet.append(cabecalho)
    for linha in linhas:
        worksheet.append(linha)
    if segunda_aba is not None:
        worksheet = workbook.create_sheet("Segunda aba")
        worksheet.append(segunda_aba[0])
        for linha in segunda_aba[1]:
            worksheet.append(linha)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def arquivo_planilha(nome="organograma.xlsx", conteudo=None):
    conteudo = conteudo if conteudo is not None else conteudo_organograma()
    return SimpleUploadedFile(
        nome,
        conteudo,
        content_type=(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        ),
    )


def conteudo_lotacao(linhas=None, segunda_aba=None):
    linhas = linhas or [
        (
            "SEMAD",
            "Secretaria de Administração",
            "Maria da Silva",
            "903.961.004-57",
            "  auxiliar administrativo  ",
            "SEMAD",
            "Secretaria de Administração",
        )
    ]
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "ADMINISTRACAO"
    header = (
        "Sigla",
        "Secretaria/Órgão",
        "Nome do Servidor",
        "CPF",
        "Ocupação/Cargo",
        "Sigla Setor",
        "Setor de Lotação",
    )
    worksheet.append(("", "Patriarca:", "[Município Lotação]"))
    worksheet.append(header)
    for linha in linhas:
        worksheet.append(linha)
    if segunda_aba is not None:
        worksheet = workbook.create_sheet("SEGUNDA SECRETARIA")
        worksheet.append(header)
        for linha in segunda_aba:
            worksheet.append(linha)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


@pytest.mark.django_db(transaction=True)
class TestAutenticacao:
    @pytest.mark.parametrize(
        "url",
        [
            PATRIARCAS_URL,
            IMPORTACOES_URL,
            ITENS_LOTACAO_URL,
            PROCESSOS_URL,
            CARGAS_URL,
            HISTORICO_URL,
            EVENTOS_INTEGRACAO_URL,
            f"{BASE_URL}/status-carga/",
            f"{BASE_URL}/status-progresso/",
            f"{BASE_URL}/tipos-carga/",
        ],
    )
    def test_listas_exigem_autenticacao(self, client_anonimo_carga, url):
        response = client_anonimo_carga.get(url)
        assert response.status_code == 401


@pytest.mark.django_db(transaction=True)
class TestPermissoesDeclaradas:
    @pytest.mark.parametrize(
        ("viewset", "model_name"),
        [
            (PatriarcaViewSet, "patriarca"),
            (StatusCargaViewSet, "statuscarga"),
            (StatusProgressoViewSet, "statusprogresso"),
            (TipoCargaViewSet, "tipocarga"),
            (CargaOrgLotViewSet, "tokenenviocarga"),
            (HistoricoCargaViewSet, "detalhestatuscarga"),
            (ItemOrganogramaViewSet, "itemorganograma"),
        ],
    )
    def test_crud_possui_mapa_completo(self, viewset, model_name):
        expected = {
            "list": f"view_{model_name}",
            "retrieve": f"view_{model_name}",
            "create": f"add_{model_name}",
            "update": f"change_{model_name}",
            "partial_update": f"change_{model_name}",
            "destroy": f"delete_{model_name}",
        }
        if viewset is PatriarcaViewSet:
            expected["download_organograma"] = "view_importacaoplanilha"
        if viewset is ItemOrganogramaViewSet:
            expected["impacto_exclusao"] = "view_itemorganograma"
        assert viewset.permission_map == expected

    def test_importacao_inclui_upload_e_download(self):
        assert ImportacaoPlanilhaViewSet.permission_map["upload"] == (
            "add_importacaoplanilha"
        )
        assert ImportacaoPlanilhaViewSet.permission_map[
            "criar_organograma_inicial"
        ] == "add_importacaoplanilha"
        assert ImportacaoPlanilhaViewSet.permission_map["download"] == (
            "view_importacaoplanilha"
        )
        assert ImportacaoPlanilhaViewSet.permission_map[
            "download_organograma_atualizado"
        ] == "view_importacaoplanilha"
        assert ImportacaoPlanilhaViewSet.permission_map["relatorio"] == (
            "view_importacaoplanilha"
        )

    def test_processo_e_somente_leitura(self):
        assert ProcessoCargaInicialViewSet.permission_map == {
            "list": "view_patriarca",
            "retrieve": "view_patriarca",
        }

    def test_itens_lotacao_sao_somente_leitura(self):
        assert ItemLotacaoViewSet.permission_map == {
            "list": "view_importacaoplanilha",
            "retrieve": "view_importacaoplanilha",
        }

    def test_conjunto_lotacao_possui_permissoes_explicitas(self):
        assert ConjuntoLotacaoViewSet.permission_map == {
            "list": "view_importacaoplanilha",
            "retrieve": "view_importacaoplanilha",
            "create": "add_importacaoplanilha",
            "adicionar_importacao": "add_importacaoplanilha",
            "homologar": "change_importacaoplanilha",
            "ativar": "change_importacaoplanilha",
            "cancelar": "change_importacaoplanilha",
            "download_consolidado": "view_importacaoplanilha",
            "historico": "view_importacaoplanilha",
            "resumo_organograma": "view_importacaoplanilha",
        }

    def test_envio_organograma_possui_permissoes_explicitas(self):
        assert EnvioOrganogramaProdestViewSet.permission_map == {
            "list": "view_importacaoplanilha",
            "retrieve": "view_importacaoplanilha",
            "create": "change_importacaoplanilha",
            "configuracao": "view_importacaoplanilha",
            "ultimas_prodest": "view_importacaoplanilha",
            "testar_conexao": "change_importacaoplanilha",
            "consultar_status": "change_importacaoplanilha",
            "consultar_relatorio": "change_importacaoplanilha",
            "acompanhamento": "view_importacaoplanilha",
            "eventos": "view_importacaoplanilha",
        }

    def test_eventos_integracao_sao_somente_leitura(self):
        assert EventoIntegracaoProdestViewSet.permission_map == {
            "list": "view_importacaoplanilha",
            "retrieve": "view_importacaoplanilha",
        }


@pytest.mark.django_db(transaction=True)
class TestPatriarca:
    def test_crud_completo(
        self,
        client_gestor_carga_lot,
        gestor_carga_lot,
        status_novo,
        status_processando,
    ):
        external_id = uuid4()
        create_response = client_gestor_carga_lot.post(
            PATRIARCAS_URL,
            {
                "idexternopatriarca": str(external_id),
                "strsiglapatriarca": "  pmc  ",
                "strnome": "  Prefeitura Municipal de Cariacica  ",
            },
            format="json",
        )

        assert create_response.status_code == 201
        pk = create_response.data["idpatriarca"]
        assert create_response.data["strsiglapatriarca"] == "PMC"
        assert create_response.data["idstatusprogresso_id"] == status_novo.pk
        assert create_response.data["created_by_id"] == gestor_carga_lot.pk
        process = create_response.data["processo_ativo"]
        assert process["strfaseatual"] == FaseProcessoCargaInicial.NOVO
        assert (
            process["strstatusorganograma"]
            == StatusOrganogramaProcesso.NAO_INICIADO
        )
        assert process["strstatuslotacao"] == StatusLotacaoProcesso.BLOQUEADA
        assert process["historico"][0]["stracao"] == "PROCESSO_CRIADO"

        list_response = client_gestor_carga_lot.get(PATRIARCAS_URL)
        assert list_response.status_code == 200
        assert list_response.data["count"] == 1

        retrieve_response = client_gestor_carga_lot.get(f"{PATRIARCAS_URL}{pk}/")
        assert retrieve_response.status_code == 200

        patch_response = client_gestor_carga_lot.patch(
            f"{PATRIARCAS_URL}{pk}/",
            {"idstatusprogresso_id": status_processando.pk},
            format="json",
        )
        assert patch_response.status_code == 200
        assert patch_response.data["status"]["strdescricao"] == "Em progresso"
        assert patch_response.data["datalteracao"] is not None

        delete_response = client_gestor_carga_lot.delete(f"{PATRIARCAS_URL}{pk}/")
        assert delete_response.status_code == 204
        assert not Patriarca.objects.filter(pk=pk).exists()

    def test_rejeita_sigla_duplicada_sem_diferenciar_caixa(
        self, client_gestor_carga_lot, status_novo, patriarca
    ):
        response = client_gestor_carga_lot.post(
            PATRIARCAS_URL,
            {
                "idexternopatriarca": str(uuid4()),
                "strsiglapatriarca": "pmvv",
                "strnome": "Outra instituição",
            },
            format="json",
        )
        assert response.status_code == 400
        assert "strsiglapatriarca" in response.data["errors"]


@pytest.mark.django_db(transaction=True)
class TestProcessoCargaInicial:
    def test_lista_filtra_e_detalha_processo_ativo(
        self,
        client_gestor_carga_lot,
        patriarca,
    ):
        process = ProcessoCargaInicial.objects.create(idpatriarca=patriarca)

        list_response = client_gestor_carga_lot.get(
            PROCESSOS_URL,
            {"patriarca": patriarca.pk, "ativo": "true"},
        )

        assert list_response.status_code == 200
        assert list_response.data["count"] == 1
        assert (
            list_response.data["results"][0]["idprocessocargainicial"]
            == process.pk
        )
        detail_response = client_gestor_carga_lot.get(
            f"{PROCESSOS_URL}{process.pk}/"
        )
        assert detail_response.status_code == 200
        assert detail_response.data["strfaseatual"] == "NOVO"

    def test_rejeita_filtros_invalidos(
        self,
        client_gestor_carga_lot,
    ):
        invalid_patriarch = client_gestor_carga_lot.get(
            PROCESSOS_URL,
            {"patriarca": "abc"},
        )
        invalid_active = client_gestor_carga_lot.get(
            PROCESSOS_URL,
            {"ativo": "sim"},
        )

        assert invalid_patriarch.status_code == 400
        assert "patriarca" in invalid_patriarch.data["errors"]
        assert invalid_active.status_code == 400
        assert "ativo" in invalid_active.data["errors"]


@pytest.mark.django_db(transaction=True)
class TestCatalogos:
    @pytest.mark.parametrize(
        ("url", "id_field", "description_field", "payload"),
        [
            (
                f"{BASE_URL}/status-progresso/",
                "idstatusprogresso",
                "strdescricao",
                {"idstatusprogresso": 9, "strdescricao": "Recebido"},
            ),
            (
                f"{BASE_URL}/tipos-carga/",
                "idtipocarga",
                "strdescricao",
                {"idtipocarga": 9, "strdescricao": "Teste"},
            ),
            (
                f"{BASE_URL}/status-carga/",
                "idstatuscarga",
                "strdescricao",
                {
                    "idstatuscarga": 20,
                    "strdescricao": "Recebida",
                    "flgsucesso": 0,
                },
            ),
        ],
    )
    def test_crud_completo(
        self,
        client_gestor_carga_lot,
        url,
        id_field,
        description_field,
        payload,
    ):
        create_response = client_gestor_carga_lot.post(url, payload, format="json")
        assert create_response.status_code == 201
        pk = create_response.data[id_field]

        assert client_gestor_carga_lot.get(url).status_code == 200
        assert client_gestor_carga_lot.get(f"{url}{pk}/").status_code == 200

        patch_response = client_gestor_carga_lot.patch(
            f"{url}{pk}/",
            {description_field: "Atualizado"},
            format="json",
        )
        assert patch_response.status_code == 200
        assert patch_response.data[description_field] == "Atualizado"

        delete_response = client_gestor_carga_lot.delete(f"{url}{pk}/")
        assert delete_response.status_code == 204


@pytest.mark.django_db(transaction=True)
class TestImportacaoPlanilha:
    def test_cria_organograma_inicial_e_importa_o_xlsx_gerado(
        self,
        client_gestor_carga_lot,
        gestor_carga_lot,
        patriarca,
        tipo_organograma,
    ):
        response = client_gestor_carga_lot.post(
            f"{IMPORTACOES_URL}criar-organograma-inicial/",
            {
                "idpatriarca_id": patriarca.pk,
                "strreferencia": "Estrutura inicial padrão",
            },
            format="json",
        )

        assert response.status_code == 201
        importacao_id = response.data["idimportacaoplanilha"]
        assert response.data["strnomearquivo"] == (
            "organograma_inicial_PMVV.xlsx"
        )
        assert response.data["strreferencia"] == "Estrutura inicial padrão"
        assert response.data["strstatusimportacao"] == (
            StatusImportacaoLocal.INCORPORADA
        )
        assert response.data["flgpossuialteracoes"] is False
        assert response.data["flgcriarunidadesespelho"] is False
        assert response.data["total_itens_organograma"] == 9
        assert response.data["created_by_id"] == gestor_carga_lot.pk
        assert response.data["relatorio_validacao"]["inttotallinhas"] == 9
        assert response.data["relatorio_validacao"]["inttotalerros"] == 0

        importacao = ImportacaoPlanilha.objects.get(pk=importacao_id)
        assert importacao.idtipocarga_id == tipo_organograma.pk
        assert hashlib.sha256(bytes(importacao.binconteudo)).hexdigest() == (
            importacao.strhasharquivo
        )
        items = list(
            importacao.itens_organograma.select_related("idpai").order_by(
                "strnumero"
            )
        )
        assert len(items) == 9
        assert not any(item.flggeradoautomaticamente for item in items)
        root = next(item for item in items if item.idpai_id is None)
        assert (root.strnome, root.strsigla) == (
            "PREFEITURA MUNICIPAL DE VILA VELHA",
            "PMVV",
        )
        for acronym in ("GAB", "SEFAZ", "SESA", "SEMED"):
            organ = next(
                item
                for item in items
                if item.strsigla == acronym and item.idpai_id == root.pk
            )
            mirror = next(
                item
                for item in items
                if item.strsigla == acronym and item.idpai_id == organ.pk
            )
            assert mirror.strnome == organ.strnome

        download = client_gestor_carga_lot.get(
            f"{IMPORTACOES_URL}{importacao_id}/download/"
        )
        assert download.status_code == 200
        assert download.content == bytes(importacao.binconteudo)
        workbook = load_workbook(BytesIO(download.content), data_only=True)
        assert workbook.sheetnames == ["Organograma"]
        assert list(workbook["Organograma"].values) == [
            ("Numero", "Nome", "Sigla", "Pai_Sigla", "Responsavel", "Cargo"),
            (
                "1",
                "Prefeitura Municipal de Vila Velha",
                "PMVV",
                None,
                None,
                None,
            ),
            ("1.1", "Gabinete do Prefeito", "GAB", "PMVV", None, None),
            ("1.1.1", "Gabinete do Prefeito", "GAB", "GAB", None, None),
            (
                "1.2",
                "Secretaria Municipal de Fazenda",
                "SEFAZ",
                "PMVV",
                None,
                None,
            ),
            (
                "1.2.1",
                "Secretaria Municipal de Fazenda",
                "SEFAZ",
                "SEFAZ",
                None,
                None,
            ),
            (
                "1.3",
                "Secretaria Municipal de Saúde",
                "SESA",
                "PMVV",
                None,
                None,
            ),
            (
                "1.3.1",
                "Secretaria Municipal de Saúde",
                "SESA",
                "SESA",
                None,
                None,
            ),
            (
                "1.4",
                "Secretaria Municipal de Educação",
                "SEMED",
                "PMVV",
                None,
                None,
            ),
            (
                "1.4.1",
                "Secretaria Municipal de Educação",
                "SEMED",
                "SEMED",
                None,
                None,
            ),
        ]

    def test_criar_organograma_inicial_exige_patriarca_valido(
        self,
        client_gestor_carga_lot,
    ):
        response = client_gestor_carga_lot.post(
            f"{IMPORTACOES_URL}criar-organograma-inicial/",
            {"idpatriarca_id": 999999},
            format="json",
        )

        assert response.status_code == 400
        assert "idpatriarca_id" in response.data["errors"]
        assert not ImportacaoPlanilha.objects.exists()

    def test_upload_lista_detalha_e_download(
        self,
        client_gestor_carga_lot,
        gestor_carga_lot,
        patriarca,
        tipo_organograma,
    ):
        conteudo = conteudo_organograma()
        upload_response = client_gestor_carga_lot.post(
            f"{IMPORTACOES_URL}upload/",
            {
                "idpatriarca_id": patriarca.pk,
                "idtipocarga_id": tipo_organograma.pk,
                "strreferencia": "Processo E-Docs 2026-ABC123",
                "arquivo": arquivo_planilha(conteudo=conteudo),
            },
            format="multipart",
        )

        assert upload_response.status_code == 201
        pk = upload_response.data["idimportacaoplanilha"]
        assert upload_response.data["strnomearquivo"] == "organograma.xlsx"
        assert upload_response.data["inttamanhoarquivo"] == len(conteudo)
        assert upload_response.data["strhasharquivo"] == hashlib.sha256(
            conteudo
        ).hexdigest()
        assert upload_response.data["status"]["strdescricao"] == "Novo"
        assert upload_response.data["strreferencia"] == (
            "Processo E-Docs 2026-ABC123"
        )
        assert upload_response.data["created_by_id"] == gestor_carga_lot.pk
        assert "binconteudo" not in upload_response.data
        assert (
            upload_response.data["strstatusimportacao"]
            == StatusImportacaoLocal.INCORPORADA
        )
        assert (
            upload_response.data["processo"]["strfaseatual"]
            == FaseProcessoCargaInicial.PREPARANDO_ORGANOGRAMA
        )
        assert (
            upload_response.data["relatorio_validacao"]["strresultado"]
            == ResultadoRelatorioValidacao.VALIDO
        )

        importacao = ImportacaoPlanilha.objects.get(pk=pk)
        assert bytes(importacao.binconteudo) == conteudo
        assert importacao.relatorio_validacao.inttotalerros == 0

        list_response = client_gestor_carga_lot.get(IMPORTACOES_URL)
        assert list_response.status_code == 200
        assert list_response.data["count"] == 1
        listed_import = list_response.data["results"][0]
        assert "binconteudo" not in listed_import
        report_summary = listed_import["relatorio_validacao"]
        assert report_summary["strresultado"] == ResultadoRelatorioValidacao.VALIDO
        assert "jsonresumo" not in report_summary
        assert "ocorrencias" not in report_summary
        assert "casos" not in report_summary
        assert "historico" not in listed_import["processo"]
        assert "processo_ativo" not in listed_import["patriarca"]

        filtered_response = client_gestor_carga_lot.get(
            IMPORTACOES_URL,
            {"patriarca": patriarca.pk, "tipo_carga": tipo_organograma.pk},
        )
        assert filtered_response.status_code == 200
        assert filtered_response.data["count"] == 1

        empty_response = client_gestor_carga_lot.get(
            IMPORTACOES_URL,
            {"patriarca": patriarca.pk + 1},
        )
        assert empty_response.status_code == 200
        assert empty_response.data["count"] == 0

        detail_response = client_gestor_carga_lot.get(f"{IMPORTACOES_URL}{pk}/")
        assert detail_response.status_code == 200
        assert detail_response.data["patriarca"]["idpatriarca"] == patriarca.pk

        download_response = client_gestor_carga_lot.get(
            f"{IMPORTACOES_URL}{pk}/download/"
        )
        assert download_response.status_code == 200
        assert download_response.content == conteudo
        assert download_response["Content-Length"] == str(len(conteudo))
        assert "organograma.xlsx" in download_response["Content-Disposition"]
        assert download_response["X-Content-Type-Options"] == "nosniff"

        report_response = client_gestor_carga_lot.get(
            f"{IMPORTACOES_URL}{pk}/relatorio/"
        )
        assert report_response.status_code == 200
        assert report_response.data["strresultado"] == "VALIDO"
        assert "jsonresumo" in report_response.data
        assert report_response.data["ocorrencias"] == []

    def test_listagem_nao_expande_casos_de_validacao(
        self,
        client_gestor_carga_lot,
        patriarca,
        tipo_organograma,
        django_assert_max_num_queries,
    ):
        uploaded = client_gestor_carga_lot.post(
            f"{IMPORTACOES_URL}upload/",
            {
                "idpatriarca_id": patriarca.pk,
                "idtipocarga_id": tipo_organograma.pk,
                "arquivo": arquivo_planilha(),
            },
            format="multipart",
        )
        assert uploaded.status_code == 201
        report = RelatorioValidacaoImportacao.objects.get(
            idimportacaoplanilha_id=uploaded.data["idimportacaoplanilha"]
        )
        CasoValidacaoImportacao.objects.bulk_create(
            [
                CasoValidacaoImportacao(
                    idrelatoriovalidacao=report,
                    strchaveagrupamento=f"teste:{index}",
                    strcodigo="caso_teste",
                    strseveridade="AVISO",
                    strtitulo=f"Caso {index}",
                    strmensagem="Caso criado para testar a listagem resumida.",
                )
                for index in range(50)
            ]
        )

        with django_assert_max_num_queries(20):
            response = client_gestor_carga_lot.get(
                IMPORTACOES_URL,
                {
                    "patriarca": patriarca.pk,
                    "tipo_carga": tipo_organograma.pk,
                },
            )

        assert response.status_code == 200
        item = response.data["results"][0]
        assert item["total_casos_pendentes"] == 50
        assert "casos" not in item["relatorio_validacao"]

    def test_rejeita_filtro_nao_numerico(self, client_gestor_carga_lot):
        response = client_gestor_carga_lot.get(
            IMPORTACOES_URL,
            {"patriarca": "qualquer"},
        )
        assert response.status_code == 400
        assert "patriarca" in response.data["errors"]

    def test_post_padrao_tambem_realiza_upload(
        self,
        client_gestor_carga_lot,
        patriarca,
        tipo_organograma,
    ):
        response = client_gestor_carga_lot.post(
            IMPORTACOES_URL,
            {
                "idpatriarca_id": patriarca.pk,
                "idtipocarga_id": tipo_organograma.pk,
                "arquivo": arquivo_planilha(),
            },
            format="multipart",
        )
        assert response.status_code == 201

    def test_rejeita_extensao_invalida(
        self,
        client_gestor_carga_lot,
        patriarca,
        tipo_organograma,
    ):
        response = client_gestor_carga_lot.post(
            f"{IMPORTACOES_URL}upload/",
            {
                "idpatriarca_id": patriarca.pk,
                "idtipocarga_id": tipo_organograma.pk,
                "arquivo": arquivo_planilha(nome="dados.exe"),
            },
            format="multipart",
        )
        assert response.status_code == 400
        assert "arquivo" in response.data["errors"]

    def test_rejeita_arquivo_acima_do_limite(
        self,
        client_gestor_carga_lot,
        patriarca,
        tipo_organograma,
        settings,
    ):
        settings.CARGA_ORG_LOT_MAX_UPLOAD_BYTES = 4
        response = client_gestor_carga_lot.post(
            f"{IMPORTACOES_URL}upload/",
            {
                "idpatriarca_id": patriarca.pk,
                "idtipocarga_id": tipo_organograma.pk,
                "arquivo": arquivo_planilha(conteudo=b"12345"),
            },
            format="multipart",
        )
        assert response.status_code == 400
        assert "arquivo" in response.data["errors"]

    @pytest.mark.parametrize(
        ("cabecalho", "linhas"),
        [
            (
                ("Numero", "Nome", "Sigla", "Pai_Sigla"),
                [
                    (1, "Prefeitura", "PMVV", ""),
                    (2, "Secretaria", "SEC", "PMVV"),
                    (3, "Gerência", "GER", "SEC"),
                ],
            ),
            (
                ("Nome", "Sigla", "Pai_Sigla"),
                [
                    ("Prefeitura", "PMVV", ""),
                    ("Secretaria", "SEC", "PMVV"),
                    ("Gerência", "GER", "SEC"),
                ],
            ),
            (
                (
                    "pode ser isso",
                    "ESTRUTURA ORGANIZACIONAL DO MUNICÍPIO",
                    "SIGLAS ATUAIS",
                ),
                [
                    ("", "Prefeitura", "PMVV"),
                    (1, "Secretaria", "SEC"),
                    ("1.1", "Gerência", "GER"),
                ],
            ),
        ],
    )
    def test_aceita_formatos_flexiveis_de_hierarquia(
        self,
        client_gestor_carga_lot,
        patriarca,
        tipo_organograma,
        cabecalho,
        linhas,
    ):
        response = client_gestor_carga_lot.post(
            f"{IMPORTACOES_URL}upload/",
            {
                "idpatriarca_id": patriarca.pk,
                "idtipocarga_id": tipo_organograma.pk,
                "arquivo": arquivo_planilha(
                    conteudo=conteudo_organograma(linhas, cabecalho)
                ),
            },
            format="multipart",
        )

        assert response.status_code == 201
        assert ImportacaoPlanilha.objects.count() == 1

    @pytest.mark.parametrize(
        "linhas",
        [
            [
                (1, "Prefeitura", "PMVV"),
                ("1.1", "Secretaria de Administração", "SEMAD"),
                ("1.1.1", "Secretaria de Administração", "SEMAD"),
            ],
            [
                (1, "Prefeitura", "PMVV"),
                ("1.1", "Órgão A", "ORGA"),
                ("1.1.1", "Unidade Comum A", "COMUM"),
                ("1.2", "Órgão B", "ORGB"),
                ("1.2.1", "Unidade Comum B", "COMUM"),
            ],
        ],
    )
    def test_aceita_sigla_repetida_com_pai_ou_em_outro_ramo(
        self,
        client_gestor_carga_lot,
        patriarca,
        tipo_organograma,
        linhas,
    ):
        response = client_gestor_carga_lot.post(
            f"{IMPORTACOES_URL}upload/",
            {
                "idpatriarca_id": patriarca.pk,
                "idtipocarga_id": tipo_organograma.pk,
                "arquivo": arquivo_planilha(
                    conteudo=conteudo_organograma(linhas=linhas)
                ),
            },
            format="multipart",
        )

        assert response.status_code == 201

    def test_cria_unidades_espelho_quando_opcao_esta_marcada(
        self,
        client_gestor_carga_lot,
        patriarca,
        tipo_organograma,
    ):
        content = conteudo_organograma(
            linhas=[
                (1, "Prefeitura", "PMVV"),
                ("1.1", "Secretaria de Administração", "SEMAD"),
                ("1.2", "Secretaria de Saúde", "SESA"),
            ]
        )

        response = client_gestor_carga_lot.post(
            f"{IMPORTACOES_URL}upload/",
            {
                "idpatriarca_id": patriarca.pk,
                "idtipocarga_id": tipo_organograma.pk,
                "criar_unidades_espelho": "true",
                "arquivo": arquivo_planilha(conteudo=content),
            },
            format="multipart",
        )

        assert response.status_code == 201
        assert response.data["flgcriarunidadesespelho"] is True
        assert response.data["flgpossuialteracoes"] is True
        assert response.data["strstatusimportacao"] == StatusImportacaoLocal.ALTERADA
        assert response.data["organograma_atualizado_download_url"] is not None
        assert response.data["total_itens_organograma"] == 5
        summary = response.data["relatorio_validacao"]["jsonresumo"]
        assert summary["criar_unidades_espelho"] is True
        assert summary["total_unidades_espelho_criadas"] == 2

        generated = ItemOrganograma.objects.filter(
            idimportacaoplanilha_id=response.data["idimportacaoplanilha"],
            flggeradoautomaticamente=True,
        ).order_by("strnumero")
        assert list(
            generated.values_list(
                "strnome",
                "strsigla",
                "idpai__strsigla",
                "strnumero",
                "strabaorigem",
                "intlinhaorigem",
            )
        ) == [
            (
                "SECRETARIA DE ADMINISTRAÇÃO",
                "SEMAD",
                "SEMAD",
                "1.1.1",
                "Gerada automaticamente",
                0,
            ),
            (
                "SECRETARIA DE SAÚDE",
                "SESA",
                "SESA",
                "1.2.1",
                "Gerada automaticamente",
                0,
            ),
        ]
        items_response = client_gestor_carga_lot.get(
            ITENS_ORGANOGRAMA_URL,
            {"importacao": response.data["idimportacaoplanilha"]},
        )
        assert items_response.status_code == 200
        assert sum(
            item["flggeradoautomaticamente"]
            for item in items_response.data["results"]
        ) == 2

    def test_nao_duplica_unidade_espelho_que_ja_existe(
        self,
        client_gestor_carga_lot,
        patriarca,
        tipo_organograma,
    ):
        content = conteudo_organograma(
            linhas=[
                (1, "Prefeitura", "PMVV"),
                ("1.1", "Secretaria de Administração", "SEMAD"),
                ("1.1.1", "Secretaria de Administração", "SEMAD"),
            ]
        )

        response = client_gestor_carga_lot.post(
            f"{IMPORTACOES_URL}upload/",
            {
                "idpatriarca_id": patriarca.pk,
                "idtipocarga_id": tipo_organograma.pk,
                "criar_unidades_espelho": True,
                "arquivo": arquivo_planilha(conteudo=content),
            },
            format="multipart",
        )

        assert response.status_code == 201
        assert response.data["total_itens_organograma"] == 3
        assert response.data["flgcriarunidadesespelho"] is True
        assert response.data["flgpossuialteracoes"] is False
        assert response.data["strstatusimportacao"] == (
            StatusImportacaoLocal.INCORPORADA
        )
        assert not ItemOrganograma.objects.filter(
            flggeradoautomaticamente=True
        ).exists()
        summary = response.data["relatorio_validacao"]["jsonresumo"]
        assert summary["total_unidades_espelho_criadas"] == 0

    def test_nao_cria_unidade_espelho_sem_a_opcao(
        self,
        client_gestor_carga_lot,
        patriarca,
        tipo_organograma,
    ):
        content = conteudo_organograma(
            linhas=[
                (1, "Prefeitura", "PMVV"),
                ("1.1", "Secretaria de Administração", "SEMAD"),
            ]
        )

        response = client_gestor_carga_lot.post(
            f"{IMPORTACOES_URL}upload/",
            {
                "idpatriarca_id": patriarca.pk,
                "idtipocarga_id": tipo_organograma.pk,
                "arquivo": arquivo_planilha(conteudo=content),
            },
            format="multipart",
        )

        assert response.status_code == 201
        assert response.data["flgcriarunidadesespelho"] is False
        assert response.data["total_itens_organograma"] == 2
        assert not ItemOrganograma.objects.filter(
            flggeradoautomaticamente=True
        ).exists()

    def test_rejeita_opcao_de_unidade_espelho_em_importacao_de_lotacao(
        self,
        client_gestor_carga_lot,
        patriarca,
        tipo_organograma,
        tipo_lotacao,
    ):
        organogram = client_gestor_carga_lot.post(
            f"{IMPORTACOES_URL}upload/",
            {
                "idpatriarca_id": patriarca.pk,
                "idtipocarga_id": tipo_organograma.pk,
                "arquivo": arquivo_planilha(),
            },
            format="multipart",
        )
        assert organogram.status_code == 201

        response = client_gestor_carga_lot.post(
            f"{IMPORTACOES_URL}upload/",
            {
                "idpatriarca_id": patriarca.pk,
                "idtipocarga_id": tipo_lotacao.pk,
                "idimportacaoorganograma_id": organogram.data[
                    "idimportacaoplanilha"
                ],
                "criar_unidades_espelho": True,
                "arquivo": arquivo_planilha(nome="lotacao.xlsx"),
            },
            format="multipart",
        )

        assert response.status_code == 400
        assert "criar_unidades_espelho" in response.data["errors"]

    def test_rejeita_conflito_parcial_ao_criar_unidade_espelho(
        self,
        client_gestor_carga_lot,
        patriarca,
        tipo_organograma,
    ):
        content = conteudo_organograma(
            linhas=[
                (1, "Prefeitura", "PMVV"),
                ("1.1", "Secretaria de Administração", "SEMAD"),
                ("1.1.1", "Secretaria de Administração", "OUTRA"),
            ]
        )

        response = client_gestor_carga_lot.post(
            f"{IMPORTACOES_URL}upload/",
            {
                "idpatriarca_id": patriarca.pk,
                "idtipocarga_id": tipo_organograma.pk,
                "criar_unidades_espelho": True,
                "arquivo": arquivo_planilha(conteudo=content),
            },
            format="multipart",
        )

        assert response.status_code == 400
        report = response.data["errors"]["arquivo"]
        assert "conflito_unidade_espelho" in {
            error["codigo"] for error in report["erros"]
        }
        importacao = ImportacaoPlanilha.objects.get(
            pk=report["idimportacaoplanilha"]
        )
        assert importacao.flgcriarunidadesespelho is True
        assert not importacao.itens_organograma.exists()

    def test_resolve_pai_com_mesma_sigla_sem_numeracao(
        self,
        client_gestor_carga_lot,
        patriarca,
        tipo_organograma,
    ):
        content = conteudo_organograma(
            cabecalho=("Nome", "Sigla", "Pai_Sigla"),
            linhas=[
                ("Prefeitura", "PMVV", ""),
                ("Secretaria de Administração", "SEMAD", "PMVV"),
                ("Secretaria de Administração", "SEMAD", "SEMAD"),
            ],
        )

        response = client_gestor_carga_lot.post(
            f"{IMPORTACOES_URL}upload/",
            {
                "idpatriarca_id": patriarca.pk,
                "idtipocarga_id": tipo_organograma.pk,
                "arquivo": arquivo_planilha(conteudo=content),
            },
            format="multipart",
        )

        assert response.status_code == 201
        repeated = list(ItemOrganograma.objects.filter(strsigla="SEMAD"))
        assert len(repeated) == 2
        organ = next(item for item in repeated if item.idpai.strsigla == "PMVV")
        unit = next(item for item in repeated if item.idpai_id == organ.pk)
        assert unit.strnome == organ.strnome

    @pytest.mark.parametrize(
        ("linhas", "codigo"),
        [
            (
                [
                    (1, "Prefeitura", "PMVV"),
                    ("1.1", "Órgão A", "ORG"),
                    ("1.2", "Órgão B", "ORG"),
                ],
                "sigla_duplicada_entre_irmaos",
            ),
            (
                [
                    (1, "Prefeitura", "PMVV"),
                    ("1.1", "Órgão A", "ORGA"),
                    ("1.1.1", "Unidade A", "UNI"),
                    ("1.1.2", "Unidade B", "UNI"),
                ],
                "sigla_duplicada_no_orgao",
            ),
            (
                [
                    (1, "Prefeitura", "PMVV"),
                    ("1.1", "Órgão com sigla inválida", "PMVV"),
                ],
                "sigla_igual_ao_pai_no_nivel_2",
            ),
        ],
    )
    def test_rejeita_siglas_invalidas_no_escopo_hierarquico(
        self,
        client_gestor_carga_lot,
        patriarca,
        tipo_organograma,
        linhas,
        codigo,
    ):
        response = client_gestor_carga_lot.post(
            f"{IMPORTACOES_URL}upload/",
            {
                "idpatriarca_id": patriarca.pk,
                "idtipocarga_id": tipo_organograma.pk,
                "arquivo": arquivo_planilha(
                    conteudo=conteudo_organograma(linhas=linhas)
                ),
            },
            format="multipart",
        )

        assert response.status_code == 400
        errors = response.data["errors"]["arquivo"]["erros"]
        assert codigo in {error["codigo"] for error in errors}

    def test_rejeita_sigla_repetida_em_ramos_do_mesmo_orgao(
        self,
        client_gestor_carga_lot,
        patriarca,
        tipo_organograma,
    ):
        content = conteudo_organograma(
            linhas=[
                (1, "Prefeitura", "PMVV"),
                ("1.1", "Órgão A", "ORGA"),
                ("1.1.1", "Diretoria A", "DIRA"),
                ("1.1.1.1", "Unidade Comum A", "COMUM"),
                ("1.1.2", "Diretoria B", "DIRB"),
                ("1.1.2.1", "Unidade Comum B", "COMUM"),
            ]
        )

        response = client_gestor_carga_lot.post(
            f"{IMPORTACOES_URL}upload/",
            {
                "idpatriarca_id": patriarca.pk,
                "idtipocarga_id": tipo_organograma.pk,
                "arquivo": arquivo_planilha(conteudo=content),
            },
            format="multipart",
        )

        assert response.status_code == 400
        errors = response.data["errors"]["arquivo"]["erros"]
        matching = [
            error
            for error in errors
            if error["codigo"] == "sigla_duplicada_no_orgao"
        ]
        assert len(matching) == 2

    @pytest.mark.parametrize(
        ("name", "acronym", "number", "code"),
        [
            ("N" * 101, "ORGA", "1.1", "nome_orgao_muito_longo"),
            ("Órgão A", "A" * 21, "1.1", "sigla_orgao_muito_longa"),
            ("N" * 101, "UNIA", "1.1.1", "nome_unidade_muito_longo"),
            ("Unidade A", "A" * 26, "1.1.1", "sigla_unidade_muito_longa"),
        ],
    )
    def test_rejeita_tamanhos_superiores_aos_limites_da_prodest(
        self,
        client_gestor_carga_lot,
        patriarca,
        tipo_organograma,
        name,
        acronym,
        number,
        code,
    ):
        rows = [(1, "Prefeitura", "PMVV"), ("1.1", "Órgão A", "ORGA")]
        if number == "1.1":
            rows[1] = (number, name, acronym)
        else:
            rows.append((number, name, acronym))

        response = client_gestor_carga_lot.post(
            f"{IMPORTACOES_URL}upload/",
            {
                "idpatriarca_id": patriarca.pk,
                "idtipocarga_id": tipo_organograma.pk,
                "arquivo": arquivo_planilha(
                    conteudo=conteudo_organograma(linhas=rows)
                ),
            },
            format="multipart",
        )

        assert response.status_code == 400
        errors = response.data["errors"]["arquivo"]["erros"]
        assert code in {error["codigo"] for error in errors}

    def test_aceita_tamanhos_exatamente_nos_limites_da_prodest(
        self,
        client_gestor_carga_lot,
        patriarca,
        tipo_organograma,
    ):
        response = client_gestor_carga_lot.post(
            f"{IMPORTACOES_URL}upload/",
            {
                "idpatriarca_id": patriarca.pk,
                "idtipocarga_id": tipo_organograma.pk,
                "arquivo": arquivo_planilha(
                    conteudo=conteudo_organograma(
                        linhas=[
                            (1, "Prefeitura", "PMVV"),
                            ("1.1", "O" * 100, "A" * 20),
                            ("1.1.1", "U" * 100, "B" * 25),
                        ]
                    )
                ),
            },
            format="multipart",
        )

        assert response.status_code == 201

    def test_importa_responsavel_e_cargo_quando_informados(
        self,
        client_gestor_carga_lot,
        patriarca,
        tipo_organograma,
    ):
        content = conteudo_organograma(
            cabecalho=("Numero", "Nome", "Sigla", "Responsável", "Cargo/Função"),
            linhas=[
                (1, "Prefeitura", "PMVV", "Maria da Silva", "Prefeita"),
                (
                    "1.1",
                    "Secretaria de Administração",
                    "SEMAD",
                    "João Souza",
                    "Secretário Municipal",
                ),
            ],
        )

        response = client_gestor_carga_lot.post(
            f"{IMPORTACOES_URL}upload/",
            {
                "idpatriarca_id": patriarca.pk,
                "idtipocarga_id": tipo_organograma.pk,
                "arquivo": arquivo_planilha(conteudo=content),
            },
            format="multipart",
        )

        assert response.status_code == 201
        root = ItemOrganograma.objects.get(strsigla="PMVV")
        child = ItemOrganograma.objects.get(strsigla="SEMAD")
        assert root.strresponsavel == "MARIA DA SILVA"
        assert root.strcargo == "PREFEITA"
        assert child.strresponsavel == "JOÃO SOUZA"
        assert child.strcargo == "SECRETÁRIO MUNICIPAL"

    def test_rejeita_importacao_e_retorna_todos_os_erros(
        self,
        client_gestor_carga_lot,
        patriarca,
        tipo_organograma,
    ):
        conteudo = conteudo_organograma(
            linhas=[
                (1, "Prefeitura", "PMVV"),
                ("1.1", "Secretaria", "SEC REPETIDA"),
                ("1.2", "Secretaria", "SEC REPETIDA"),
                ("1.3.1", "", "SEM NOME"),
            ]
        )
        response = client_gestor_carga_lot.post(
            f"{IMPORTACOES_URL}upload/",
            {
                "idpatriarca_id": patriarca.pk,
                "idtipocarga_id": tipo_organograma.pk,
                "arquivo": arquivo_planilha(conteudo=conteudo),
            },
            format="multipart",
        )

        assert response.status_code == 400
        report = response.data["errors"]["arquivo"]
        assert report["leitura_completa"] is True
        assert report["total_erros"] >= 6
        assert {error["codigo"] for error in report["erros"]} >= {
            "nome_obrigatorio",
            "nome_duplicado_entre_irmaos",
            "sigla_duplicada_entre_irmaos",
            "sigla_com_espaco",
            "numero_pai_inexistente",
        }
        importacao = ImportacaoPlanilha.objects.get(
            pk=report["idimportacaoplanilha"]
        )
        assert importacao.strstatusimportacao == StatusImportacaoLocal.REJEITADA
        assert not importacao.itens_organograma.exists()
        persisted_report = importacao.relatorio_validacao
        assert persisted_report.strresultado == ResultadoRelatorioValidacao.INVALIDO
        assert persisted_report.inttotalerros == report["total_erros"]
        assert (
            OcorrenciaValidacaoImportacao.objects.filter(
                idrelatoriovalidacao=persisted_report
            ).count()
            == report["total_erros"]
        )
        process = ProcessoCargaInicial.objects.get(idpatriarca=patriarca)
        assert (
            process.strstatusorganograma
            == StatusOrganogramaProcesso.COM_PENDENCIAS
        )

    def test_exige_selecao_quando_organograma_tem_varias_abas(
        self,
        client_gestor_carga_lot,
        patriarca,
        tipo_organograma,
    ):
        conteudo = conteudo_organograma(
            segunda_aba=(
                ("Numero", "Nome", "Sigla"),
                [(1, "Outro patriarca", "OUTRO")],
            )
        )
        response = client_gestor_carga_lot.post(
            f"{IMPORTACOES_URL}upload/",
            {
                "idpatriarca_id": patriarca.pk,
                "idtipocarga_id": tipo_organograma.pk,
                "arquivo": arquivo_planilha(conteudo=conteudo),
            },
            format="multipart",
        )

        assert response.status_code == 409
        report = response.data["errors"]["arquivo"]
        assert report["codigo"] == "selecao_aba_obrigatoria"
        assert [aba["nome"] for aba in report["abas"]] == [
            "Organograma",
            "Segunda aba",
        ]
        assert not ImportacaoPlanilha.objects.exists()

    def test_importacao_rejeitada_nao_e_base_para_lotacao(
        self,
        client_gestor_carga_lot,
        patriarca,
        tipo_organograma,
        tipo_lotacao,
    ):
        rejected = ImportacaoPlanilha.objects.create(
            idpatriarca=patriarca,
            idtipocarga=tipo_organograma,
            idstatusprogresso=patriarca.idstatusprogresso,
            strnomearquivo="organograma-rejeitado.xlsx",
            strtipomime="application/octet-stream",
            inttamanhoarquivo=3,
            strhasharquivo=hashlib.sha256(b"abc").hexdigest(),
            strstatusimportacao=StatusImportacaoLocal.REJEITADA,
            binconteudo=b"abc",
        )

        response = client_gestor_carga_lot.post(
            f"{IMPORTACOES_URL}upload/",
            {
                "idpatriarca_id": patriarca.pk,
                "idtipocarga_id": tipo_lotacao.pk,
                "idimportacaoorganograma_id": rejected.pk,
                "arquivo": arquivo_planilha(nome="lotacao.xlsx"),
            },
            format="multipart",
        )

        assert response.status_code == 400
        assert "idimportacaoorganograma_id" in response.data["errors"]

    def test_importa_somente_aba_selecionada_e_persiste_hierarquia(
        self,
        client_gestor_carga_lot,
        patriarca,
        tipo_organograma,
    ):
        conteudo = conteudo_organograma(
            segunda_aba=(
                ("Numero", "Nome", "Sigla"),
                [(1, "Outro patriarca", "OUTRO")],
            )
        )
        response = client_gestor_carga_lot.post(
            f"{IMPORTACOES_URL}upload/",
            {
                "idpatriarca_id": patriarca.pk,
                "idtipocarga_id": tipo_organograma.pk,
                "aba": "Organograma",
                "arquivo": arquivo_planilha(conteudo=conteudo),
            },
            format="multipart",
        )

        assert response.status_code == 201
        assert response.data["strabaimportada"] == "Organograma"
        assert response.data["total_itens_organograma"] == 2
        items = ItemOrganograma.objects.order_by("intlinhaorigem")
        assert [item.strsigla for item in items] == ["PMVV", "SEMAD"]
        assert items[0].idpai_id is None
        assert items[1].idpai_id == items[0].pk

    def test_lotacao_exige_organograma_do_mesmo_patriarca(
        self,
        client_gestor_carga_lot,
        patriarca,
        tipo_organograma,
        tipo_lotacao,
    ):
        organograma = client_gestor_carga_lot.post(
            f"{IMPORTACOES_URL}upload/",
            {
                "idpatriarca_id": patriarca.pk,
                "idtipocarga_id": tipo_organograma.pk,
                "arquivo": arquivo_planilha(),
            },
            format="multipart",
        )
        assert organograma.status_code == 201
        sem_referencia = client_gestor_carga_lot.post(
            f"{IMPORTACOES_URL}upload/",
            {
                "idpatriarca_id": patriarca.pk,
                "idtipocarga_id": tipo_lotacao.pk,
                "arquivo": arquivo_planilha(nome="lotacao.xlsx"),
            },
            format="multipart",
        )
        assert sem_referencia.status_code == 400

        response = client_gestor_carga_lot.post(
            f"{IMPORTACOES_URL}upload/",
            {
                "idpatriarca_id": patriarca.pk,
                "idtipocarga_id": tipo_lotacao.pk,
                "idimportacaoorganograma_id": organograma.data[
                    "idimportacaoplanilha"
                ],
                "arquivo": arquivo_planilha(
                    nome="lotacao.xlsx",
                    conteudo=conteudo_lotacao(),
                ),
            },
            format="multipart",
        )
        assert response.status_code == 201
        assert response.data["idimportacaoorganograma_id"] == organograma.data[
            "idimportacaoplanilha"
        ]
        assert response.data["total_itens_lotacao"] == 1
        process = ProcessoCargaInicial.objects.get(idpatriarca=patriarca)
        assert (
            process.strfaseatual
            == FaseProcessoCargaInicial.PREPARANDO_LOTACOES
        )
        assert process.strstatusorganograma == StatusOrganogramaProcesso.PRONTO
        assert process.strstatuslotacao == StatusLotacaoProcesso.EM_PREPARACAO

    def test_lotacao_processa_todas_as_abas_e_persiste_itens_normalizados(
        self,
        client_gestor_carga_lot,
        patriarca,
        tipo_organograma,
        tipo_lotacao,
    ):
        organograma = client_gestor_carga_lot.post(
            f"{IMPORTACOES_URL}upload/",
            {
                "idpatriarca_id": patriarca.pk,
                "idtipocarga_id": tipo_organograma.pk,
                "arquivo": arquivo_planilha(),
            },
            format="multipart",
        )
        segunda_linha = (
            "SEMAD",
            "  secretaria de administracao  ",
            "João Souza",
            1234567890,
            "Analista Administrativo",
            "SEMAD",
            "secretaria de administracao",
        )

        response = client_gestor_carga_lot.post(
            f"{IMPORTACOES_URL}upload/",
            {
                "idpatriarca_id": patriarca.pk,
                "idtipocarga_id": tipo_lotacao.pk,
                "idimportacaoorganograma_id": organograma.data[
                    "idimportacaoplanilha"
                ],
                "arquivo": arquivo_planilha(
                    nome="lotacao.xlsx",
                    conteudo=conteudo_lotacao(segunda_aba=[segunda_linha]),
                ),
            },
            format="multipart",
        )

        assert response.status_code == 201
        assert response.data["total_itens_lotacao"] == 2
        report = response.data["relatorio_validacao"]
        assert report["strresultado"] == ResultadoRelatorioValidacao.VALIDO
        assert report["strversaoregras"] == "lotacao-v3"
        assert report["inttotalabas"] == 2
        assert report["inttotallinhas"] == 2
        assert report["inttotalvalidas"] == 2
        assert report["inttotalnormalizacoes"] >= 2

        items = ItemLotacao.objects.order_by("strabaorigem", "intlinhaorigem")
        assert list(items.values_list("strcpf", flat=True)) == [
            "90396100457",
            "01234567890",
        ]
        assert items[0].strcargooriginal == "auxiliar administrativo"
        assert items[0].strcargonormalizado == "AUXILIAR ADMINISTRATIVO"
        assert items[0].idorgaoorganograma_id == items[0].idunidadeorganograma_id

        listed = client_gestor_carga_lot.get(
            ITENS_LOTACAO_URL,
            {"importacao": response.data["idimportacaoplanilha"]},
        )
        assert listed.status_code == 200
        assert listed.data["count"] == 2

    def test_conjunto_lotacao_acrescenta_homologa_ativa_e_exporta(
        self,
        client_gestor_carga_lot,
        patriarca,
        tipo_organograma,
        tipo_lotacao,
    ):
        organograma = client_gestor_carga_lot.post(
            f"{IMPORTACOES_URL}upload/",
            {
                "idpatriarca_id": patriarca.pk,
                "idtipocarga_id": tipo_organograma.pk,
                "arquivo": arquivo_planilha(),
            },
            format="multipart",
        )
        first_upload = client_gestor_carga_lot.post(
            f"{IMPORTACOES_URL}upload/",
            {
                "idpatriarca_id": patriarca.pk,
                "idtipocarga_id": tipo_lotacao.pk,
                "idimportacaoorganograma_id": organograma.data[
                    "idimportacaoplanilha"
                ],
                "strreferencia": "Versão discutida com o RH",
                "arquivo": arquivo_planilha(
                    nome="lotacao-inicial.xlsx",
                    conteudo=conteudo_lotacao(),
                ),
            },
            format="multipart",
        )
        assert first_upload.status_code == 201
        set_id = first_upload.data["idconjuntolotacao_id"]
        conjunto = ConjuntoLotacao.objects.get(pk=set_id)
        assert conjunto.strsituacao == SituacaoConjuntoLotacao.PRONTO_PARA_HOMOLOGACAO
        assert first_upload.data["strtipooperacaolotacao"] == "INICIAL"

        second_row = (
            "SEMAD",
            "Secretaria de Administração",
            "João Souza",
            "01234567890",
            "Analista Administrativo",
            "SEMAD",
            "Secretaria de Administração",
        )
        addition = client_gestor_carga_lot.post(
            f"{CONJUNTOS_LOTACAO_URL}{set_id}/importacoes/",
            {
                "idtipocarga_id": tipo_lotacao.pk,
                "strreferencia": "Complemento recebido da secretaria",
                "arquivo": arquivo_planilha(
                    nome="lotacao-complementar.xlsx",
                    conteudo=conteudo_lotacao(linhas=[second_row]),
                ),
            },
            format="multipart",
        )
        assert addition.status_code == 201
        assert addition.data["idconjuntolotacao_id"] == set_id
        assert addition.data["strtipooperacaolotacao"] == "ACRESCIMO"
        assert conjunto.importacoes.count() == 2
        assert ItemLotacao.objects.filter(
            idimportacaoplanilha__idconjuntolotacao=conjunto
        ).count() == 2

        listed = client_gestor_carga_lot.get(
            CONJUNTOS_LOTACAO_URL,
            {
                "organograma": organograma.data["idimportacaoplanilha"],
            },
        )
        assert listed.status_code == 200
        assert listed.data["count"] == 1
        assert listed.data["results"][0]["total_importacoes"] == 2
        assert listed.data["results"][0]["total_registros"] == 2

        summary = client_gestor_carga_lot.get(
            f"{CONJUNTOS_LOTACAO_URL}resumo-organograma/",
            {"organograma": organograma.data["idimportacaoplanilha"]},
        )
        assert summary.status_code == 200
        assert summary.data["total_conjuntos"] == 1
        assert summary.data["total_importacoes"] == 2
        assert summary.data["conjunto_ativo"] is None

        homologated = client_gestor_carga_lot.post(
            f"{CONJUNTOS_LOTACAO_URL}{set_id}/homologar/",
            {"justificativa": "Conteúdo confirmado com o RH municipal."},
            format="json",
        )
        assert homologated.status_code == 200
        assert homologated.data["strsituacao"] == "HOMOLOGADO"

        activated = client_gestor_carga_lot.post(
            f"{CONJUNTOS_LOTACAO_URL}{set_id}/ativar/",
            {"justificativa": "Primeira versão homologada."},
            format="json",
        )
        assert activated.status_code == 200
        assert activated.data["strsituacao"] == "ATIVO"

        closed_addition = client_gestor_carga_lot.post(
            f"{CONJUNTOS_LOTACAO_URL}{set_id}/importacoes/",
            {
                "idtipocarga_id": tipo_lotacao.pk,
                "arquivo": arquivo_planilha(
                    nome="nao-pode-alterar-ativo.xlsx",
                    conteudo=conteudo_lotacao(linhas=[second_row]),
                ),
            },
            format="multipart",
        )
        assert closed_addition.status_code == 400

        protected_original = client_gestor_carga_lot.delete(
            (
                f"{IMPORTACOES_URL}"
                f"{first_upload.data['idimportacaoplanilha']}/"
            )
        )
        assert protected_original.status_code == 400
        assert ImportacaoPlanilha.objects.filter(
            pk=first_upload.data["idimportacaoplanilha"]
        ).exists()

        current_items = client_gestor_carga_lot.get(
            ITENS_LOTACAO_URL,
            {
                "conjunto_lotacao": set_id,
                "vigente": "true",
            },
        )
        assert current_items.status_code == 200
        assert current_items.data["count"] == 2

        history = client_gestor_carga_lot.get(
            f"{CONJUNTOS_LOTACAO_URL}{set_id}/historico/"
        )
        assert history.status_code == 200
        assert history.data["count"] >= 4
        assert {
            event["stracao"] for event in history.data["results"]
        } >= {
            "CONJUNTO_CRIADO",
            "PLANILHA_IMPORTADA",
            "CONJUNTO_HOMOLOGADO",
            "CONJUNTO_ATIVADO",
        }

        exported = client_gestor_carga_lot.get(
            f"{CONJUNTOS_LOTACAO_URL}{set_id}/download-consolidado/"
        )
        assert exported.status_code == 200
        workbook = load_workbook(BytesIO(exported.content))
        assert workbook.sheetnames == [
            "LOTACOES_CONSOLIDADAS",
            "REGISTROS_DESCARTADOS",
            "RESUMO",
            "HISTORICO",
        ]
        assert workbook["LOTACOES_CONSOLIDADAS"].max_row == 3

        original = client_gestor_carga_lot.get(first_upload.data["download_url"])
        assert original.status_code == 200
        assert original.content == ImportacaoPlanilha.objects.get(
            pk=first_upload.data["idimportacaoplanilha"]
        ).binconteudo

    def test_ativar_novo_conjunto_substitui_anterior_e_preserva_historico(
        self,
        client_gestor_carga_lot,
        patriarca,
        tipo_organograma,
        tipo_lotacao,
    ):
        organograma = client_gestor_carga_lot.post(
            f"{IMPORTACOES_URL}upload/",
            {
                "idpatriarca_id": patriarca.pk,
                "idtipocarga_id": tipo_organograma.pk,
                "arquivo": arquivo_planilha(),
            },
            format="multipart",
        )

        def create_ready_set(reference, filename, person_name):
            created = client_gestor_carga_lot.post(
                CONJUNTOS_LOTACAO_URL,
                {
                    "idpatriarca_id": patriarca.pk,
                    "idimportacaoorganograma_id": organograma.data[
                        "idimportacaoplanilha"
                    ],
                    "strreferencia": reference,
                },
                format="json",
            )
            assert created.status_code == 201
            row = (
                "SEMAD",
                "Secretaria de Administração",
                person_name,
                "90396100457",
                "Auxiliar Administrativo",
                "SEMAD",
                "Secretaria de Administração",
            )
            uploaded = client_gestor_carga_lot.post(
                (
                    f"{CONJUNTOS_LOTACAO_URL}"
                    f"{created.data['idconjuntolotacao']}/importacoes/"
                ),
                {
                    "idtipocarga_id": tipo_lotacao.pk,
                    "arquivo": arquivo_planilha(
                        nome=filename,
                        conteudo=conteudo_lotacao(linhas=[row]),
                    ),
                },
                format="multipart",
            )
            assert uploaded.status_code == 201
            expected_operation = (
                "SUBSTITUICAO"
                if ConjuntoLotacao.objects.filter(
                    idimportacaoorganograma_id=organograma.data[
                        "idimportacaoplanilha"
                    ],
                    strsituacao=SituacaoConjuntoLotacao.ATIVO,
                ).exclude(pk=created.data["idconjuntolotacao"]).exists()
                else "INICIAL"
            )
            assert uploaded.data["strtipooperacaolotacao"] == expected_operation
            homologated = client_gestor_carga_lot.post(
                (
                    f"{CONJUNTOS_LOTACAO_URL}"
                    f"{created.data['idconjuntolotacao']}/homologar/"
                ),
                {"justificativa": "Conferido com a prefeitura."},
                format="json",
            )
            assert homologated.status_code == 200
            return created.data["idconjuntolotacao"]

        first_id = create_ready_set("Alternativa inicial", "primeira.xlsx", "Maria")
        first_active = client_gestor_carga_lot.post(
            f"{CONJUNTOS_LOTACAO_URL}{first_id}/ativar/",
            {"justificativa": "Primeira versão vigente."},
            format="json",
        )
        assert first_active.status_code == 200

        replacement_id = create_ready_set(
            "Nova consolidação",
            "substituta.xlsx",
            "Joana",
        )
        replacement = client_gestor_carga_lot.post(
            f"{CONJUNTOS_LOTACAO_URL}{replacement_id}/ativar/",
            {"justificativa": "Nova consolidação aprovada pela prefeitura."},
            format="json",
        )
        assert replacement.status_code == 200
        assert replacement.data["strsituacao"] == "ATIVO"
        assert replacement.data["idconjuntoanterior_id"] == first_id

        first = ConjuntoLotacao.objects.get(pk=first_id)
        assert first.strsituacao == SituacaoConjuntoLotacao.SUBSTITUIDO
        assert first.importacoes.count() == 1
        assert first.historico.filter(stracao="CONJUNTO_SUBSTITUIDO").exists()
        assert ConjuntoLotacao.objects.filter(
            idimportacaoorganograma_id=organograma.data["idimportacaoplanilha"],
            strsituacao=SituacaoConjuntoLotacao.ATIVO,
        ).count() == 1

        first_detail = client_gestor_carga_lot.get(
            f"{CONJUNTOS_LOTACAO_URL}{first_id}/"
        )
        assert first_detail.status_code == 200
        assert first_detail.data["idconjuntosubstituto"] == replacement_id

        old_items = client_gestor_carga_lot.get(
            ITENS_LOTACAO_URL,
            {"conjunto_lotacao": first_id, "vigente": "false"},
        )
        assert old_items.status_code == 200
        assert old_items.data["count"] == 1

        cancelled = client_gestor_carga_lot.post(
            f"{CONJUNTOS_LOTACAO_URL}{replacement_id}/cancelar/",
            {"justificativa": "Prefeitura solicitou nova revisão."},
            format="json",
        )
        assert cancelled.status_code == 200
        assert cancelled.data["strsituacao"] == "CANCELADO"
        assert ConjuntoLotacao.objects.get(pk=replacement_id).importacoes.count() == 1

    def test_lotacao_rejeita_cpf_com_digitos_verificadores_invalidos(
        self,
        client_gestor_carga_lot,
        patriarca,
        tipo_organograma,
        tipo_lotacao,
    ):
        organograma = client_gestor_carga_lot.post(
            f"{IMPORTACOES_URL}upload/",
            {
                "idpatriarca_id": patriarca.pk,
                "idtipocarga_id": tipo_organograma.pk,
                "arquivo": arquivo_planilha(),
            },
            format="multipart",
        )
        linha = (
            "SEMAD",
            "Secretaria de Administração",
            "Maria da Silva",
            "123.456.789-01",
            "Auxiliar Administrativo",
            "SEMAD",
            "Secretaria de Administração",
        )

        response = client_gestor_carga_lot.post(
            f"{IMPORTACOES_URL}upload/",
            {
                "idpatriarca_id": patriarca.pk,
                "idtipocarga_id": tipo_lotacao.pk,
                "idimportacaoorganograma_id": organograma.data[
                    "idimportacaoplanilha"
                ],
                "arquivo": arquivo_planilha(
                    nome="lotacao-cpf-invalido.xlsx",
                    conteudo=conteudo_lotacao(linhas=[linha]),
                ),
            },
            format="multipart",
        )

        assert response.status_code == 201
        assert response.data["strstatusimportacao"] == "AGUARDANDO_REVISAO"
        assert {
            error["codigo"]
            for error in response.data["relatorio_validacao"]["jsonresumo"]["erros"]
        } == {"cpf_digitos_verificadores_invalidos"}
        staged = ItemLotacaoImportado.objects.get(
            idimportacaoplanilha_id=response.data["idimportacaoplanilha"]
        )
        assert staged.strsituacao == StatusItemLotacaoImportado.COM_ERROS
        assert not ItemLotacao.objects.filter(
            idimportacaoplanilha_id=response.data["idimportacaoplanilha"]
        ).exists()

    def test_lotacao_avisa_cpf_repetido_inclusive_entre_abas(
        self,
        client_gestor_carga_lot,
        patriarca,
        tipo_organograma,
        tipo_lotacao,
    ):
        organograma = client_gestor_carga_lot.post(
            f"{IMPORTACOES_URL}upload/",
            {
                "idpatriarca_id": patriarca.pk,
                "idtipocarga_id": tipo_organograma.pk,
                "arquivo": arquivo_planilha(),
            },
            format="multipart",
        )
        repetida = (
            "SEMAD",
            "Secretaria de Administração",
            "Maria da Silva",
            "90396100457",
            "Outro cargo",
            "SEMAD",
            "Secretaria de Administração",
        )

        response = client_gestor_carga_lot.post(
            f"{IMPORTACOES_URL}upload/",
            {
                "idpatriarca_id": patriarca.pk,
                "idtipocarga_id": tipo_lotacao.pk,
                "idimportacaoorganograma_id": organograma.data[
                    "idimportacaoplanilha"
                ],
                "arquivo": arquivo_planilha(
                    nome="lotacao-cpf-repetido.xlsx",
                    conteudo=conteudo_lotacao(segunda_aba=[repetida]),
                ),
            },
            format="multipart",
        )

        assert response.status_code == 201
        report = response.data["relatorio_validacao"]
        assert report["inttotalavisos"] == 2
        assert {
            warning["strcodigo"] for warning in report["ocorrencias"]
            if warning["strseveridade"] == "AVISO"
        } == {"cpf_duplicado_planilha"}
        assert response.data["strstatusimportacao"] == "AGUARDANDO_REVISAO"
        assert response.data["total_itens_lotacao"] == 0
        assert response.data["total_itens_lotacao_importados"] == 2
        conjunto = ConjuntoLotacao.objects.get(
            pk=response.data["idconjuntolotacao_id"]
        )
        assert conjunto.strsituacao == SituacaoConjuntoLotacao.COM_PENDENCIAS
        cases = report["casos"]
        assert len(cases) == 1
        assert len(cases[0]["item_ids"]) == 2

        resolution = client_gestor_carga_lot.post(
            f"{CASOS_VALIDACAO_URL}{cases[0]['idcasovalidacao']}/resolver/",
            {
                "acao": "ACEITAR_COMO_INFORMADO",
                "justificativa": "O servidor possui duas lotações válidas.",
            },
            format="json",
        )
        assert resolution.status_code == 200
        assert resolution.data["strstatus"] == "ACEITO_COMO_INFORMADO"
        importacao = ImportacaoPlanilha.objects.get(
            pk=response.data["idimportacaoplanilha"]
        )
        assert (
            importacao.strstatusimportacao
            == StatusImportacaoLocal.PRONTA_PARA_INCORPORAR
        )
        conjunto.refresh_from_db()
        assert conjunto.strsituacao == SituacaoConjuntoLotacao.EM_VALIDACAO

        inherited_item = ItemLotacaoImportado.objects.filter(
            idimportacaoplanilha=importacao
        ).first()
        inherited_item.strcpf = "903.961.004-57"
        inherited_item.save(update_fields=["strcpf", "updated_at"])

        incorporated = client_gestor_carga_lot.post(
            f"{IMPORTACOES_URL}{importacao.pk}/incorporar-lotacao/",
            {},
            format="json",
        )
        assert incorporated.status_code == 200
        assert incorporated.data["strstatusimportacao"] == "INCORPORADA"
        assert incorporated.data["total_itens_lotacao"] == 2
        inherited_item.refresh_from_db()
        assert inherited_item.strcpf == "90396100457"
        conjunto.refresh_from_db()
        assert (
            conjunto.strsituacao
            == SituacaoConjuntoLotacao.PRONTO_PARA_HOMOLOGACAO
        )
        assert set(
            ItemLotacaoImportado.objects.filter(
                idimportacaoplanilha=importacao
            ).values_list("strsituacao", flat=True)
        ) == {StatusItemLotacaoImportado.INCORPORADO}

    def test_correcao_manual_normaliza_cpf_formatado_antes_de_incorporar(
        self,
        client_gestor_carga_lot,
        patriarca,
        tipo_organograma,
        tipo_lotacao,
    ):
        organograma = client_gestor_carga_lot.post(
            f"{IMPORTACOES_URL}upload/",
            {
                "idpatriarca_id": patriarca.pk,
                "idtipocarga_id": tipo_organograma.pk,
                "arquivo": arquivo_planilha(),
            },
            format="multipart",
        )
        repetida = (
            "SEMAD",
            "Secretaria de Administração",
            "Maria da Silva",
            "90396100457",
            "Outro cargo",
            "SEMAD",
            "Secretaria de Administração",
        )
        uploaded = client_gestor_carga_lot.post(
            f"{IMPORTACOES_URL}upload/",
            {
                "idpatriarca_id": patriarca.pk,
                "idtipocarga_id": tipo_lotacao.pk,
                "idimportacaoorganograma_id": organograma.data[
                    "idimportacaoplanilha"
                ],
                "arquivo": arquivo_planilha(
                    nome="lotacao-correcao-cpf.xlsx",
                    conteudo=conteudo_lotacao(segunda_aba=[repetida]),
                ),
            },
            format="multipart",
        )
        case = uploaded.data["relatorio_validacao"]["casos"][0]
        item_ids = case["item_ids"]

        corrected = client_gestor_carga_lot.post(
            f"{CASOS_VALIDACAO_URL}{case['idcasovalidacao']}/resolver/",
            {
                "acao": "CORRIGIR_MANUALMENTE",
                "alteracoes": [
                    {
                        "iditemlotacaoimportado": item_ids[0],
                        "campos": {"strcpf": "980.545.250-60"},
                    },
                    {
                        "iditemlotacaoimportado": item_ids[1],
                        "campos": {"strcpf": "359.644.720-87"},
                    },
                ],
            },
            format="json",
        )

        assert corrected.status_code == 200
        importacao_id = uploaded.data["idimportacaoplanilha"]
        assert set(
            ItemLotacaoImportado.objects.filter(
                idimportacaoplanilha_id=importacao_id
            ).values_list("strcpf", flat=True)
        ) == {"98054525060", "35964472087"}
        resolution = ResolucaoValidacao.objects.get(
            idcasovalidacao_id=case["idcasovalidacao"]
        )
        assert {
            item["strcpf"] for item in resolution.jsondepois
        } == {"98054525060", "35964472087"}

        incorporated = client_gestor_carga_lot.post(
            f"{IMPORTACOES_URL}{importacao_id}/incorporar-lotacao/",
            {},
            format="json",
        )

        assert incorporated.status_code == 200
        assert set(
            ItemLotacao.objects.filter(
                idimportacaoplanilha_id=importacao_id
            ).values_list("strcpf", flat=True)
        ) == {"98054525060", "35964472087"}

    def test_incorporacao_converte_erro_de_modelo_em_resposta_400(
        self,
        client_gestor_carga_lot,
        patriarca,
        tipo_organograma,
        tipo_lotacao,
        monkeypatch,
    ):
        organograma = client_gestor_carga_lot.post(
            f"{IMPORTACOES_URL}upload/",
            {
                "idpatriarca_id": patriarca.pk,
                "idtipocarga_id": tipo_organograma.pk,
                "arquivo": arquivo_planilha(),
            },
            format="multipart",
        )
        repetida = (
            "SEMAD",
            "Secretaria de Administração",
            "Maria da Silva",
            "90396100457",
            "Outro cargo",
            "SEMAD",
            "Secretaria de Administração",
        )
        uploaded = client_gestor_carga_lot.post(
            f"{IMPORTACOES_URL}upload/",
            {
                "idpatriarca_id": patriarca.pk,
                "idtipocarga_id": tipo_lotacao.pk,
                "idimportacaoorganograma_id": organograma.data[
                    "idimportacaoplanilha"
                ],
                "arquivo": arquivo_planilha(
                    nome="lotacao-erro-modelo.xlsx",
                    conteudo=conteudo_lotacao(segunda_aba=[repetida]),
                ),
            },
            format="multipart",
        )
        case = uploaded.data["relatorio_validacao"]["casos"][0]
        accepted = client_gestor_carga_lot.post(
            f"{CASOS_VALIDACAO_URL}{case['idcasovalidacao']}/resolver/",
            {
                "acao": "ACEITAR_COMO_INFORMADO",
                "justificativa": "As duas lotações são intencionais.",
            },
            format="json",
        )
        assert accepted.status_code == 200
        importacao_id = uploaded.data["idimportacaoplanilha"]

        def invalid_model(_instance):
            raise DjangoValidationError(
                {"strnomeservidor": ["Valor rejeitado pelo modelo."]}
            )

        monkeypatch.setattr(ItemLotacao, "full_clean", invalid_model)

        response = client_gestor_carga_lot.post(
            f"{IMPORTACOES_URL}{importacao_id}/incorporar-lotacao/",
            {},
            format="json",
        )

        assert response.status_code == 400
        assert not ItemLotacao.objects.filter(
            idimportacaoplanilha_id=importacao_id
        ).exists()
        importacao = ImportacaoPlanilha.objects.get(pk=importacao_id)
        assert (
            importacao.strstatusimportacao
            == StatusImportacaoLocal.PRONTA_PARA_INCORPORAR
        )

    def test_lotacao_aceita_cpf_repetido_em_orgao_ou_unidade_diferente(
        self,
        client_gestor_carga_lot,
        patriarca,
        tipo_organograma,
        tipo_lotacao,
    ):
        organograma = client_gestor_carga_lot.post(
            f"{IMPORTACOES_URL}upload/",
            {
                "idpatriarca_id": patriarca.pk,
                "idtipocarga_id": tipo_organograma.pk,
                "arquivo": arquivo_planilha(
                    conteudo=conteudo_organograma(
                        linhas=[
                            (1, "Prefeitura Municipal de Vila Velha", "PMVV"),
                            ("1.1", "Secretaria de Administração", "SEMAD"),
                            ("1.1.1", "Secretaria de Administração", "SEMAD"),
                            ("1.1.2", "Recursos Humanos", "RH"),
                            ("1.2", "Secretaria de Saúde", "SESA"),
                            ("1.2.1", "Secretaria de Saúde", "SESA"),
                        ]
                    )
                ),
            },
            format="multipart",
        )
        cpf = "90396100457"
        linhas = [
            (
                "SEMAD",
                "Administração",
                "Maria da Silva",
                cpf,
                "Auxiliar Administrativo",
                "SEMAD",
                "Administração",
            ),
            (
                "SEMAD",
                "Administração",
                "Maria da Silva",
                cpf,
                "Auxiliar Administrativo",
                "RH",
                "RH",
            ),
            (
                "SESA",
                "Saúde",
                "Maria da Silva",
                cpf,
                "Auxiliar Administrativo",
                "SESA",
                "Saúde",
            ),
        ]

        response = client_gestor_carga_lot.post(
            f"{IMPORTACOES_URL}upload/",
            {
                "idpatriarca_id": patriarca.pk,
                "idtipocarga_id": tipo_lotacao.pk,
                "idimportacaoorganograma_id": organograma.data[
                    "idimportacaoplanilha"
                ],
                "arquivo": arquivo_planilha(
                    nome="lotacao-cpf-em-lotacoes-distintas.xlsx",
                    conteudo=conteudo_lotacao(linhas=linhas),
                ),
            },
            format="multipart",
        )

        assert response.status_code == 201
        assert response.data["strstatusimportacao"] == "INCORPORADA"
        assert response.data["total_itens_lotacao"] == 3
        report = response.data["relatorio_validacao"]
        assert report["inttotalavisos"] == 0
        assert not any(
            occurrence["strcodigo"] == "cpf_duplicado_planilha"
            for occurrence in report["ocorrencias"]
        )

    def test_lotacao_rejeita_apenas_nome_servidor_ausente_quando_siglas_validas(
        self,
        client_gestor_carga_lot,
        patriarca,
        tipo_organograma,
        tipo_lotacao,
    ):
        organograma = client_gestor_carga_lot.post(
            f"{IMPORTACOES_URL}upload/",
            {
                "idpatriarca_id": patriarca.pk,
                "idtipocarga_id": tipo_organograma.pk,
                "arquivo": arquivo_planilha(),
            },
            format="multipart",
        )
        linha = (
            "SEMAD",
            "Nome divergente do órgão",
            "",
            "90396100457",
            "Auxiliar Administrativo",
            "SEMAD",
            "Nome divergente da unidade",
        )

        response = client_gestor_carga_lot.post(
            f"{IMPORTACOES_URL}upload/",
            {
                "idpatriarca_id": patriarca.pk,
                "idtipocarga_id": tipo_lotacao.pk,
                "idimportacaoorganograma_id": organograma.data[
                    "idimportacaoplanilha"
                ],
                "arquivo": arquivo_planilha(
                    nome="lotacao-nomes-invalidos.xlsx",
                    conteudo=conteudo_lotacao(linhas=[linha]),
                ),
            },
            format="multipart",
        )

        assert response.status_code == 201
        codes = {
            error["codigo"]
            for error in response.data["relatorio_validacao"]["jsonresumo"]["erros"]
        }
        assert codes == {"nome_servidor_ausente"}
        assert response.data["strstatusimportacao"] == "AGUARDANDO_REVISAO"
        assert response.data["total_casos_pendentes"] == 1

    def test_lotacao_permite_corrigir_nome_servidor_e_incorporar(
        self,
        client_gestor_carga_lot,
        patriarca,
        tipo_organograma,
        tipo_lotacao,
    ):
        organograma = client_gestor_carga_lot.post(
            f"{IMPORTACOES_URL}upload/",
            {
                "idpatriarca_id": patriarca.pk,
                "idtipocarga_id": tipo_organograma.pk,
                "arquivo": arquivo_planilha(),
            },
            format="multipart",
        )
        linha = (
            "SEMAD",
            "Nome divergente do órgão",
            "",
            "90396100457",
            "Auxiliar Administrativo",
            "SEMAD",
            "Nome divergente da unidade",
        )
        uploaded = client_gestor_carga_lot.post(
            f"{IMPORTACOES_URL}upload/",
            {
                "idpatriarca_id": patriarca.pk,
                "idtipocarga_id": tipo_lotacao.pk,
                "idimportacaoorganograma_id": organograma.data[
                    "idimportacaoplanilha"
                ],
                "arquivo": arquivo_planilha(
                    nome="lotacao-para-revisao.xlsx",
                    conteudo=conteudo_lotacao(linhas=[linha]),
                ),
            },
            format="multipart",
        )
        assert uploaded.status_code == 201
        cases = {
            case["strcodigo"]: case
            for case in uploaded.data["relatorio_validacao"]["casos"]
        }
        item_id = cases["nome_servidor_ausente"]["item_ids"][0]

        refused = client_gestor_carga_lot.post(
            (
                f"{CASOS_VALIDACAO_URL}"
                f"{cases['nome_servidor_ausente']['idcasovalidacao']}/resolver/"
            ),
            {
                "acao": "ACEITAR_COMO_INFORMADO",
                "justificativa": "Não corrigir.",
            },
            format="json",
        )
        assert refused.status_code == 400

        corrected = client_gestor_carga_lot.post(
            (
                f"{CASOS_VALIDACAO_URL}"
                f"{cases['nome_servidor_ausente']['idcasovalidacao']}/resolver/"
            ),
            {
                "acao": "CORRIGIR_MANUALMENTE",
                "alteracoes": [
                    {
                        "iditemlotacaoimportado": item_id,
                        "campos": {"strnomeservidor": "Maria da Silva"},
                    }
                ],
            },
            format="json",
        )
        assert corrected.status_code == 200
        assert corrected.data["strstatus"] == "CORRIGIDO"

        importacao_id = uploaded.data["idimportacaoplanilha"]
        importacao = ImportacaoPlanilha.objects.get(pk=importacao_id)
        assert (
            importacao.strstatusimportacao
            == StatusImportacaoLocal.PRONTA_PARA_INCORPORAR
        )
        incorporated = client_gestor_carga_lot.post(
            f"{IMPORTACOES_URL}{importacao_id}/incorporar-lotacao/",
            {},
            format="json",
        )
        assert incorporated.status_code == 200
        item = ItemLotacao.objects.get(idimportacaoplanilha_id=importacao_id)
        assert item.strnomeservidor == "Maria da Silva"
        assert item.strnomeorgao == "Nome divergente do órgão"
        assert item.strnomeunidade == "Nome divergente da unidade"

    def test_descartar_linha_resolve_suas_criticas_sem_registros_validos(
        self,
        client_gestor_carga_lot,
        patriarca,
        tipo_organograma,
        tipo_lotacao,
    ):
        organograma = client_gestor_carga_lot.post(
            f"{IMPORTACOES_URL}upload/",
            {
                "idpatriarca_id": patriarca.pk,
                "idtipocarga_id": tipo_organograma.pk,
                "arquivo": arquivo_planilha(),
            },
            format="multipart",
        )
        linha = (
            "SEMAD",
            "Nome divergente",
            "",
            "123",
            "",
            "INEXISTENTE",
            "Unidade inexistente",
        )
        uploaded = client_gestor_carga_lot.post(
            f"{IMPORTACOES_URL}upload/",
            {
                "idpatriarca_id": patriarca.pk,
                "idtipocarga_id": tipo_lotacao.pk,
                "idimportacaoorganograma_id": organograma.data[
                    "idimportacaoplanilha"
                ],
                "arquivo": arquivo_planilha(
                    nome="lotacao-descartada.xlsx",
                    conteudo=conteudo_lotacao(linhas=[linha]),
                ),
            },
            format="multipart",
        )
        cases = uploaded.data["relatorio_validacao"]["casos"]
        discarded = client_gestor_carga_lot.post(
            f"{CASOS_VALIDACAO_URL}{cases[0]['idcasovalidacao']}/resolver/",
            {
                "acao": "DESCARTAR_LINHAS",
                "justificativa": "Registro será corrigido na planilha de origem.",
            },
            format="json",
        )
        assert discarded.status_code == 200

        importacao = ImportacaoPlanilha.objects.get(
            pk=uploaded.data["idimportacaoplanilha"]
        )
        assert (
            importacao.strstatusimportacao
            == StatusImportacaoLocal.SEM_REGISTROS_VALIDOS
        )
        assert not importacao.relatorio_validacao.casos.filter(
            strstatus="PENDENTE"
        ).exists()
        assert set(
            importacao.relatorio_validacao.casos.values_list(
                "strstatus", flat=True
            )
        ) == {"LINHAS_DESCARTADAS"}
        assert importacao.itens_lotacao_importados.get().strsituacao == "DESCARTADO"

        incorporated = client_gestor_carga_lot.post(
            f"{IMPORTACOES_URL}{importacao.pk}/incorporar-lotacao/",
            {},
            format="json",
        )
        assert incorporated.status_code == 400
        assert not importacao.itens_lotacao.exists()

    def test_lotacao_invalida_le_tudo_rejeita_e_nao_persiste_itens(
        self,
        client_gestor_carga_lot,
        patriarca,
        tipo_organograma,
        tipo_lotacao,
    ):
        organograma = client_gestor_carga_lot.post(
            f"{IMPORTACOES_URL}upload/",
            {
                "idpatriarca_id": patriarca.pk,
                "idtipocarga_id": tipo_organograma.pk,
                "arquivo": arquivo_planilha(),
            },
            format="multipart",
        )
        invalid_rows = [
            (
                "SEMAD",
                "Secretaria de Administração",
                "Pessoa Um",
                "123",
                "",
                "INEXISTENTE",
                "Unidade inexistente",
            ),
            (
                "OUTRO",
                "Outro órgão",
                "Pessoa Dois",
                "987.654.321-00",
                "Analista",
                "SEMAD",
                "Secretaria de Administração",
            ),
        ]

        response = client_gestor_carga_lot.post(
            f"{IMPORTACOES_URL}upload/",
            {
                "idpatriarca_id": patriarca.pk,
                "idtipocarga_id": tipo_lotacao.pk,
                "idimportacaoorganograma_id": organograma.data[
                    "idimportacaoplanilha"
                ],
                "arquivo": arquivo_planilha(
                    nome="lotacao-invalida.xlsx",
                    conteudo=conteudo_lotacao(linhas=invalid_rows),
                ),
            },
            format="multipart",
        )

        assert response.status_code == 201
        report = response.data["relatorio_validacao"]["jsonresumo"]
        assert report["leitura_completa"] is True
        assert report["total_linhas_lidas"] == 2
        assert {error["codigo"] for error in report["erros"]} >= {
            "cpf_invalido",
            "cargo_ausente",
            "unidade_invalida",
            "orgao_invalido",
        }
        importacao = ImportacaoPlanilha.objects.get(
            pk=response.data["idimportacaoplanilha"]
        )
        assert (
            importacao.strstatusimportacao
            == StatusImportacaoLocal.AGUARDANDO_REVISAO
        )
        assert not importacao.itens_lotacao.exists()
        assert importacao.itens_lotacao_importados.count() == 2
        assert CasoValidacaoImportacao.objects.filter(
            idrelatoriovalidacao=importacao.relatorio_validacao,
            strstatus="PENDENTE",
        ).exists()
        process = ProcessoCargaInicial.objects.get(idpatriarca=patriarca)
        assert process.strstatuslotacao == StatusLotacaoProcesso.COM_PENDENCIAS

    def test_novo_organograma_desatualiza_lotacoes_existentes(
        self,
        client_gestor_carga_lot,
        patriarca,
        tipo_organograma,
        tipo_lotacao,
    ):
        first_organogram = client_gestor_carga_lot.post(
            f"{IMPORTACOES_URL}upload/",
            {
                "idpatriarca_id": patriarca.pk,
                "idtipocarga_id": tipo_organograma.pk,
                "arquivo": arquivo_planilha(nome="organograma-1.xlsx"),
            },
            format="multipart",
        )
        assert first_organogram.status_code == 201

        lotation = client_gestor_carga_lot.post(
            f"{IMPORTACOES_URL}upload/",
            {
                "idpatriarca_id": patriarca.pk,
                "idtipocarga_id": tipo_lotacao.pk,
                "idimportacaoorganograma_id": first_organogram.data[
                    "idimportacaoplanilha"
                ],
                "arquivo": arquivo_planilha(
                    nome="lotacao.xlsx",
                    conteudo=conteudo_lotacao(),
                ),
            },
            format="multipart",
        )
        assert lotation.status_code == 201

        second_organogram = client_gestor_carga_lot.post(
            f"{IMPORTACOES_URL}upload/",
            {
                "idpatriarca_id": patriarca.pk,
                "idtipocarga_id": tipo_organograma.pk,
                "arquivo": arquivo_planilha(nome="organograma-2.xlsx"),
            },
            format="multipart",
        )

        assert second_organogram.status_code == 201
        process = ProcessoCargaInicial.objects.get(idpatriarca=patriarca)
        assert (
            process.strfaseatual
            == FaseProcessoCargaInicial.PREPARANDO_ORGANOGRAMA
        )
        assert process.strstatuslotacao == StatusLotacaoProcesso.DESATUALIZADA
        assert list(process.historico.values_list("stracao", flat=True)) == [
            "PROCESSO_CRIADO",
            "ORGANOGRAMA_IMPORTADO",
            "LOTACAO_IMPORTADA",
            "ORGANOGRAMA_IMPORTADO",
        ]

    def test_interrompe_quando_arquivo_nao_pode_ser_lido(
        self,
        client_gestor_carga_lot,
        patriarca,
        tipo_organograma,
    ):
        response = client_gestor_carga_lot.post(
            f"{IMPORTACOES_URL}upload/",
            {
                "idpatriarca_id": patriarca.pk,
                "idtipocarga_id": tipo_organograma.pk,
                "arquivo": arquivo_planilha(conteudo=b"arquivo corrompido"),
            },
            format="multipart",
        )

        assert response.status_code == 400
        report = response.data["errors"]["arquivo"]
        assert report["leitura_completa"] is False
        assert report["erros"][0]["codigo"] == "arquivo_ilegivel"
        importacao = ImportacaoPlanilha.objects.get(
            pk=report["idimportacaoplanilha"]
        )
        assert importacao.strstatusimportacao == StatusImportacaoLocal.REJEITADA
        assert importacao.relatorio_validacao.flgleituracompleta is False

    def test_download_do_ultimo_organograma_pelo_patriarca(
        self,
        client_gestor_carga_lot,
        patriarca,
        tipo_organograma,
    ):
        antigo = b"organograma antigo"
        recente = b"organograma mais recente"
        for nome, conteudo in [
            ("antigo.xlsx", antigo),
            ("recente.xlsx", recente),
        ]:
            ImportacaoPlanilha.objects.create(
                idpatriarca=patriarca,
                idtipocarga=tipo_organograma,
                idstatusprogresso=patriarca.idstatusprogresso,
                strnomearquivo=nome,
                strtipomime=(
                    "application/vnd.openxmlformats-officedocument."
                    "spreadsheetml.sheet"
                ),
                inttamanhoarquivo=len(conteudo),
                strhasharquivo=hashlib.sha256(conteudo).hexdigest(),
                binconteudo=conteudo,
            )

        detail_response = client_gestor_carga_lot.get(
            f"{PATRIARCAS_URL}{patriarca.pk}/"
        )
        assert detail_response.status_code == 200
        assert detail_response.data["organograma_download_url"].endswith(
            f"{PATRIARCAS_URL}{patriarca.pk}/organograma/download/"
        )

        response = client_gestor_carga_lot.get(
            f"{PATRIARCAS_URL}{patriarca.pk}/organograma/download/"
        )
        assert response.status_code == 200
        assert response.content == recente
        assert "recente.xlsx" in response["Content-Disposition"]

    def test_download_de_organograma_inexistente_retorna_404(
        self,
        client_gestor_carga_lot,
        patriarca,
    ):
        response = client_gestor_carga_lot.get(
            f"{PATRIARCAS_URL}{patriarca.pk}/organograma/download/"
        )

        assert response.status_code == 404
        assert response.data["success"] is False

    def test_atualiza_status_e_exclui_importacao(
        self,
        client_gestor_carga_lot,
        patriarca,
        tipo_organograma,
        status_processando,
    ):
        importacao = ImportacaoPlanilha.objects.create(
            idpatriarca=patriarca,
            idtipocarga=tipo_organograma,
            idstatusprogresso=patriarca.idstatusprogresso,
            strnomearquivo="organograma.xlsx",
            strtipomime="application/octet-stream",
            inttamanhoarquivo=3,
            strhasharquivo=hashlib.sha256(b"abc").hexdigest(),
            binconteudo=b"abc",
        )

        patch_response = client_gestor_carga_lot.patch(
            f"{IMPORTACOES_URL}{importacao.pk}/",
            {"idstatusprogresso_id": status_processando.pk},
            format="json",
        )
        assert patch_response.status_code == 200
        assert patch_response.data["status"]["strdescricao"] == "Em progresso"

        delete_response = client_gestor_carga_lot.delete(
            f"{IMPORTACOES_URL}{importacao.pk}/"
        )
        assert delete_response.status_code == 204

    def test_referencia_e_opcional_e_pode_ser_atualizada(
        self,
        client_gestor_carga_lot,
        patriarca,
        tipo_organograma,
    ):
        upload_response = client_gestor_carga_lot.post(
            f"{IMPORTACOES_URL}upload/",
            {
                "idpatriarca_id": patriarca.pk,
                "idtipocarga_id": tipo_organograma.pk,
                "arquivo": arquivo_planilha(),
            },
            format="multipart",
        )
        assert upload_response.status_code == 201
        assert upload_response.data["strreferencia"] == ""

        importacao_id = upload_response.data["idimportacaoplanilha"]
        patch_response = client_gestor_carga_lot.patch(
            f"{IMPORTACOES_URL}{importacao_id}/",
            {"strreferencia": "Planilha revisada conforme reunião de 15/10."},
            format="json",
        )

        assert patch_response.status_code == 200
        assert patch_response.data["strreferencia"] == (
            "Planilha revisada conforme reunião de 15/10."
        )

    def test_exclusao_da_importacao_remove_fisicamente_os_itens(
        self,
        client_gestor_carga_lot,
        patriarca,
        tipo_organograma,
    ):
        response = client_gestor_carga_lot.post(
            f"{IMPORTACOES_URL}upload/",
            {
                "idpatriarca_id": patriarca.pk,
                "idtipocarga_id": tipo_organograma.pk,
                "arquivo": arquivo_planilha(),
            },
            format="multipart",
        )
        importacao_id = response.data["idimportacaoplanilha"]
        item_ids = list(
            ItemOrganograma.objects.filter(
                idimportacaoplanilha_id=importacao_id
            ).values_list("pk", flat=True)
        )

        delete_response = client_gestor_carga_lot.delete(
            f"{IMPORTACOES_URL}{importacao_id}/"
        )

        assert delete_response.status_code == 204
        assert not ImportacaoPlanilha.objects.filter(pk=importacao_id).exists()
        assert not ItemOrganograma.objects.filter(pk__in=item_ids).exists()


@pytest.mark.django_db(transaction=True)
class TestItemOrganograma:
    def test_crud_completo(
        self,
        client_gestor_carga_lot,
        patriarca,
        tipo_organograma,
    ):
        upload = client_gestor_carga_lot.post(
            f"{IMPORTACOES_URL}upload/",
            {
                "idpatriarca_id": patriarca.pk,
                "idtipocarga_id": tipo_organograma.pk,
                "arquivo": arquivo_planilha(),
            },
            format="multipart",
        )
        importacao_id = upload.data["idimportacaoplanilha"]
        root = ItemOrganograma.objects.get(
            idimportacaoplanilha_id=importacao_id,
            idpai__isnull=True,
        )

        create_response = client_gestor_carga_lot.post(
            ITENS_ORGANOGRAMA_URL,
            {
                "idimportacaoplanilha_id": importacao_id,
                "idpai_id": root.pk,
                "strnome": "Gerência de Pessoas",
                "strsigla": "gep",
                "strnumero": "1.2",
                "strresponsavel": "Maria de Souza",
                "strcargo": "Gerente",
            },
            format="json",
        )
        assert create_response.status_code == 201
        pk = create_response.data["iditemorganograma"]
        assert create_response.data["strsigla"] == "GEP"
        assert create_response.data["strresponsavel"] == "Maria de Souza"
        assert create_response.data["strcargo"] == "Gerente"
        assert create_response.data["strreferenciapai"] == root.strsigla
        assert create_response.data["strabaorigem"] == "Edição manual"
        assert create_response.data["intlinhaorigem"] == 0
        importacao = ImportacaoPlanilha.objects.get(pk=importacao_id)
        assert importacao.strstatusimportacao == StatusImportacaoLocal.ALTERADA
        assert importacao.flgpossuialteracoes is True
        assert importacao.datprimeiraalteracao is not None
        assert importacao.datultimaalteracao is not None
        process = importacao.idprocessocargainicial
        process.refresh_from_db()
        assert (
            process.strfaseatual
            == FaseProcessoCargaInicial.PREPARANDO_ORGANOGRAMA
        )
        assert (
            process.strstatusorganograma
            == StatusOrganogramaProcesso.EM_PREPARACAO
        )

        list_response = client_gestor_carga_lot.get(
            ITENS_ORGANOGRAMA_URL,
            {"importacao": importacao_id},
        )
        assert list_response.status_code == 200
        assert list_response.data["count"] == 3
        assert client_gestor_carga_lot.get(
            f"{ITENS_ORGANOGRAMA_URL}{pk}/"
        ).status_code == 200

        patch_response = client_gestor_carga_lot.patch(
            f"{ITENS_ORGANOGRAMA_URL}{pk}/",
            {
                "strnome": "Gerência de Gestão de Pessoas",
                "strresponsavel": "Joana de Souza",
                "strcargo": "Gerente de Gestão",
            },
            format="json",
        )
        assert patch_response.status_code == 200
        assert patch_response.data["strnome"] == "Gerência de Gestão de Pessoas"
        assert patch_response.data["strresponsavel"] == "Joana de Souza"
        assert patch_response.data["strcargo"] == "Gerente de Gestão"

        delete_response = client_gestor_carga_lot.delete(
            f"{ITENS_ORGANOGRAMA_URL}{pk}/"
        )
        assert delete_response.status_code == 204
        assert ImportacaoPlanilha.objects.filter(pk=importacao_id).exists()
        assert list(
            process.historico.values_list("stracao", flat=True)
        )[-3:] == [
            "ITEM_ORGANOGRAMA_INCLUIDO",
            "ITEM_ORGANOGRAMA_ALTERADO",
            "ITEM_ORGANOGRAMA_EXCLUIDO",
        ]

        import_response = client_gestor_carga_lot.get(
            f"{IMPORTACOES_URL}{importacao_id}/"
        )
        updated_url = import_response.data["organograma_atualizado_download_url"]
        assert updated_url.endswith(
            f"{IMPORTACOES_URL}{importacao_id}/organograma-atualizado/download/"
        )
        export_response = client_gestor_carga_lot.get(updated_url)
        assert export_response.status_code == 200
        workbook = load_workbook(BytesIO(export_response.content), data_only=True)
        worksheet = workbook["Organograma"]
        assert list(worksheet.values) == [
            ("Numero", "Nome", "Sigla", "Pai_Sigla", "Responsavel", "Cargo"),
            (
                "1",
                "PREFEITURA MUNICIPAL DE VILA VELHA",
                "PMVV",
                None,
                None,
                None,
            ),
            (
                "1.1",
                "SECRETARIA DE ADMINISTRAÇÃO",
                "SEMAD",
                "PMVV",
                None,
                None,
            ),
        ]

    @pytest.mark.parametrize(
        ("parent_type", "field", "value"),
        [
            ("organ", "strnome", "N" * 101),
            ("organ", "strsigla", "A" * 21),
            ("unit", "strnome", "N" * 101),
            ("unit", "strsigla", "A" * 26),
        ],
    )
    def test_crud_aplica_tamanhos_maximos_da_prodest(
        self,
        client_gestor_carga_lot,
        patriarca,
        tipo_organograma,
        parent_type,
        field,
        value,
    ):
        upload = client_gestor_carga_lot.post(
            f"{IMPORTACOES_URL}upload/",
            {
                "idpatriarca_id": patriarca.pk,
                "idtipocarga_id": tipo_organograma.pk,
                "arquivo": arquivo_planilha(),
            },
            format="multipart",
        )
        importacao_id = upload.data["idimportacaoplanilha"]
        root = ItemOrganograma.objects.get(
            idimportacaoplanilha_id=importacao_id,
            idpai__isnull=True,
        )
        organ = ItemOrganograma.objects.get(
            idimportacaoplanilha_id=importacao_id,
            idpai=root,
        )
        parent = root if parent_type == "organ" else organ
        payload = {
            "idimportacaoplanilha_id": importacao_id,
            "idpai_id": parent.pk,
            "strnome": "Nova estrutura",
            "strsigla": "NOVA",
        }
        payload[field] = value

        response = client_gestor_carga_lot.post(
            ITENS_ORGANOGRAMA_URL,
            payload,
            format="json",
        )

        assert response.status_code == 400
        assert field in response.data["errors"]

    def test_crud_rejeita_sigla_repetida_em_ramos_do_mesmo_orgao(
        self,
        client_gestor_carga_lot,
        patriarca,
        tipo_organograma,
    ):
        upload = client_gestor_carga_lot.post(
            f"{IMPORTACOES_URL}upload/",
            {
                "idpatriarca_id": patriarca.pk,
                "idtipocarga_id": tipo_organograma.pk,
                "arquivo": arquivo_planilha(),
            },
            format="multipart",
        )
        importacao_id = upload.data["idimportacaoplanilha"]
        root = ItemOrganograma.objects.get(
            idimportacaoplanilha_id=importacao_id,
            idpai__isnull=True,
        )
        organ = ItemOrganograma.objects.get(
            idimportacaoplanilha_id=importacao_id,
            idpai=root,
        )

        branches = []
        for suffix in ("A", "B"):
            response = client_gestor_carga_lot.post(
                ITENS_ORGANOGRAMA_URL,
                {
                    "idimportacaoplanilha_id": importacao_id,
                    "idpai_id": organ.pk,
                    "strnome": f"Diretoria {suffix}",
                    "strsigla": f"DIR{suffix}",
                },
                format="json",
            )
            assert response.status_code == 201
            branches.append(response.data["iditemorganograma"])

        first = client_gestor_carga_lot.post(
            ITENS_ORGANOGRAMA_URL,
            {
                "idimportacaoplanilha_id": importacao_id,
                "idpai_id": branches[0],
                "strnome": "Unidade Comum A",
                "strsigla": "COMUM",
            },
            format="json",
        )
        duplicate = client_gestor_carga_lot.post(
            ITENS_ORGANOGRAMA_URL,
            {
                "idimportacaoplanilha_id": importacao_id,
                "idpai_id": branches[1],
                "strnome": "Unidade Comum B",
                "strsigla": "COMUM",
            },
            format="json",
        )

        assert first.status_code == 201
        assert duplicate.status_code == 400
        assert "strsigla" in duplicate.data["errors"]

    def test_exclusao_de_subarvore_exige_confirmacao_e_informa_impacto(
        self,
        client_gestor_carga_lot,
        patriarca,
        tipo_organograma,
    ):
        upload = client_gestor_carga_lot.post(
            f"{IMPORTACOES_URL}upload/",
            {
                "idpatriarca_id": patriarca.pk,
                "idtipocarga_id": tipo_organograma.pk,
                "arquivo": arquivo_planilha(),
            },
            format="multipart",
        )
        importacao_id = upload.data["idimportacaoplanilha"]
        root = ItemOrganograma.objects.get(
            idimportacaoplanilha_id=importacao_id,
            idpai__isnull=True,
        )

        impact_response = client_gestor_carga_lot.get(
            f"{ITENS_ORGANOGRAMA_URL}{root.pk}/impacto-exclusao/"
        )
        assert impact_response.status_code == 200
        assert impact_response.data["total_descendentes"] == 1
        assert impact_response.data["total_excluidos"] == 2

        unconfirmed = client_gestor_carga_lot.delete(
            f"{ITENS_ORGANOGRAMA_URL}{root.pk}/"
        )
        assert unconfirmed.status_code == 409
        assert (
            unconfirmed.data["errors"]["code"]
            == "confirmacao_exclusao_subarvore_obrigatoria"
        )
        assert unconfirmed.data["errors"]["impacto"]["total_excluidos"] == 2
        assert ItemOrganograma.objects.filter(
            idimportacaoplanilha_id=importacao_id
        ).count() == 2

        confirmed = client_gestor_carga_lot.delete(
            f"{ITENS_ORGANOGRAMA_URL}{root.pk}/?confirmar_subarvore=true",
        )
        assert confirmed.status_code == 204
        assert not ItemOrganograma.objects.filter(
            idimportacaoplanilha_id=importacao_id
        ).exists()
        process = ProcessoCargaInicial.objects.get(idpatriarca=patriarca)
        history = process.historico.get(
            stracao="SUBARVORE_ORGANOGRAMA_EXCLUIDA"
        )
        assert history.jsondetalhes["impacto"]["total_excluidos"] == 2

    def test_edicao_de_item_desatualiza_lotacao(
        self,
        client_gestor_carga_lot,
        patriarca,
        tipo_organograma,
        tipo_lotacao,
    ):
        organogram = client_gestor_carga_lot.post(
            f"{IMPORTACOES_URL}upload/",
            {
                "idpatriarca_id": patriarca.pk,
                "idtipocarga_id": tipo_organograma.pk,
                "arquivo": arquivo_planilha(),
            },
            format="multipart",
        )
        organogram_id = organogram.data["idimportacaoplanilha"]
        lotation = client_gestor_carga_lot.post(
            f"{IMPORTACOES_URL}upload/",
            {
                "idpatriarca_id": patriarca.pk,
                "idtipocarga_id": tipo_lotacao.pk,
                "idimportacaoorganograma_id": organogram_id,
                "arquivo": arquivo_planilha(
                    nome="lotacao.xlsx",
                    conteudo=conteudo_lotacao(),
                ),
            },
            format="multipart",
        )
        assert lotation.status_code == 201
        item = ItemOrganograma.objects.filter(
            idimportacaoplanilha_id=organogram_id
        ).exclude(idpai=None).get()

        response = client_gestor_carga_lot.patch(
            f"{ITENS_ORGANOGRAMA_URL}{item.pk}/",
            {"strnome": "Secretaria de Administração Atualizada"},
            format="json",
        )

        assert response.status_code == 200
        process = ProcessoCargaInicial.objects.get(idpatriarca=patriarca)
        assert process.strstatuslotacao == StatusLotacaoProcesso.DESATUALIZADA

    def test_importacao_impede_exclusao_do_patriarca(
        self,
        client_gestor_carga_lot,
        patriarca,
        tipo_organograma,
    ):
        ImportacaoPlanilha.objects.create(
            idpatriarca=patriarca,
            idtipocarga=tipo_organograma,
            idstatusprogresso=patriarca.idstatusprogresso,
            strnomearquivo="organograma.xlsx",
            strtipomime="application/octet-stream",
            inttamanhoarquivo=3,
            strhasharquivo=hashlib.sha256(b"abc").hexdigest(),
            binconteudo=b"abc",
        )
        response = client_gestor_carga_lot.delete(
            f"{PATRIARCAS_URL}{patriarca.pk}/"
        )
        assert response.status_code == 400
        assert Patriarca.objects.filter(pk=patriarca.pk).exists()


@pytest.mark.django_db(transaction=True)
class TestEnvioOrganogramaProdest:
    class FakeCargaRHClient:
        submitted_payload = None
        submitted_dry_run = None

        def __init__(self, **kwargs):
            self.context = kwargs

        def submit_organogram(self, patriarch_id, payload, *, dry_run):
            self.__class__.submitted_payload = payload
            self.__class__.submitted_dry_run = dry_run
            return "aaaaaaaa-aaaa-4aaa-8aaa-aaaaaaaaaaaa"

        def get_status(self, carga_id):
            return {
                "guid": str(carga_id),
                "dryRun": self.__class__.submitted_dry_run,
                "finalizada": True,
                "mensagemErro": None,
                "etapas": [
                    {
                        "etapa": 3,
                        "status": 3,
                        "tentativas": 1,
                        "mensagemErro": None,
                    }
                ],
            }

        def get_report(self, carga_id):
            return {
                "cargaId": str(carga_id),
                "tipo": "organograma",
                "modo": "Simulacao",
                "parcial": False,
                "totais": {"incluidos": 2},
            }

        def list_initial_imports(self, patriarch_id, *, limit):
            return [
                {
                    "guid": "aaaaaaaa-aaaa-4aaa-8aaa-aaaaaaaaaaaa",
                    "patriarcaId": str(patriarch_id),
                    "finalizada": True,
                    "dryRun": True,
                }
            ][:limit]

    def _valid_import(self, client, patriarca, tipo_organograma):
        response = client.post(
            f"{IMPORTACOES_URL}upload/",
            {
                "idpatriarca_id": patriarca.pk,
                "idtipocarga_id": tipo_organograma.pk,
                "arquivo": arquivo_planilha(
                    conteudo=conteudo_organograma(
                        linhas=[
                            (1, "Prefeitura", "PMVV"),
                            ("1.1", "Secretaria de Administração", "SEMAD"),
                            ("1.1.1", "Secretaria de Administração", "SEMAD"),
                            ("1.1.1.1", "Gerência de Pessoas", "GEP"),
                        ]
                    )
                ),
            },
            format="multipart",
        )
        assert response.status_code == 201
        return response.data["idimportacaoplanilha"]

    def test_simula_acompanha_e_obtem_relatorio(
        self,
        client_gestor_carga_lot,
        patriarca,
        tipo_organograma,
        monkeypatch,
    ):
        monkeypatch.setattr(
            "apps.carga_org_lot.services.organogram_submission.CargaRHClient",
            self.FakeCargaRHClient,
        )
        importacao_id = self._valid_import(
            client_gestor_carga_lot,
            patriarca,
            tipo_organograma,
        )

        response = client_gestor_carga_lot.post(
            ENVIOS_ORGANOGRAMA_URL,
            {
                "idimportacaoplanilha_id": importacao_id,
                "dry_run": True,
            },
            format="json",
        )

        assert response.status_code == 202
        envio_id = response.data["idenvioorganograma"]
        assert response.data["modo"] == "SIMULACAO"
        assert response.data["strstatus"] == StatusEnvioOrganograma.SOLICITANDO
        assert "strtoken" not in response.data
        payload = self.FakeCargaRHClient.submitted_payload
        assert payload["patriarcaId"] == str(patriarca.idexternopatriarca)
        assert payload["orgaos"] == [
            {
                "nome": "SECRETARIA DE ADMINISTRAÇÃO",
                "sigla": "SEMAD",
                "unidades": [
                    {
                        "nome": "SECRETARIA DE ADMINISTRAÇÃO",
                        "sigla": "SEMAD",
                        "unidadePai": None,
                    },
                    {
                        "nome": "GERÊNCIA DE PESSOAS",
                        "sigla": "GEP",
                        "unidadePai": "SEMAD",
                    },
                ],
            }
        ]
        process = ProcessoCargaInicial.objects.get(idpatriarca=patriarca)
        assert process.strfaseatual == FaseProcessoCargaInicial.SIMULANDO
        assert (
            process.strstatusorganograma
            == StatusOrganogramaProcesso.PROCESSANDO_NA_PRODEST
        )

        status_response = client_gestor_carga_lot.post(
            f"{ENVIOS_ORGANOGRAMA_URL}{envio_id}/consultar-status/"
        )
        assert status_response.status_code == 200
        assert status_response.data["strstatus"] == StatusEnvioOrganograma.CONCLUIDO
        assert status_response.data["finalizada_prodest"] is True
        assert status_response.data["relatorio_disponivel"] is True
        assert status_response.data["etapas_prodest"][0]["etapa_nome"] == (
            "Organograma"
        )
        assert status_response.data["etapas_prodest"][0]["status_nome"] == (
            "Finalizada"
        )
        process.refresh_from_db()
        assert process.strfaseatual == FaseProcessoCargaInicial.PRONTO_PARA_ENVIO
        assert process.strstatusorganograma == StatusOrganogramaProcesso.PRONTO

        report_response = client_gestor_carga_lot.post(
            f"{ENVIOS_ORGANOGRAMA_URL}{envio_id}/consultar-relatorio/"
        )
        assert report_response.status_code == 200
        assert report_response.data["jsonrelatorioprodest"]["parcial"] is False

    def test_envio_definitivo_concluido_libera_preparacao_das_lotacoes(
        self,
        client_gestor_carga_lot,
        patriarca,
        tipo_organograma,
        monkeypatch,
    ):
        monkeypatch.setattr(
            "apps.carga_org_lot.services.organogram_submission.CargaRHClient",
            self.FakeCargaRHClient,
        )
        importacao_id = self._valid_import(
            client_gestor_carga_lot,
            patriarca,
            tipo_organograma,
        )
        created = client_gestor_carga_lot.post(
            ENVIOS_ORGANOGRAMA_URL,
            {
                "idimportacaoplanilha_id": importacao_id,
                "dry_run": False,
            },
            format="json",
        )
        assert created.status_code == 202

        completed = client_gestor_carga_lot.post(
            f"{ENVIOS_ORGANOGRAMA_URL}"
            f"{created.data['idenvioorganograma']}/consultar-status/"
        )

        assert completed.status_code == 200
        process = ProcessoCargaInicial.objects.get(idpatriarca=patriarca)
        assert (
            process.strfaseatual
            == FaseProcessoCargaInicial.PREPARANDO_LOTACOES
        )
        assert process.strstatusorganograma == StatusOrganogramaProcesso.CONCLUIDO
        assert process.strstatuslotacao == StatusLotacaoProcesso.NAO_INICIADA

    def test_rejeita_segunda_carga_inicial_ativa_do_mesmo_patriarca(
        self,
        client_gestor_carga_lot,
        patriarca,
        tipo_organograma,
        monkeypatch,
    ):
        monkeypatch.setattr(
            "apps.carga_org_lot.services.organogram_submission.CargaRHClient",
            self.FakeCargaRHClient,
        )
        importacao_id = self._valid_import(
            client_gestor_carga_lot,
            patriarca,
            tipo_organograma,
        )
        payload = {
            "idimportacaoplanilha_id": importacao_id,
            "dry_run": True,
        }
        first = client_gestor_carga_lot.post(
            ENVIOS_ORGANOGRAMA_URL,
            payload,
            format="json",
        )
        payload["dry_run"] = False
        second = client_gestor_carga_lot.post(
            ENVIOS_ORGANOGRAMA_URL,
            payload,
            format="json",
        )

        assert first.status_code == 202
        assert second.status_code == 400
        assert EnvioOrganogramaProdest.objects.count() == 1

    def test_consulta_ultimas_cargas_iniciais_na_prodest_e_concilia_local(
        self,
        client_gestor_carga_lot,
        patriarca,
        tipo_organograma,
        monkeypatch,
    ):
        monkeypatch.setattr(
            "apps.carga_org_lot.services.organogram_submission.CargaRHClient",
            self.FakeCargaRHClient,
        )
        importacao_id = self._valid_import(
            client_gestor_carga_lot,
            patriarca,
            tipo_organograma,
        )
        created = client_gestor_carga_lot.post(
            ENVIOS_ORGANOGRAMA_URL,
            {"idimportacaoplanilha_id": importacao_id, "dry_run": True},
            format="json",
        )
        monkeypatch.setattr(
            "apps.carga_org_lot.views.CargaRHClient",
            self.FakeCargaRHClient,
        )

        response = client_gestor_carga_lot.get(
            f"{ENVIOS_ORGANOGRAMA_URL}ultimas-prodest/",
            {"patriarca": patriarca.pk, "limit": 10},
        )

        assert response.status_code == 200
        assert response.data["patriarca_id"] == patriarca.pk
        assert len(response.data["cargas"]) == 1
        assert response.data["cargas"][0]["idenvioorganograma_local"] == (
            created.data["idenvioorganograma"]
        )

    def test_consulta_ultimas_cargas_rejeita_limite_acima_de_dez(
        self,
        client_gestor_carga_lot,
        patriarca,
    ):
        response = client_gestor_carga_lot.get(
            f"{ENVIOS_ORGANOGRAMA_URL}ultimas-prodest/",
            {"patriarca": patriarca.pk, "limit": 11},
        )

        assert response.status_code == 400

    def test_falha_de_comunicacao_fica_registrada_sem_token(
        self,
        client_gestor_carga_lot,
        patriarca,
        tipo_organograma,
        monkeypatch,
    ):
        class FailingClient:
            def __init__(self, **kwargs):
                self.context = kwargs

            def submit_organogram(self, patriarch_id, payload, *, dry_run):
                raise CargaRHRequestError(
                    "Não foi possível estabelecer comunicação com a PRODEST."
                )

        monkeypatch.setattr(
            "apps.carga_org_lot.services.organogram_submission.CargaRHClient",
            FailingClient,
        )
        importacao_id = self._valid_import(
            client_gestor_carga_lot,
            patriarca,
            tipo_organograma,
        )

        response = client_gestor_carga_lot.post(
            ENVIOS_ORGANOGRAMA_URL,
            {
                "idimportacaoplanilha_id": importacao_id,
                "dry_run": True,
            },
            format="json",
        )

        assert response.status_code == 202
        envio = EnvioOrganogramaProdest.objects.get()
        assert response.data["idenvioorganograma"] == envio.pk
        assert envio.strstatus == StatusEnvioOrganograma.ERRO
        assert envio.idexternocarga is None
        assert envio.jsonstatusprodest == {
            "falha_solicitacao": {
                "http_status": None,
                "problem": {},
            }
        }
        assert "token" not in str(response.data).lower()
        process = ProcessoCargaInicial.objects.get(idpatriarca=patriarca)
        assert process.strstatusorganograma == StatusOrganogramaProcesso.ERRO

    def test_status_de_erro_da_prodest_atualiza_processo(
        self,
        client_gestor_carga_lot,
        patriarca,
        tipo_organograma,
        monkeypatch,
    ):
        class ErrorStatusClient(self.FakeCargaRHClient):
            def get_status(self, carga_id):
                return {
                    "guid": str(carga_id),
                    "dryRun": True,
                    "finalizada": True,
                    "mensagemErro": None,
                    "etapas": [
                        {
                            "etapa": 3,
                            "status": 4,
                            "tentativas": 1,
                            "mensagemErro": "Sigla rejeitada pela PRODEST.",
                        }
                    ],
                }

        monkeypatch.setattr(
            "apps.carga_org_lot.services.organogram_submission.CargaRHClient",
            ErrorStatusClient,
        )
        importacao_id = self._valid_import(
            client_gestor_carga_lot,
            patriarca,
            tipo_organograma,
        )
        created = client_gestor_carga_lot.post(
            ENVIOS_ORGANOGRAMA_URL,
            {
                "idimportacaoplanilha_id": importacao_id,
                "dry_run": True,
            },
            format="json",
        )

        response = client_gestor_carga_lot.post(
            f"{ENVIOS_ORGANOGRAMA_URL}"
            f"{created.data['idenvioorganograma']}/consultar-status/"
        )

        assert response.status_code == 200
        assert response.data["strstatus"] == StatusEnvioOrganograma.ERRO
        assert "Sigla rejeitada" in response.data["strmensagemerro"]
        process = ProcessoCargaInicial.objects.get(idpatriarca=patriarca)
        assert process.strstatusorganograma == StatusOrganogramaProcesso.ERRO

    def test_conclusao_nao_valida_organograma_alterado_durante_processamento(
        self,
        client_gestor_carga_lot,
        patriarca,
        tipo_organograma,
        monkeypatch,
    ):
        monkeypatch.setattr(
            "apps.carga_org_lot.services.organogram_submission.CargaRHClient",
            self.FakeCargaRHClient,
        )
        importacao_id = self._valid_import(
            client_gestor_carga_lot,
            patriarca,
            tipo_organograma,
        )
        created = client_gestor_carga_lot.post(
            ENVIOS_ORGANOGRAMA_URL,
            {
                "idimportacaoplanilha_id": importacao_id,
                "dry_run": True,
            },
            format="json",
        )
        unit = ItemOrganograma.objects.get(
            idimportacaoplanilha_id=importacao_id,
            strsigla="GEP",
        )
        changed = client_gestor_carga_lot.patch(
            f"{ITENS_ORGANOGRAMA_URL}{unit.pk}/",
            {"strnome": "Gerência de Pessoas Atualizada"},
            format="json",
        )
        assert changed.status_code == 200

        response = client_gestor_carga_lot.post(
            f"{ENVIOS_ORGANOGRAMA_URL}"
            f"{created.data['idenvioorganograma']}/consultar-status/"
        )

        assert response.status_code == 200
        assert (
            response.data["strstatus"]
            == StatusEnvioOrganograma.CONCLUIDO_DESATUALIZADO
        )
        process = ProcessoCargaInicial.objects.get(idpatriarca=patriarca)
        assert (
            process.strfaseatual
            == FaseProcessoCargaInicial.PREPARANDO_ORGANOGRAMA
        )
        assert (
            process.strstatusorganograma
            == StatusOrganogramaProcesso.EM_PREPARACAO
        )

    def test_configuracao_nao_expoe_credenciais(self, client_gestor_carga_lot):
        response = client_gestor_carga_lot.get(
            f"{ENVIOS_ORGANOGRAMA_URL}configuracao/"
        )

        assert response.status_code == 200
        assert set(response.data) == {
            "configurado",
            "ambiente",
            "ausentes",
            "invalidas",
        }
        serialized = str(response.data).lower()
        assert "client_secret" not in serialized
        assert "access_token" not in serialized

    def test_conexao_autenticada_nao_expoe_token(
        self,
        client_gestor_carga_lot,
        monkeypatch,
    ):
        class AuthenticatedClient:
            def __init__(self, **kwargs):
                self.context = kwargs

            def test_connection(self):
                return True

        monkeypatch.setattr(
            "apps.carga_org_lot.views.CargaRHClient",
            AuthenticatedClient,
        )

        response = client_gestor_carga_lot.post(
            f"{ENVIOS_ORGANOGRAMA_URL}testar-conexao/"
        )

        assert response.status_code == 200
        assert response.data == {
            "conectado": True,
            "ambiente": settings.CARGARH_ENVIRONMENT,
            "token_armazenado": "cache",
            "token_exposto": False,
        }


@pytest.mark.django_db(transaction=True)
class TestEventoIntegracaoProdest:
    def test_lista_eventos_seguros_e_filtra_por_patriarca(
        self,
        client_gestor_carga_lot,
        gestor_carga_lot,
        patriarca,
    ):
        event = EventoIntegracaoProdest.objects.create(
            idpatriarca=patriarca,
            stroperacao="OBTER_TOKEN",
            strmetodo="POST",
            strendpoint="https://auth.example/connect/token",
            inttentativa=1,
            strresultado="ERRO_HTTP",
            inthttpstatus=401,
            intduracaoms=125,
            strmensagem="O Acesso Cidadão recusou a autenticação.",
            jsondetalhes={"title": "Unauthorized", "status": 401},
            created_by_id=gestor_carga_lot.pk,
            created_by_name=gestor_carga_lot.username,
            updated_by_id=gestor_carga_lot.pk,
            updated_by_name=gestor_carga_lot.username,
        )

        response = client_gestor_carga_lot.get(
            EVENTOS_INTEGRACAO_URL,
            {"patriarca": patriarca.pk},
        )

        assert response.status_code == 200
        assert response.data["count"] == 1
        item = response.data["results"][0]
        assert item["ideventointegracao"] == event.pk
        assert item["operacao_nome"] == "Autenticação no Acesso Cidadão"
        assert item["resultado_nome"] == "Recusada pelo serviço externo"
        assert item["sucesso"] is False
        assert "token" not in str(item["jsondetalhes"]).lower()

    def test_rejeita_filtro_numerico_invalido(
        self,
        client_gestor_carga_lot,
    ):
        response = client_gestor_carga_lot.get(
            EVENTOS_INTEGRACAO_URL,
            {"patriarca": "invalido"},
        )

        assert response.status_code == 400


@pytest.mark.django_db(transaction=True)
class TestCargaEHistorico:
    def test_crud_carga_sem_expor_token(
        self,
        client_gestor_carga_lot,
        tipo_organograma,
        status_processando,
    ):
        create_response = client_gestor_carga_lot.post(
            CARGAS_URL,
            {
                "strtoken": "segredo",
                "idtipocarga_id": tipo_organograma.pk,
                "idstatusprogresso_id": status_processando.pk,
            },
            format="json",
        )
        assert create_response.status_code == 201
        pk = create_response.data["idtokenenviocarga"]
        assert create_response.data["possui_token"] is True
        assert "strtoken" not in create_response.data

        assert client_gestor_carga_lot.get(CARGAS_URL).status_code == 200
        assert client_gestor_carga_lot.get(f"{CARGAS_URL}{pk}/").status_code == 200
        assert client_gestor_carga_lot.patch(
            f"{CARGAS_URL}{pk}/",
            {"idstatusprogresso_id": status_processando.pk},
            format="json",
        ).status_code == 200
        assert client_gestor_carga_lot.delete(f"{CARGAS_URL}{pk}/").status_code == 204

    def test_crud_historico(
        self,
        client_gestor_carga_lot,
        carga,
        status_carga_recebida,
    ):
        create_response = client_gestor_carga_lot.post(
            HISTORICO_URL,
            {
                "idtokenenviocarga_id": carga.pk,
                "idstatuscarga_id": status_carga_recebida.pk,
                "strmensagem": "Arquivo recebido.",
            },
            format="json",
        )
        assert create_response.status_code == 201
        pk = create_response.data["iddetalhestatuscarga"]

        assert client_gestor_carga_lot.get(HISTORICO_URL).status_code == 200
        assert client_gestor_carga_lot.get(f"{HISTORICO_URL}{pk}/").status_code == 200
        patch_response = client_gestor_carga_lot.patch(
            f"{HISTORICO_URL}{pk}/",
            {"strmensagem": "Atualizado."},
            format="json",
        )
        assert patch_response.status_code == 200
        assert patch_response.data["strmensagem"] == "Atualizado."
        assert client_gestor_carga_lot.delete(
            f"{HISTORICO_URL}{pk}/"
        ).status_code == 204


@pytest.mark.django_db(transaction=True)
class TestAutorizacao:
    def test_portal_admin_pode_listar(self, client_portal_admin):
        assert client_portal_admin.get(PATRIARCAS_URL).status_code == 200

    def test_sem_role_e_negado(self, client_sem_role_carga):
        response = client_sem_role_carga.get(PATRIARCAS_URL)
        assert response.status_code in (401, 403)
